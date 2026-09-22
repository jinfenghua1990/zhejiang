"""读取吉客云客户端的官方 XLSX / CSV 导出文件。

文件格式会随客户启用的模块、报表和吉客云版本变化。这里仅做安全读取、
表头识别和报表类型判断；不把无法确认的列硬映射为业务字段。
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable


ALLOWED_EXTENSIONS = {".xlsx", ".csv"}
MAX_XLSX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024


@dataclass(frozen=True)
class ParsedJackyunExport:
    report_type: str
    sheet_name: str
    headers: list[str]
    rows: list[dict[str, str]]

    @property
    def status(self) -> str:
        return "parsed" if self.report_type != "unknown" else "needs_mapping"


def _normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip().lower()
    return re.sub(r"[\s_\-()（）\[\]【】/\\:：]+", "", text)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value).strip()


def _matches(headers: set[str], aliases: Iterable[str]) -> bool:
    for alias in aliases:
        key = _normalize_header(alias)
        if any(key == header or key in header for header in headers):
            return True
    return False


def detect_report_type(headers: list[str]) -> str:
    """根据官方导出常见中文列名做保守判断；不满足关键条件就返回 unknown。"""
    normalized = {_normalize_header(header) for header in headers if header}
    has_product_ref = _matches(normalized, ("货品编号", "货品编码", "商品编号", "商品编码", "SKU编码", "规格编码"))
    has_order_ref = _matches(normalized, ("订单编号", "订单号", "线上订单编号", "原始订单号"))

    if _matches(normalized, ("售后单号", "退款单号", "退货单号", "换货单号")):
        return "aftersales"
    if _matches(normalized, ("入库单号", "入库单据号", "入库单编号")):
        return "inbound"
    if (
        _matches(normalized, ("申请单号", "入库申请单号", "入库申请编号"))
        and _matches(normalized, ("入库数量", "申请数量", "入库类型"))
        and has_product_ref
    ):
        return "inbound"
    if _matches(normalized, ("出库单号", "出库单据号", "出库单编号", "发货单号")):
        return "outbound"
    if _matches(normalized, ("采购单号", "采购订单号", "采购单编号")):
        return "purchase"
    if has_order_ref and _matches(normalized, ("订单状态", "下单时间", "付款时间", "店铺名称", "平台名称")):
        return "sales"
    if has_product_ref and _matches(normalized, ("库存数量", "可用库存", "库存量", "可用量", "实际库存")):
        return "inventory"
    if has_product_ref and _matches(normalized, ("货品名称", "商品名称", "规格名称", "货品名")):
        return "products"
    if _matches(normalized, ("仓库编号", "仓库编码")) and _matches(normalized, ("仓库名称", "仓库名")):
        return "warehouses"
    return "unknown"


def _header_score(values: list[object]) -> int:
    nonempty = [_normalize_header(v) for v in values if _normalize_header(v)]
    if not nonempty:
        return -1
    known = sum(
        _matches({value}, (
            "货品", "商品", "订单", "库存", "采购", "入库", "出库", "售后", "仓库",
            "编号", "名称", "时间", "状态", "数量", "金额",
        ))
        for value in nonempty
    )
    return known * 10 + min(len(nonempty), 20)


def _find_header_row(table: list[list[object]]) -> int:
    candidates = table[:20]
    if not candidates:
        raise ValueError("文件没有可读取的工作表内容")
    index = max(range(len(candidates)), key=lambda item: _header_score(candidates[item]))
    score = _header_score(candidates[index])
    if score < 1:
        raise ValueError("未找到有效表头")
    return index


def _unique_headers(values: list[object]) -> list[str]:
    result: list[str] = []
    used: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = _cell_text(value) or f"列{index}"
        count = used.get(base, 0) + 1
        used[base] = count
        result.append(base if count == 1 else f"{base}#{count}")
    return result


def _rows_to_records(
    rows_iter: Iterable[Iterable[object]], *, header_index: int, max_rows: int
) -> tuple[list[str], list[dict[str, str]]]:
    headers: list[str] | None = None
    rows: list[dict[str, str]] = []
    for source_index, source_row in enumerate(rows_iter):
        if source_index < header_index:
            continue
        raw_row = list(source_row)
        if source_index == header_index:
            headers = _unique_headers(raw_row)
            continue
        if headers is None:
            raise ValueError("未找到有效表头")
        payload: dict[str, str] = {}
        for index, raw_value in enumerate(raw_row[:len(headers)]):
            value = _cell_text(raw_value)
            if value:
                payload[headers[index]] = value
        if not payload:
            continue
        if len(rows) >= max_rows:
            raise ValueError(f"数据行超过单次导入上限（{max_rows} 行）")
        rows.append(payload)
    if headers is None:
        raise ValueError("未找到有效表头")
    return headers, rows


def _validate_xlsx(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            total = sum(member.file_size for member in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ValueError("文件不是有效的 XLSX") from exc
    if total > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ValueError("XLSX 解压后的内容超过安全上限")


def _parse_xlsx(content: bytes, *, max_rows: int) -> ParsedJackyunExport:
    from openpyxl import load_workbook

    _validate_xlsx(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取 XLSX 文件") from exc

    best: tuple[int, str, list[list[object]]] | None = None
    try:
        for worksheet in workbook.worksheets:
            sample = [list(row) for row in worksheet.iter_rows(max_row=20, values_only=True)]
            score = max((_header_score(row) for row in sample), default=-1)
            if best is None or score > best[0]:
                best = (score, worksheet.title, sample)
        if best is None or best[0] < 1:
            raise ValueError("未找到含表头的工作表")
        _, sheet_name, _ = best
        worksheet = workbook[sheet_name]
        header_index = _find_header_row(best[2])
        headers, rows = _rows_to_records(
            worksheet.iter_rows(values_only=True), header_index=header_index, max_rows=max_rows
        )
    finally:
        workbook.close()

    return ParsedJackyunExport(detect_report_type(headers), sheet_name, headers, rows)


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "utf-8"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("CSV 编码无法识别，请导出 UTF-8 或 GBK/GB18030 文件")


def _parse_csv(content: bytes, *, max_rows: int) -> ParsedJackyunExport:
    text = _decode_csv(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    table = [list(row) for row in csv.reader(io.StringIO(text), dialect=dialect)]
    header_index = _find_header_row(table)
    headers, rows = _rows_to_records(table, header_index=header_index, max_rows=max_rows)
    return ParsedJackyunExport(detect_report_type(headers), "CSV", headers, rows)


def parse_jackyun_export(content: bytes, original_name: str, *, max_rows: int) -> ParsedJackyunExport:
    if not content:
        raise ValueError("空文件")
    extension = Path(original_name or "").suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("仅支持吉客云导出的 XLSX 或 CSV 文件；请不要上传客户端程序或 PDF")
    if extension == ".xlsx":
        return _parse_xlsx(content, max_rows=max_rows)
    return _parse_csv(content, max_rows=max_rows)
