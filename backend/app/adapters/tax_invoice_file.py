"""读取税务系统官方导出的发票清单。

这里只做保守的表头识别和字段抽取。原始行由导入服务完整保存；没有明确表头或
发票号码的行进入待核对，不把缺失数据推断成“未开票”。
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field as dataclass_field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable


ALLOWED_EXTENSIONS = {".xlsx", ".csv"}
MAX_XLSX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "invoice_number": ("发票号码", "发票号", "发票编号", "数电发票号码", "电子发票号码", "invoice number", "invoice_no"),
    "invoice_code": ("发票代码", "票代码", "数电发票代码", "电子发票代码", "invoice code", "invoice_code"),
    "invoice_type": ("发票类型", "发票种类", "票种", "发票票种", "invoice type"),
    "issue_date": ("开票日期", "开具日期", "开票时间", "发票日期", "issue date", "invoice date"),
    "seller_name": ("销售方名称", "销方名称", "开票方名称", "卖方名称", "seller name"),
    "seller_tax_id": ("销售方识别号", "销方识别号", "开票方税号", "卖方税号", "销售方税号", "seller tax id"),
    "buyer_name": ("购买方名称", "购方名称", "受票方名称", "买方名称", "buyer name"),
    "buyer_tax_id": ("购买方识别号", "购方识别号", "受票方税号", "买方税号", "购买方税号", "buyer tax id"),
    "amount_excl_tax": ("不含税金额", "金额(不含税)", "金额（不含税）", "金额", "amount excl tax", "amount"),
    "tax_amount": ("税额", "税金", "合计税额", "tax amount", "tax"),
    "total_amount": ("价税合计", "价税合计金额", "含税金额", "金额(含税)", "金额（含税）", "发票金额", "total amount"),
    "currency": ("币种", "货币", "currency"),
    "direction": ("发票方向", "进销项", "发票属性", "发票来源", "invoice direction"),
    "status": ("发票状态", "有效状态", "作废标志", "是否作废", "状态", "invoice status"),
    "related_order_ref": ("关联订单号", "订单号", "订单编号", "采购单号", "采购订单号", "业务单号", "来源单号", "原始订单号"),
}


@dataclass(frozen=True)
class ParsedTaxInvoiceExport:
    sheet_name: str
    headers: list[str]
    rows: list[dict[str, str]]
    mapping: dict[str, str]
    direction_hint: str
    ignored_row_count: int = 0
    summary_sheet_name: str = ""
    summary_headers: list[str] = dataclass_field(default_factory=list)
    summary_mapping: dict[str, str] = dataclass_field(default_factory=dict)
    summary_validation: dict[str, Any] = dataclass_field(default_factory=dict)
    detail_sheet_name: str = ""
    detail_headers: list[str] = dataclass_field(default_factory=list)
    detail_mapping: dict[str, str] = dataclass_field(default_factory=dict)
    detail_rows: list[dict[str, str]] = dataclass_field(default_factory=list)


def normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip().lower()
    return re.sub(r"[\s_\-()（）\[\]【】/\\:：]+", "", text)


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value).strip()


def _decimal_value(value: object) -> Decimal | None:
    text = cell_text(value).replace(",", "").replace("￥", "").replace("¥", "")
    if not text or text in {"-", "--", "/", "—", "――", "－"}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1].strip()
    try:
        number = Decimal(text)
    except Exception:
        return None
    return -number if negative else number


def resolve_mapping(headers: list[str]) -> dict[str, str]:
    normalized = [(header, normalize_header(header)) for header in headers if header]
    mapping: dict[str, str] = {}
    for field, aliases in FIELD_ALIASES.items():
        alias_keys = [normalize_header(alias) for alias in aliases]
        if field == "invoice_number":
            # 数电发票同时可能保留一个空的“发票号码”列；优先使用真正有值的
            # 数电号码列，服务层仍会兼容旧批次的原始映射。
            for preferred in ("数电发票号码", "电子发票号码", "发票号码", "发票号", "发票编号"):
                preferred_key = normalize_header(preferred)
                exact_preferred = next(
                    (header for header, key in normalized if key == preferred_key), None
                )
                if exact_preferred:
                    mapping[field] = exact_preferred
                    break
            if field in mapping:
                continue
        # 先精确命中，避免“金额”之类的短别名抢到相邻字段。
        exact = next((header for header, key in normalized if key in alias_keys), None)
        if exact:
            mapping[field] = exact
            continue
        for alias_key in alias_keys:
            if len(alias_key) < 3:
                continue
            partial = next((header for header, key in normalized if alias_key in key), None)
            if partial:
                mapping[field] = partial
                break
    return mapping


def _header_score(values: list[object]) -> int:
    headers = [cell_text(value) for value in values]
    normalized = [normalize_header(header) for header in headers if header]
    if not normalized:
        return -1
    tokens = ("发票", "开票", "税额", "价税", "税号", "购方", "销方", "金额", "红字", "作废", "票种")
    score = sum(any(token in header for token in tokens) for header in normalized) * 10
    mapping = resolve_mapping(headers)
    score += len(mapping) * 2
    return score + min(len(normalized), 20)


def _find_header_row(table: list[list[object]]) -> int:
    candidates = table[:20]
    if not candidates:
        raise ValueError("文件没有可读取的内容")
    index = max(range(len(candidates)), key=lambda item: _header_score(candidates[item]))
    if _header_score(candidates[index]) < 5:
        raise ValueError("未找到可识别的发票表头")
    return index


def _unique_headers(values: list[object]) -> list[str]:
    result: list[str] = []
    used: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = cell_text(value) or f"列{index}"
        count = used.get(base, 0) + 1
        used[base] = count
        result.append(base if count == 1 else f"{base}#{count}")
    return result


def _rows_to_records(
    rows_iter: Iterable[Iterable[object]], *, header_index: int, max_rows: int
) -> tuple[list[str], list[dict[str, str]]]:
    headers: list[str] | None = None
    rows: list[dict[str, str]] = []
    for source_index, source_row in enumerate(rows_iter):
        if source_index < header_index:
            continue
        raw_row = list(source_row)
        if source_index == header_index:
            headers = _unique_headers(raw_row)
            continue
        if headers is None:
            raise ValueError("未找到有效表头")
        payload: dict[str, str] = {}
        for index, raw_value in enumerate(raw_row[:len(headers)]):
            value = cell_text(raw_value)
            if value:
                payload[headers[index]] = value
        if not payload:
            continue
        if len(rows) >= max_rows:
            raise ValueError(f"数据行超过单次导入上限（{max_rows} 行）")
        rows.append(payload)
    if headers is None:
        raise ValueError("未找到有效表头")
    return headers, rows


def _validate_xlsx(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            total = sum(member.file_size for member in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ValueError("文件不是有效的 XLSX") from exc
    if total > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ValueError("XLSX 解压后的内容超过安全上限")


def _sheet_totals(
    rows: list[dict[str, str]], mapping: dict[str, str], *, unique_invoices: bool
) -> tuple[int, dict[str, Decimal | None]]:
    invoice_keys: set[str] = set()
    totals: dict[str, Decimal | None] = {
        "amount_excl_tax": None,
        "tax_amount": None,
        "total_amount": None,
    }
    for row in rows:
        invoice_code = row.get(mapping.get("invoice_code", ""), "")
        invoice_number = row.get(mapping.get("invoice_number", ""), "")
        invoice_key = f"{invoice_code}|{invoice_number}".strip("|")
        if not invoice_key:
            continue
        already_seen = bool(invoice_key and invoice_key in invoice_keys)
        if invoice_key:
            invoice_keys.add(invoice_key)
        if unique_invoices and already_seen:
            continue
        for field in totals:
            column = mapping.get(field)
            value = _decimal_value(row.get(column, "")) if column else None
            if value is None:
                continue
            totals[field] = (totals[field] or Decimal("0")) + value
    return len(invoice_keys), totals


def _validate_summary(
    base_rows: list[dict[str, str]],
    base_mapping: dict[str, str],
    summary_rows: list[dict[str, str]],
    summary_mapping: dict[str, str],
    summary_sheet_name: str,
) -> dict[str, Any]:
    if not summary_sheet_name:
        return {
            "status": "not_available",
            "reason": "未找到信息汇总表",
        }
    if not summary_rows:
        return {
            "status": "not_available",
            "sheetName": summary_sheet_name,
            "reason": "信息汇总表没有可校验数据行",
        }

    base_count, base_totals = _sheet_totals(base_rows, base_mapping, unique_invoices=True)
    summary_count, summary_totals = _sheet_totals(summary_rows, summary_mapping, unique_invoices=False)
    fields = ("amount_excl_tax", "tax_amount", "total_amount")
    labels = {
        "amount_excl_tax": "金额",
        "tax_amount": "税额",
        "total_amount": "价税合计",
    }
    mismatches: list[str] = []
    if base_count != summary_count:
        mismatches.append(f"发票张数：基础信息={base_count}，信息汇总表={summary_count}")
    for field in fields:
        base_value = base_totals[field]
        summary_value = summary_totals[field]
        if base_value is None or summary_value is None:
            continue
        if abs(base_value - summary_value) > Decimal("0.02"):
            mismatches.append(
                f"{labels[field]}：基础信息={base_value.quantize(Decimal('0.01'))}，"
                f"信息汇总表={summary_value.quantize(Decimal('0.01'))}"
            )
    return {
        "status": "mismatch" if mismatches else "matched",
        "sheetName": summary_sheet_name,
        "baseInvoiceCount": base_count,
        "summaryInvoiceCount": summary_count,
        "baseTotals": {
            field: str(value.quantize(Decimal("0.01"))) if value is not None else None
            for field, value in base_totals.items()
        },
        "summaryTotals": {
            field: str(value.quantize(Decimal("0.01"))) if value is not None else None
            for field, value in summary_totals.items()
        },
        "mismatches": mismatches,
    }


def _parse_xlsx(content: bytes, *, max_rows: int) -> tuple[
    str,
    list[str],
    list[dict[str, str]],
    int,
    str,
    list[str],
    dict[str, str],
    dict[str, Any],
    str,
    list[str],
    dict[str, str],
    list[dict[str, str]],
]:
    from openpyxl import load_workbook

    _validate_xlsx(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取 XLSX 文件") from exc
    def parse_worksheet(worksheet) -> tuple[list[str], list[dict[str, str]]]:
        sample = [list(row) for row in worksheet.iter_rows(max_row=20, values_only=True)]
        headers, rows = _rows_to_records(
            worksheet.iter_rows(values_only=True),
            header_index=_find_header_row(sample),
            max_rows=max_rows,
        )
        return headers, rows

    def role(title: str) -> str:
        normalized = normalize_header(title)
        if "发票基础信息" in normalized:
            return "primary"
        if "信息汇总" in normalized:
            return "summary"
        if "发票明细" in normalized or "明细信息" in normalized:
            return "detail"
        return "other"

    try:
        primary = [worksheet for worksheet in workbook.worksheets if role(worksheet.title) == "primary"]
        if len(primary) > 1:
            raise ValueError("税局文件包含多个「发票基础信息」Sheet，无法确定主表")
        if primary:
            primary_sheet = primary[0]
        else:
            raise ValueError("税局 XLSX 必须包含「发票基础信息」Sheet，不能使用「信息汇总表」作为发票主数据")

        headers, rows = parse_worksheet(primary_sheet)
        primary_mapping = resolve_mapping(headers)
        filtered_rows: list[dict[str, str]] = []
        ignored_row_count = 0
        total_markers = {"合计行", "总计", "合计"}
        for row in rows:
            invoice_code = row.get(primary_mapping.get("invoice_code", ""), "")
            invoice_number = row.get(primary_mapping.get("invoice_number", ""), "")
            row_text = " ".join(row.values())
            if any(token in row_text for token in total_markers) and (
                (not invoice_code and not invoice_number)
                or invoice_code in total_markers
                or invoice_number in total_markers
            ):
                ignored_row_count += 1
                continue
            filtered_rows.append(row)
        rows = filtered_rows

        summary_sheets = [worksheet for worksheet in workbook.worksheets if role(worksheet.title) == "summary"]
        summary_headers: list[str] = []
        summary_rows: list[dict[str, str]] = []
        summary_sheet_names: list[str] = []
        for worksheet in summary_sheets:
            try:
                current_headers, current_rows = parse_worksheet(worksheet)
            except ValueError:
                continue
            if not summary_headers:
                summary_headers = current_headers
            summary_rows.extend(current_rows)
            summary_sheet_names.append(worksheet.title)

        detail_sheets = [worksheet for worksheet in workbook.worksheets if role(worksheet.title) == "detail"]
        detail_headers: list[str] = []
        detail_rows: list[dict[str, str]] = []
        detail_sheet_names: list[str] = []
        for worksheet in detail_sheets:
            try:
                current_headers, current_rows = parse_worksheet(worksheet)
            except ValueError:
                continue
            if not detail_headers:
                detail_headers = current_headers
            detail_rows.extend(current_rows)
            detail_sheet_names.append(worksheet.title)
    finally:
        workbook.close()
    mapping = resolve_mapping(headers)
    summary_mapping = resolve_mapping(summary_headers)
    summary_validation = _validate_summary(
        rows,
        mapping,
        summary_rows,
        summary_mapping,
        "、".join(summary_sheet_names),
    )
    return (
        primary_sheet.title,
        headers,
        rows,
        ignored_row_count,
        "、".join(summary_sheet_names),
        summary_headers,
        summary_mapping,
        summary_validation,
        "、".join(detail_sheet_names),
        detail_headers,
        resolve_mapping(detail_headers),
        detail_rows,
    )


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "utf-8"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("CSV 编码无法识别，请导出 UTF-8 或 GBK/GB18030 文件")


def _parse_csv(content: bytes, *, max_rows: int) -> tuple[str, list[str], list[dict[str, str]]]:
    text = _decode_csv(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    table = [list(row) for row in csv.reader(io.StringIO(text), dialect=dialect)]
    header_index = _find_header_row(table)
    headers, rows = _rows_to_records(table, header_index=header_index, max_rows=max_rows)
    return "CSV", headers, rows


def parse_tax_invoice_export(content: bytes, original_name: str, *, max_rows: int) -> ParsedTaxInvoiceExport:
    if not content:
        raise ValueError("空文件")
    extension = Path(original_name or "").suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("仅支持税务系统导出的 XLSX 或 CSV 文件")
    if extension == ".xlsx":
        (
            sheet_name,
            headers,
            rows,
            ignored_row_count,
            summary_sheet_name,
            summary_headers,
            summary_mapping,
            summary_validation,
            detail_sheet_name,
            detail_headers,
            detail_mapping,
            detail_rows,
        ) = _parse_xlsx(content, max_rows=max_rows)
    else:
        sheet_name, headers, rows = _parse_csv(content, max_rows=max_rows)
        summary_sheet_name = ""
        summary_headers = []
        summary_mapping = {}
        summary_validation = {}
        ignored_row_count = 0
        detail_sheet_name = ""
        detail_headers = []
        detail_mapping = {}
        detail_rows = []
    mapping = resolve_mapping(headers)
    normalized_headers = " ".join(normalize_header(value) for value in headers)
    # “进销项”是税务系统常见的综合字段名，不能因为包含“销项”就误判为销项。
    if "进销项" in normalized_headers:
        direction_hint = "unknown"
    elif "进项" in normalized_headers:
        direction_hint = "input"
    elif "销项" in normalized_headers:
        direction_hint = "output"
    else:
        direction_hint = "unknown"
    return ParsedTaxInvoiceExport(
        sheet_name,
        headers,
        rows,
        mapping,
        direction_hint,
        ignored_row_count=ignored_row_count,
        summary_sheet_name=summary_sheet_name,
        summary_headers=summary_headers,
        summary_mapping=summary_mapping,
        summary_validation=summary_validation,
        detail_sheet_name=detail_sheet_name,
        detail_headers=detail_headers,
        detail_mapping=detail_mapping,
        detail_rows=detail_rows,
    )
