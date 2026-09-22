"""销售清单导入的真实工作簿结构兼容性测试。"""
import io

from openpyxl import Workbook, load_workbook

from app.services import sales_file_import_service as service


def _xlsx(rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def test_single_sheet_order_rows_can_also_supply_item_rows():
    content = _xlsx([
        ["订单编号", "销售渠道", "订单状态", "货品编号", "货品名称", "数量", "单价", "金额"],
        ["JY-001", "PDD", "已完成", "SKU-001", "咖啡豆", 2, 10, 20],
        [None, None, None, "SKU-002", "滤纸", 1, 5, 5],
    ])
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    headers = service._row_headers(ws)

    assert service._is_order_sheet(headers)
    assert service._is_item_sheet(headers)
    groups = service._collect_item_groups(ws, headers, same_sheet=True)
    assert list(groups) == ["JY-001"]
    assert [line["货品编号"] for line in groups["JY-001"]] == ["SKU-001", "SKU-002"]
    wb.close()


def test_upload_archive_keeps_the_exact_received_bytes(tmp_path, monkeypatch):
    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    content = b"received-from-lan-browser"
    target = service._archive_upload(content, "销售单查询202609161358.xlsx", "abc123")

    assert target.relative_to(tmp_path).parts[0] == "sales-imports"
    assert target.read_bytes() == content


def test_invalid_sales_workbook_returns_an_actionable_format_error(client, tmp_path, monkeypatch):
    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    response = client.post(
        "/api/v1/sales-file/import",
        files={"file": ("坏文件.xlsx", b"not-an-xlsx", "application/octet-stream")},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "文件格式无法解析，请上传有效的 .xlsx 文件"


# ---------- 货品档案校验：明细编码必须已建档 ----------

_ITEM_HEADER = ["订单编号", "销售渠道", "订单状态", "货品编号", "货品名称", "数量", "单价", "金额"]


def _make_sku(db_session, jackyun_sku_id: str, sku_code: str):
    from app.models.catalog import ProductSku

    sku = ProductSku(
        jackyun_sku_id=jackyun_sku_id,
        sku_code=sku_code,
        sku_name=f"货品 {sku_code}",
        product_type="single",
    )
    db_session.add(sku)
    db_session.flush()
    return sku


def _cleanup_sales_fixtures(db_session) -> None:
    """成功导入路径内部会 commit（绕过事务回滚），测试痕迹必须显式清理。"""
    from app.models.catalog import ProductSku
    from app.models.sales import SalesOrder, SalesOrderItem

    order_ids = [
        row[0]
        for row in db_session.query(SalesOrder.id).filter(SalesOrder.order_no.like("JYTEST-%")).all()
    ]
    if order_ids:
        db_session.query(SalesOrderItem).filter(SalesOrderItem.order_id.in_(order_ids)).delete(
            synchronize_session=False,
        )
    db_session.query(SalesOrder).filter(SalesOrder.order_no.like("JYTEST-%")).delete(
        synchronize_session=False,
    )
    db_session.query(ProductSku).filter(ProductSku.jackyun_sku_id.like("JYTEST-SKU-%")).delete(
        synchronize_session=False,
    )
    db_session.commit()


def test_import_rejects_unknown_sku_codes_without_writing_anything(client, db_session, tmp_path, monkeypatch):
    """文件含未建档编码 → 整单拒绝（400），订单与明细一个都不写库。"""
    from app.models.sales import SalesOrder, SalesOrderItem

    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    _make_sku(db_session, "JYTEST-SKU-KNOWN", "JYTEST-SKU-KNOWN")

    content = _xlsx([
        _ITEM_HEADER,
        ["JYTEST-101", "PDD", "已完成", "JYTEST-SKU-KNOWN", "已知货品", 2, 10, 20],
        ["JYTEST-102", "PDD", "已完成", "JYTEST-SKU-UNKNOWN", "未知货品", 1, 5, 5],
    ])
    response = client.post(
        "/api/v1/sales-file/import",
        files={"file": ("销售单查询.xlsx", content, "application/octet-stream")},
    )

    assert response.status_code == 400, response.text
    assert response.json()["detail"] == (
        "货品档案缺少以下货品编码，请先建档或修正后再导入：JYTEST-SKU-UNKNOWN"
    )
    assert db_session.query(SalesOrder).filter(SalesOrder.order_no.like("JYTEST-%")).count() == 0
    assert (
        db_session.query(SalesOrderItem)
        .join(SalesOrder, SalesOrderItem.order_id == SalesOrder.id)
        .filter(SalesOrder.order_no.like("JYTEST-%"))
        .count()
        == 0
    )


def test_import_accepts_sku_code_with_case_mismatch(client, db_session, tmp_path, monkeypatch):
    """编码大小写与档案不一致但存在 → 忽略大小写比对，放行写入。"""
    from app.models.sales import SalesOrder, SalesOrderItem

    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    _make_sku(db_session, "JYTEST-SKU-CASE", "jytest-sku-case")

    content = _xlsx([
        _ITEM_HEADER,
        ["JYTEST-201", "PDD", "已完成", "JYTEST-SKU-CASE", "大小写不一致货品", 3, 7, 21],
    ])
    response = client.post(
        "/api/v1/sales-file/import",
        files={"file": ("销售单查询.xlsx", content, "application/octet-stream")},
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["ordersImported"] == 1
    assert body["itemsImported"] == 1
    order = db_session.query(SalesOrder).filter_by(order_no="JYTEST-201").one()
    item = db_session.query(SalesOrderItem).filter_by(order_id=order.id).one()
    assert item.sku_code == "JYTEST-SKU-CASE"
    assert item.sku_id is not None
    _cleanup_sales_fixtures(db_session)


def test_import_accepts_all_known_sku_codes(client, db_session, tmp_path, monkeypatch):
    """全部编码已建档 → 正常导入订单与明细。"""
    from app.models.sales import SalesOrder, SalesOrderItem

    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    _make_sku(db_session, "JYTEST-SKU-A1", "JYTEST-SKU-A1")
    _make_sku(db_session, "JYTEST-SKU-A2", "JYTEST-SKU-A2")

    content = _xlsx([
        _ITEM_HEADER,
        ["JYTEST-301", "PDD", "已完成", "JYTEST-SKU-A1", "已知货品一", 2, 10, 20],
        ["JYTEST-301", "PDD", "已完成", "JYTEST-SKU-A2", "已知货品二", 1, 5, 5],
    ])
    response = client.post(
        "/api/v1/sales-file/import",
        files={"file": ("销售单查询.xlsx", content, "application/octet-stream")},
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["ordersImported"] == 1
    assert body["itemsImported"] == 2
    order = db_session.query(SalesOrder).filter_by(order_no="JYTEST-301").one()
    assert db_session.query(SalesOrderItem).filter_by(order_id=order.id).count() == 2
    _cleanup_sales_fixtures(db_session)


# ---------- 明细成本口径：文件只提供销售数量与客户实付，成本按本系统入库加权 ----------


def _cleanup_inbound_fixture(db_session) -> None:
    from app.models.jackyun import JackyunGoodsDocument, JackyunGoodsDocumentItem

    doc = db_session.query(JackyunGoodsDocument).filter_by(goodsdoc_no="JYTEST-INBOUND-COST").one_or_none()
    if doc is not None:
        db_session.query(JackyunGoodsDocumentItem).filter_by(document_id=doc.id).delete(
            synchronize_session=False,
        )
        db_session.delete(doc)
    db_session.commit()


def test_detail_cost_uses_inbound_weighted_cost_and_paid_amount(client, db_session, tmp_path, monkeypatch):
    """明细页订单成本取本系统入库加权平均成本，毛利 = 订单实付 − 订单成本。

    导入文件自带的「订单货品成本 / 毛利」列不得参与，缺入库成本时两个字段留空。
    """
    from datetime import datetime, timezone
    from decimal import Decimal

    from app.models.jackyun import JackyunGoodsDocument, JackyunGoodsDocumentItem

    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    costed = _make_sku(db_session, "JYTEST-SKU-COST", "JYTEST-SKU-COST")
    _make_sku(db_session, "JYTEST-SKU-NOCOST", "JYTEST-SKU-NOCOST")  # 无入库成本 SKU（匹配 Excel 行）
    inbound = JackyunGoodsDocument(
        document_type="inbound", goodsdoc_no="JYTEST-INBOUND-COST",
        document_at=datetime(2026, 1, 5, tzinfo=timezone.utc),
    )
    db_session.add(inbound)
    db_session.flush()
    db_session.add(JackyunGoodsDocumentItem(
        document_id=inbound.id, line_no=1, goods_no=costed.sku_code,
        quantity=Decimal("10"), unit_price_tax=Decimal("5.00"), matched_sku_id=costed.id,
    ))
    db_session.flush()

    content = _xlsx([
        _ITEM_HEADER + ["付款时间", "实付金额", "订单货品成本", "毛利"],
        ["JYTEST-COST-1", "PDD", "已完成", "JYTEST-SKU-COST", "有入库成本货品", 2, 10, 20, "2026-09-10 10:00", 26, 9999, 9999],
        ["JYTEST-COST-2", "PDD", "已完成", "JYTEST-SKU-NOCOST", "无入库成本货品", 1, 8, 8, "2026-09-11 10:00", 8, 0, 8],
    ])
    response = client.post(
        "/api/v1/sales-file/import",
        files={"file": ("销售单查询.xlsx", content, "application/octet-stream")},
    )
    assert response.status_code == 200, response.text

    detail = client.get("/api/v1/sales-file/detail", params={"q": "JYTEST-COST-"})
    assert detail.status_code == 200, detail.text
    rows = {row["orderNo"]: row for row in detail.json()["rows"]}

    assert rows["JYTEST-COST-1"]["goodsCost"] == 10.0  # 2 × 5.00，而不是文件里的 9999
    assert rows["JYTEST-COST-1"]["grossProfit"] == 16.0  # 实付 26 − 成本 10
    assert rows["JYTEST-COST-1"]["costIncomplete"] is False
    assert rows["JYTEST-COST-2"]["goodsCost"] is None  # 无入库成本不留 0
    assert rows["JYTEST-COST-2"]["grossProfit"] is None
    assert rows["JYTEST-COST-2"]["costIncomplete"] is True

    _cleanup_inbound_fixture(db_session)
    _cleanup_sales_fixtures(db_session)


def test_detail_export_uses_same_scope_and_cost_as_page(client, db_session, tmp_path, monkeypatch):
    """销售明细页「导出 Excel」：与页面同一筛选与成本口径，且不受分页限制。"""
    from datetime import datetime, timezone
    from decimal import Decimal
    from io import BytesIO

    from openpyxl import load_workbook

    from app.models.jackyun import JackyunGoodsDocument, JackyunGoodsDocumentItem

    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    costed = _make_sku(db_session, "JYTEST-SKU-COST", "JYTEST-SKU-COST")
    inbound = JackyunGoodsDocument(
        document_type="inbound", goodsdoc_no="JYTEST-INBOUND-COST",
        document_at=datetime(2026, 1, 5, tzinfo=timezone.utc),
    )
    db_session.add(inbound)
    db_session.flush()
    db_session.add(JackyunGoodsDocumentItem(
        document_id=inbound.id, line_no=1, goods_no=costed.sku_code,
        quantity=Decimal("10"), unit_price_tax=Decimal("5.00"), matched_sku_id=costed.id,
    ))
    db_session.flush()

    content = _xlsx([
        _ITEM_HEADER + ["付款时间", "实付金额", "订单货品成本", "毛利"],
        ["JYTEST-EXP-1", "PDD", "已完成", "JYTEST-SKU-COST", "有入库成本货品", 2, 10, 20, "2026-09-10 10:00", 26, 9999, 9999],
        ["JYTEST-EXP-2", "TM", "已完成", "JYTEST-SKU-COST", "有入库成本货品", 1, 10, 10, "2026-09-11 10:00", 13, 9999, 9999],
    ])
    response = client.post(
        "/api/v1/sales-file/import",
        files={"file": ("销售单查询.xlsx", content, "application/octet-stream")},
    )
    assert response.status_code == 200, response.text

    # 页面分页只给 1 行，导出必须给全部 2 行
    first = client.get("/api/v1/sales-file/detail", params={"q": "JYTEST-EXP-", "page_size": 1})
    assert first.json()["total"] == 2
    assert len(first.json()["rows"]) == 1
    page_rows = {
        row["orderNo"]: row
        for row in client.get(
            "/api/v1/sales-file/detail", params={"q": "JYTEST-EXP-", "page_size": 50}
        ).json()["rows"]
    }

    exported = client.get("/api/v1/sales-file/detail/export", params={"q": "JYTEST-EXP-"})
    assert exported.status_code == 200, exported.text
    assert exported.headers["x-export-row-count"] == "2"
    assert exported.headers["content-type"].startswith("application/vnd.openxmlformats")

    sheet = load_workbook(BytesIO(exported.content))["销售明细"]
    headers = [cell.value for cell in sheet[1]]
    rows = {
        row[0]: dict(zip(headers, row))
        for row in sheet.iter_rows(min_row=2, values_only=True)
    }
    assert set(rows) == {"JYTEST-EXP-1", "JYTEST-EXP-2"}
    assert rows["JYTEST-EXP-1"]["订单成本"] == page_rows["JYTEST-EXP-1"]["goodsCost"]  # 与页面同一成本函数
    assert rows["JYTEST-EXP-2"]["订单成本"] == page_rows["JYTEST-EXP-2"]["goodsCost"]
    assert rows["JYTEST-EXP-1"]["订单实付"] == 26
    assert rows["JYTEST-EXP-1"]["订单毛利"] == 16  # 26 − 2×5.00
    assert rows["JYTEST-EXP-2"]["订单毛利"] == 8  # 13 − 1×5.00
    assert not rows["JYTEST-EXP-1"]["成本提示"]  # 有成本的行不标注

    # 「订单汇总」表：按订单去重 + 毛利/合计写成 Excel 公式，导出后不用再手工核算
    assert exported.headers["x-export-order-count"] == "2"
    summary = load_workbook(BytesIO(exported.content))["订单汇总"]
    summary_rows = [dict(zip([c.value for c in summary[1]], row)) for row in summary.iter_rows(min_row=2, values_only=True)]
    assert len(summary_rows) == 3  # 2 单 + 合计行
    assert set(r["订单号"] for r in summary_rows[:2]) == {"JYTEST-EXP-1", "JYTEST-EXP-2"}
    assert summary_rows[0]["订单毛利"] == '=IF(L2="","",K2-L2)'
    assert summary_rows[1]["订单毛利"] == '=IF(L3="","",K3-L3)'
    assert summary_rows[2]["订单号"] == "合计"
    assert summary_rows[2]["订单实付"] == "=SUM(K2:K3)"
    assert summary_rows[2]["订单成本"] == "=SUM(L2:L3)"
    assert summary_rows[2]["订单毛利"] == "=SUM(M2:M3)"

    _cleanup_inbound_fixture(db_session)
    _cleanup_sales_fixtures(db_session)


def test_detail_export_splits_order_amount_to_sku_lines_by_cost_ratio(client, db_session, tmp_path, monkeypatch):
    """子 SKU 拆解：行成本 = 数量 × 入库加权单价，订单实付按行成本占比分摊，缺成本退回行金额占比。"""
    from datetime import datetime, timezone
    from decimal import Decimal
    from io import BytesIO

    from openpyxl import load_workbook

    from app.models.jackyun import JackyunGoodsDocument, JackyunGoodsDocumentItem

    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    cheap = _make_sku(db_session, "JYTEST-SKU-A", "JYTEST-SKU-A")
    pricey = _make_sku(db_session, "JYTEST-SKU-B", "JYTEST-SKU-B")
    nocost = _make_sku(db_session, "JYTEST-SKU-C", "JYTEST-SKU-C")
    inbound = JackyunGoodsDocument(
        document_type="inbound", goodsdoc_no="JYTEST-INBOUND-SPLIT",
        document_at=datetime(2026, 1, 5, tzinfo=timezone.utc),
    )
    db_session.add(inbound)
    db_session.flush()
    db_session.add_all([
        JackyunGoodsDocumentItem(
            document_id=inbound.id, line_no=1, goods_no=cheap.sku_code,
            quantity=Decimal("10"), unit_price_tax=Decimal("5.00"), matched_sku_id=cheap.id,
        ),
        JackyunGoodsDocumentItem(
            document_id=inbound.id, line_no=2, goods_no=pricey.sku_code,
            quantity=Decimal("10"), unit_price_tax=Decimal("20.00"), matched_sku_id=pricey.id,
        ),
    ])
    db_session.flush()

    content = _xlsx([
        _ITEM_HEADER + ["付款时间", "实付金额", "订单货品成本", "毛利"],
        ["JYTEST-SPLIT-1", "PDD", "已完成", cheap.sku_code, "便宜货品", 2, 10, 20, "2026-09-12 10:00", 26, 9999, 9999],
        [None, None, None, pricey.sku_code, "贵货品", 1, 30, 30, "2026-09-12 10:00", 26, 9999, 9999],
        ["JYTEST-SPLIT-2", "PDD", "已完成", cheap.sku_code, "便宜货品", 1, 10, 10, "2026-09-13 10:00", 12, 9999, 9999],
        [None, None, None, nocost.sku_code, "无入库成本货品", 1, 30, 30, "2026-09-13 10:00", 12, 9999, 9999],
    ])
    response = client.post(
        "/api/v1/sales-file/import",
        files={"file": ("销售单查询.xlsx", content, "application/octet-stream")},
    )
    assert response.status_code == 200, response.text

    exported = client.get("/api/v1/sales-file/detail/export", params={"q": "JYTEST-SPLIT-"})
    assert exported.status_code == 200, exported.text
    sheet = load_workbook(BytesIO(exported.content))["销售明细"]
    headers = [c.value for c in sheet[1]]
    rows = [dict(zip(headers, r)) for r in sheet.iter_rows(min_row=2, values_only=True)]
    by_key = {(r["订单号"], r["货品编号"]): r for r in rows}
    assert len(rows) == 4

    # 行成本 = 数量 × 入库加权单价：2×5.00 = 10，1×20.00 = 20
    split1_cheap = by_key[("JYTEST-SPLIT-1", cheap.sku_code)]
    split1_pricey = by_key[("JYTEST-SPLIT-1", pricey.sku_code)]
    assert split1_cheap["子SKU成本"] == 10
    assert split1_pricey["子SKU成本"] == 20
    # 订单实付 26 按成本占比 1:2 拆：8.67 + 17.33 = 26.00（尾差补到最后一行）
    assert split1_cheap["子SKU收入（按成本比例分摊）"] == 8.67
    assert split1_pricey["子SKU收入（按成本比例分摊）"] == 17.33
    assert split1_cheap["分摊依据"] == "按行成本比例"
    assert split1_cheap["子SKU毛利"] == -1.33  # 8.67 − 10

    # 整单有行缺成本 → 退回按行金额比例，且缺成本行不编成本
    assert by_key[("JYTEST-SPLIT-2", cheap.sku_code)]["子SKU收入（按成本比例分摊）"] == 3.0  # 12 × 10/40
    assert by_key[("JYTEST-SPLIT-2", nocost.sku_code)]["子SKU收入（按成本比例分摊）"] == 9.0  # 12 × 30/40
    assert by_key[("JYTEST-SPLIT-2", nocost.sku_code)]["子SKU成本"] is None
    assert by_key[("JYTEST-SPLIT-2", nocost.sku_code)]["子SKU毛利"] is None
    assert by_key[("JYTEST-SPLIT-2", nocost.sku_code)]["分摊依据"] == "按行金额比例"

    _cleanup_inbound_fixture(db_session)
    _cleanup_sales_fixtures(db_session)


# ---------- 成交口径：待发/已发/待确认收货/已完成计入，关闭/退货/取消/待审核不计入 ----------


def test_import_keeps_only_deal_statuses_and_scope_condition_matches(client, db_session, tmp_path, monkeypatch):
    """统一成交口径：所有订单保留事实；只有正常流转状态计入业绩。"""
    from app.models.sales import SalesOrder
    from app.services.sales_scope import NOT_DEAL_PREFIXES, deal_orders_condition, is_deal_status

    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    _make_sku(db_session, "JYTEST-SKU-S1", "JYTEST-SKU-S1")
    stamp = "2026-09-10 10:00"
    rows = [
        ["JYTEST-SCOPE-1", "待发货-已递交"],
        ["JYTEST-SCOPE-2", "发货在途"],
        ["JYTEST-SCOPE-3", "待确认收货"],
        ["JYTEST-SCOPE-4", "已完成"],
        ["JYTEST-SCOPE-5", "已取消-被合并"],
        ["JYTEST-SCOPE-6", "已关闭"],
        ["JYTEST-SCOPE-7", "退货成功"],
        ["JYTEST-SCOPE-8", "退款关闭"],
        ["JYTEST-SCOPE-9", "待审核"],
    ]
    content = _xlsx([
        _ITEM_HEADER + ["付款时间"],
        *[[no, "PDD", status, "JYTEST-SKU-S1", "口径货品", 1, 10, 10, stamp] for no, status in rows],
    ])
    response = client.post(
        "/api/v1/sales-file/import",
        files={"file": ("销售单查询.xlsx", content, "application/octet-stream")},
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["ordersImported"] == 4
    assert body["cancelledSkipped"] == 5

    scope = SalesOrder.order_no.like("JYTEST-SCOPE-%")
    kept = {row[0] for row in db_session.query(SalesOrder.order_no).filter(scope).all()}
    assert kept == {no for no, _status in rows}
    assert body["removedCancelled"] == 0
    assert body["cancelledPreserved"] == 5

    assert is_deal_status("待发货-已递交") and is_deal_status("发货在途")
    assert is_deal_status("待确认收货") and is_deal_status("已完成")
    assert not any(is_deal_status(prefix) for prefix in NOT_DEAL_PREFIXES)
    counted = {row[0] for row in db_session.query(SalesOrder.order_no).filter(scope, deal_orders_condition()).all()}
    assert counted == {"JYTEST-SCOPE-1", "JYTEST-SCOPE-2", "JYTEST-SCOPE-3", "JYTEST-SCOPE-4"}

    _cleanup_sales_fixtures(db_session)

def test_number_parser_is_decimal_and_datetime_is_shanghai_aware():
    from decimal import Decimal
    from zoneinfo import ZoneInfo

    assert service._num("0.1") == Decimal("0.1000")
    parsed = service._dt("2026-09-30 23:30:00")
    assert parsed is not None
    assert parsed.tzinfo is not None
    assert parsed.utcoffset() == ZoneInfo(service.settings.TZ).utcoffset(parsed)


def test_cancelled_file_fact_preserves_synced_order_and_invoice_link(client, db_session, tmp_path, monkeypatch):
    from decimal import Decimal
    from app.models.sales import SalesOrder
    from app.models.tax import TaxInvoice, TaxInvoiceLink
    from app.services import tax_invoice_service

    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    _make_sku(db_session, "JYTEST-SKU-CANCEL", "JYTEST-SKU-CANCEL")
    order = SalesOrder(
        order_no="JYTEST-CANCEL-KEEP",
        source_provider="jky_web",
        source_order_id="remote-1",
        order_status="已完成",
        paid_amount=Decimal("100"),
        raw={"_jky": {"sourceProvider": "jky_web"}},
    )
    db_session.add(order)
    db_session.flush()
    invoice = TaxInvoice(
        invoice_key="JYTEST-CANCEL-INVOICE",
        invoice_number="JYTEST-CANCEL-INVOICE",
        direction="output",
        status="issued",
        total_amount=Decimal("100"),
    )
    db_session.add(invoice)
    db_session.flush()
    link = TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="sales_order",
        target_id=order.id,
        allocated_amount=Decimal("100"),
        match_method="manual",
        confirmed=True,
    )
    db_session.add(link)
    db_session.commit()

    content = _xlsx([
        _ITEM_HEADER + ["付款时间"],
        ["JYTEST-CANCEL-KEEP", "PDD", "已取消", "JYTEST-SKU-CANCEL", "取消货品", 1, 100, 100, "2026-09-30 23:30"],
    ])
    response = client.post(
        "/api/v1/sales-file/import",
        files={"file": ("销售单查询.xlsx", content, "application/octet-stream")},
    )
    assert response.status_code == 200, response.text

    db_session.expire_all()
    kept = db_session.query(SalesOrder).filter_by(order_no="JYTEST-CANCEL-KEEP").one()
    assert kept.source_provider == "jky_web"
    assert kept.order_status == "已取消"
    assert db_session.query(TaxInvoiceLink).filter_by(invoice_id=invoice.id).count() == 1

    serialized = tax_invoice_service.serialize_invoice(
        db_session.get(TaxInvoice, invoice.id), db=db_session
    )
    assert serialized["businessMatchStatus"] == "needs_review"
    assert serialized["businessMatchedAmount"] == "0.0000"
    assert serialized["invalidLinkCount"] == 1

    db_session.query(TaxInvoiceLink).filter_by(invoice_id=invoice.id).delete(synchronize_session=False)
    db_session.query(TaxInvoice).filter_by(id=invoice.id).delete(synchronize_session=False)
    db_session.commit()
    _cleanup_sales_fixtures(db_session)
