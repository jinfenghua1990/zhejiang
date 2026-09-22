from io import BytesIO

from openpyxl import load_workbook


def test_data_export_options_are_stable(client):
    response = client.get("/api/v1/data/export-options")

    assert response.status_code == 200
    keys = [row["key"] for row in response.json()]
    assert keys[:4] == ["purchase_orders", "catalog", "inventory", "warehouses"]
    assert "inbound_documents" in keys
    assert "bundles" in keys
    assert "tax_rules" in keys
    assert "suppliers" in keys
    assert "external_orders" in keys


def test_catalog_export_returns_readable_xlsx(client):
    response = client.get("/api/v1/data/export/catalog", params={"kind": "goods"})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "filename*=" in response.headers["content-disposition"]
    assert response.headers["x-export-row-count"] == "0"

    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook.sheetnames == ["货品档案"]
    assert [cell.value for cell in workbook["货品档案"][1]][:4] == ["类型", "编码", "名称", "商品名称"]


def test_purchase_export_has_three_reconciliation_sheets(client):
    response = client.get("/api/v1/data/export/purchase_orders")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook.sheetnames == ["采购订单", "采购明细", "入库关联"]


def test_master_exports_include_roundtrip_keys(client):
    catalog = client.get("/api/v1/data/export/catalog", params={"template": "true"})
    warehouse = client.get("/api/v1/data/export/warehouses", params={"template": "true"})
    supplier = client.get("/api/v1/data/export/suppliers", params={"template": "true"})
    bundles = client.get("/api/v1/data/export/bundles", params={"template": "true"})
    tax_rules = client.get("/api/v1/data/export/tax_rules", params={"template": "true"})
    external = client.get("/api/v1/data/export/external_orders", params={"template": "true"})
    assert all(response.status_code == 200 for response in (catalog, warehouse, bundles, tax_rules, supplier, external))
    assert "档案ID" in [cell.value for cell in load_workbook(BytesIO(catalog.content), data_only=True)["货品档案"][1]]
    assert "档案ID" in [cell.value for cell in load_workbook(BytesIO(warehouse.content), data_only=True)["仓库档案"][1]]
    bundles_workbook = load_workbook(BytesIO(bundles.content), data_only=True)
    assert bundles_workbook.sheetnames == ["套装档案"]
    assert "货品类型" in [cell.value for cell in bundles_workbook["套装档案"][1]]
    assert "规则ID" in [cell.value for cell in load_workbook(BytesIO(tax_rules.content), data_only=True)["财务分类"][1]]
    assert "供应商ID" in [cell.value for cell in load_workbook(BytesIO(supplier.content), data_only=True)["供应商档案"][1]]
    assert "订单ID" in [cell.value for cell in load_workbook(BytesIO(external.content), data_only=True)["其他渠道订单"][1]]
