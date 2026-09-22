"""月度已收票对公付款清单。"""
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.models.bank import BankAccount, BankTransaction
from app.models.purchase import ExternalPurchaseOrder, Supplier
from app.models.tax import TaxInvoice, TaxInvoiceLink
from app.services import finance_corporate_payment_report_service as service


def test_build_report_links_bank_invoice_purchase(db_session):
    token = uuid4().hex[:10]
    seller = f"测试供应商-{token}"
    company = f"测试公司-{token}"

    account = BankAccount(
        account_no=f"ACCT-{token}",
        account_name=company,
        bank_name="测试银行",
        currency="CNY",
    )
    db_session.add(account)
    db_session.flush()

    supplier = Supplier(
        platform="other",
        external_shop_id=f"SUP-{token}",
        name=seller,
        tax_no=f"TAX-{token}",
    )
    db_session.add(supplier)

    po = ExternalPurchaseOrder(
        external_order_id=f"PO-{token}",
        platform="1688",
        supplier_name=seller,
        ordered_at=datetime(2026, 8, 2, tzinfo=timezone.utc),
        order_amount=Decimal("1000.00"),
        paid_amount=Decimal("1000.00"),
        currency="CNY",
        order_status="paid",
        pay_status="paid",
        raw={},
    )
    db_session.add(po)
    db_session.flush()


    txn = BankTransaction(
        account_id=account.id,
        txn_date=date(2026, 8, 15),
        direction="out",
        amount=Decimal("1000.00"),
        counterparty_name=seller,
        counterparty_account=f"CP-{token}",
        summary="采购货款",
        voucher_no=f"V-{token}",
        fingerprint=f"pytest-corp-pay-{token}",
        raw={},
    )
    db_session.add(txn)

    invoice = TaxInvoice(
        invoice_key=f"pytest-corp-pay-{token}",
        invoice_number=f"INV-{token}",
        direction="input",
        status="issued",
        issue_date=datetime(2026, 8, 10, tzinfo=timezone.utc),
        seller_name=seller,
        seller_tax_id=f"TAX-{token}",
        buyer_name=company,
        amount_excl_tax=Decimal("884.96"),
        tax_amount=Decimal("115.04"),
        total_amount=Decimal("1000.00"),
        source_system="tax_export",
        raw={},
    )
    db_session.add(invoice)
    db_session.flush()

    db_session.add(TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="bank_transaction",
        target_id=txn.id,
        allocated_amount=Decimal("1000.00"),
        match_method="manual",
        confirmed=True,
        note="pytest",
    ))
    db_session.commit()

    report = service.build_report(db_session, 2026, 8, company=company)

    assert report["summary"]["paymentCount"] == 1
    assert report["summary"]["invoiceCount"] == 1
    assert Decimal(report["summary"]["allocatedTotal"]) == Decimal("1000.00")
    assert report["invoiceRows"][0]["invoiceKey"] == invoice.invoice_key
    assert report["invoiceRows"][0]["invoiceNumber"] == invoice.invoice_number
    assert report["invoiceRows"][0]["payments"][0]["voucherNo"] == txn.voucher_no
    assert report["invoiceRows"][0]["paymentSource"] == "corporate"
    assert report["invoiceRows"][0]["paymentSourceLabel"] == "对公支付"
    assert report["invoiceRows"][0]["paymentMethod"] == "corporate"
    assert report["invoiceRows"][0]["paymentMethodLabel"] == "对公账户支出"
    assert report["invoiceRows"][0]["expenseNatureLabel"] == "货款"
    assert po.external_order_id in report["invoiceRows"][0]["purchaseOrderNos"]

    blob = service.corporate_payment_xlsx(report)
    wb = load_workbook(BytesIO(blob), read_only=True)
    assert wb.sheetnames == ["月度汇总", "已收票对公核对"]
    invoice_ws = wb["已收票对公核对"]
    headers = {cell.value: cell.column for cell in invoice_ws[1]}
    assert invoice_ws["D2"].value == invoice.invoice_number
    assert invoice_ws.cell(2, headers["发票属性"]).value == "蓝字发票"


def test_report_keeps_received_invoice_even_when_bank_link_is_unconfirmed(db_session):
    token = uuid4().hex[:10]
    seller = f"未确认供应商-{token}"

    txn = BankTransaction(
        txn_date=date(2026, 8, 5),
        direction="out",
        amount=Decimal("300.00"),
        counterparty_name=seller,
        fingerprint=f"pytest-corp-unconfirmed-{token}",
        raw={},
    )
    db_session.add(txn)
    invoice = TaxInvoice(
        invoice_key=f"pytest-corp-unconfirmed-{token}",
        invoice_number=f"INV-U-{token}",
        direction="input",
        status="issued",
        issue_date=datetime(2026, 8, 3, tzinfo=timezone.utc),
        seller_name=seller,
        total_amount=Decimal("300.00"),
        source_system="tax_export",
        raw={},
    )
    db_session.add(invoice)
    db_session.flush()
    db_session.add(TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="bank_transaction",
        target_id=txn.id,
        allocated_amount=Decimal("300.00"),
        match_method="manual",
        confirmed=False,
    ))
    db_session.commit()

    report = service.build_report(db_session, 2026, 8)

    assert [row["invoiceId"] for row in report["invoiceRows"]] == [invoice.id]
    assert report["invoiceRows"][0]["invoiceStatus"] == "unpaid"
    assert report["invoiceRows"][0]["paymentSource"] == "personal"
    assert report["invoiceRows"][0]["paymentSourceLabel"] == "个人支付（系统推定）"
    assert report["invoiceRows"][0]["manualPaymentMethod"] == ""
    assert report["invoiceRows"][0]["paymentMethod"] == "personal"
    assert report["invoiceRows"][0]["paymentMethodLabel"] == "个人垫付"
    assert report["invoiceRows"][0]["payments"] == []
    assert Decimal(report["summary"]["invoiceTotal"]) == Decimal("300.00")
    assert Decimal(report["summary"]["outstandingTotal"]) == Decimal("300.00")


def test_corporate_payment_xlsx_keeps_one_row_per_invoice_with_multiple_payments():
    report = {
        "summary": {
            "paymentCount": 2,
            "invoiceCount": 1,
            "paymentTotal": "1000.00",
            "allocatedTotal": "1000.00",
            "invoiceTotal": "1000.00",
            "outstandingTotal": "0.00",
            "paidInvoiceCount": 1,
            "partialInvoiceCount": 0,
            "unpaidInvoiceCount": 0,
        },
        "invoiceRows": [
            {
                "invoiceId": 1,
                "invoiceDate": "2026-08-10",
                "supplierName": "供应商甲",
                "supplierTaxId": "TAX-A",
                "invoiceNumber": "INV-A",
                "invoiceType": "增值税专用发票",
                "invoiceAmountExclTax": "884.96",
                "invoiceTaxAmount": "115.04",
                "invoiceTotalAmount": "1000.00",
                "invoiceCorporatePaidTotal": "1000.00",
                "invoiceOutstandingAmount": "0.00",
                "invoiceStatus": "paid",
                "purchaseOrderNos": ["PO-A", "PO-B"],
                "payments": [
                    {
                        "linkId": 10,
                        "paymentId": 20,
                        "paymentDate": "2026-08-15",
                        "paymentAccount": "ZJRC-001",
                        "paymentAccountName": "测试账户",
                        "counterpartyAccount": "CP-001",
                        "voucherNo": "V-001",
                        "summary": "货款",
                        "paymentAmount": "400.00",
                        "allocatedAmount": "400.00",
                        "paymentMatchedTotal": "400.00",
                        "paymentStatus": "matched",
                    },
                    {
                        "linkId": 11,
                        "paymentId": 21,
                        "paymentDate": "2026-08-20",
                        "paymentAccount": "ZJRC-001",
                        "paymentAccountName": "测试账户",
                        "counterpartyAccount": "CP-001",
                        "voucherNo": "V-002",
                        "summary": "货款",
                        "paymentAmount": "600.00",
                        "allocatedAmount": "600.00",
                        "paymentMatchedTotal": "600.00",
                        "paymentStatus": "matched",
                    },
                ],
            }
        ],
        "rows": [],
    }

    wb = load_workbook(BytesIO(service.corporate_payment_xlsx(report)), read_only=True)
    ws = wb["已收票对公核对"]

    assert ws.max_row == 2
    headers = {cell.value: cell.column for cell in ws[1]}
    assert ws.cell(2, headers["发票号码"]).value == "INV-A"
    assert ws.cell(2, headers["价税合计"]).value == 1000
    assert ws.cell(2, headers["对公匹配金额"]).value == 1000
    assert ws.cell(2, headers["个人支付/未对公匹配金额"]).value == 0
    assert ws.cell(2, headers["银行付款日期"]).value == "2026-08-15、2026-08-20"
    assert ws.cell(2, headers["银行流水/凭证号"]).value == "V-001、V-002"
    assert ws.cell(2, headers["关联采购订单"]).value == "PO-A、PO-B"


def test_negative_red_invoice_is_not_misclassified_as_bank_reconciled(db_session):
    """红字负数票保留在月度财务清单，但绝不能因为 outstanding 被截成 0 就显示已核对。"""
    token = uuid4().hex[:10]
    invoice = TaxInvoice(
        invoice_key=f"pytest-corp-red-{token}",
        invoice_number=f"INV-RED-{token}",
        direction="input",
        status="red",
        issue_date=datetime(2026, 7, 8, tzinfo=timezone.utc),
        seller_name=f"红冲供应商-{token}",
        amount_excl_tax=Decimal("-1343.36"),
        tax_amount=Decimal("-174.64"),
        total_amount=Decimal("-1518.00"),
        source_system="tax_export",
        raw={},
    )
    db_session.add(invoice)
    db_session.commit()

    report = service.build_report(db_session, 2026, 7)
    row = next(item for item in report["invoiceRows"] if item["invoiceId"] == invoice.id)

    assert row["invoiceTotalAmount"] == "-1518.00"
    assert row["invoiceCorporatePaidTotal"] == "0"
    assert row["invoiceOutstandingAmount"] == "0"
    assert row["invoiceStatus"] == "not_applicable"
    assert row["bankReconciliationStatus"] == "not_applicable"
    assert row["bankReconciliationApplicable"] is False
    assert row["paymentSource"] == "not_applicable"
    assert row["paymentSourceLabel"] == "不适用"
    assert "不参与银行付款核对" in row["bankReconciliationReason"]
    assert row["payments"] == []

    assert report["summary"]["paidInvoiceCount"] == 0
    assert report["summary"]["partialInvoiceCount"] == 0
    assert report["summary"]["unpaidInvoiceCount"] == 0
    assert report["summary"]["notApplicableInvoiceCount"] == 1
    assert Decimal(report["summary"]["outstandingTotal"]) == Decimal("0")

    wb = load_workbook(BytesIO(service.corporate_payment_xlsx(report)), read_only=True)
    ws = wb["已收票对公核对"]
    headers = {cell.value: cell.column for cell in ws[1]}
    assert ws.cell(2, headers["发票属性"]).value == "红字发票"
    assert ws.cell(2, headers["票据状态"]).value == "红字发票（待关联蓝字）"
    assert ws.cell(2, headers["支付方式"]).value == "不适用"
    assert ws.cell(2, headers["银行匹配状态"]).value == "不适用"


def test_small_positive_invoice_without_payment_stays_unpaid(db_session):
    """状态不能仅由“剩余金额落入容差”决定；没有银行付款就绝不能显示已核对。"""
    token = uuid4().hex[:10]
    invoice = TaxInvoice(
        invoice_key=f"pytest-corp-small-{token}",
        invoice_number=f"INV-SMALL-{token}",
        direction="input",
        status="issued",
        issue_date=datetime(2026, 7, 9, tzinfo=timezone.utc),
        seller_name=f"小额供应商-{token}",
        total_amount=Decimal("0.03"),
        source_system="tax_export",
        raw={},
    )
    db_session.add(invoice)
    db_session.commit()

    report = service.build_report(db_session, 2026, 7)
    row = next(item for item in report["invoiceRows"] if item["invoiceId"] == invoice.id)

    assert row["bankReconciliationApplicable"] is True
    assert row["bankReconciliationStatus"] == "unpaid"
    assert row["paymentSource"] == "personal"
    assert row["paymentSourceLabel"] == "个人支付（系统推定）"
    assert row["invoiceCorporatePaidTotal"] == "0"
    assert Decimal(row["invoiceOutstandingAmount"]) == Decimal("0.03")
    assert report["summary"]["paidInvoiceCount"] == 0
    assert report["summary"]["unpaidInvoiceCount"] == 1


def test_expense_nature_recognizes_meal_for_finance_view(db_session):
    """无采购单关联时，按发票明细/原始字段把餐饮归成财务可读的“餐饮费”."""
    token = uuid4().hex[:10]
    invoice = TaxInvoice(
        invoice_key=f"pytest-corp-meal-{token}",
        invoice_number=f"INV-MEAL-{token}",
        direction="input",
        status="issued",
        issue_date=datetime(2026, 7, 12, tzinfo=timezone.utc),
        seller_name=f"餐饮供应商-{token}",
        amount_excl_tax=Decimal("94.34"),
        tax_amount=Decimal("5.66"),
        total_amount=Decimal("100.00"),
        category="reimburse_operating",
        source_system="tax_export",
        raw={"货物或应税劳务、服务名称": "*餐饮服务*餐费"},
    )
    db_session.add(invoice)
    db_session.commit()

    report = service.build_report(db_session, 2026, 7)
    row = next(item for item in report["invoiceRows"] if item["invoiceId"] == invoice.id)

    assert row["expenseNature"] == "meal"
    assert row["expenseNatureLabel"] == "餐饮费"
    assert row["paymentSource"] == "personal"
    assert report["summary"]["personalInvoiceCount"] >= 1
    assert report["summary"]["expenseNatureCounts"]["餐饮费"] >= 1

def _add_adjustment_invoice(db_session, token: str, idx: int, company: str, total: str) -> TaxInvoice:
    invoice = TaxInvoice(
        invoice_key=f"pytest-corp-adj-{token}-{idx}",
        invoice_number=f"INV-ADJ-{token}-{idx}",
        direction="input",
        status="issued",
        issue_date=datetime(2026, 8, 10 + idx, tzinfo=timezone.utc),
        seller_name=f"勾选供应商-{token}",
        buyer_name=company,
        total_amount=Decimal(total),
        source_system="tax_export",
        raw={},
    )
    db_session.add(invoice)
    return invoice


def test_adjustment_persists_selected_invoices_and_recomputes_summary(db_session):
    """保存勾选版本后，报告只保留勾选发票并重算汇总；交付包与页面同口径。"""
    token = uuid4().hex[:10]
    company = f"测试公司-{token}"
    invoice_a = _add_adjustment_invoice(db_session, token, 0, company, "100.00")
    invoice_b = _add_adjustment_invoice(db_session, token, 1, company, "200.00")
    db_session.commit()

    base = service.build_report(db_session, 2026, 8, company=company)
    assert base["summary"]["invoiceCount"] == 2
    assert Decimal(base["summary"]["invoiceTotal"]) == Decimal("300.00")

    saved = service.save_corporate_payment_adjustment(
        db_session, company=company, year=2026, month=8,
        selected_keys=[invoice_a.invoice_number], actor="pytest-adjuster",
    )
    assert saved["version"] == 1
    assert saved["selectedKeys"] == [invoice_a.invoice_key]
    assert saved["selectedCount"] == 1
    assert saved["sourceCount"] == 2

    current = service.build_report(
        db_session, 2026, 8, company=company,
        selected_keys=list(saved["selectedKeys"]),
    )
    assert [row["invoiceKey"] for row in current["invoiceRows"]] == [invoice_a.invoice_key]
    assert [row["invoiceNumber"] for row in current["invoiceRows"]] == [invoice_a.invoice_number]
    assert current["rows"] == []
    assert current["summary"]["invoiceCount"] == 1
    assert Decimal(current["summary"]["invoiceTotal"]) == Decimal("100.00")
    assert Decimal(current["summary"]["outstandingTotal"]) == Decimal("100.00")
    assert current["summary"]["unpaidInvoiceCount"] == 1

    again = service.save_corporate_payment_adjustment(
        db_session, company=company, year=2026, month=8,
        selected_keys=[invoice_a.invoice_number, invoice_b.invoice_number],
    )
    assert again["version"] == 2
    latest = service.latest_adjustment(db_session, company=company, year=2026, month=8)
    assert latest is not None and latest.version == 2
    assert list(latest.selected_keys or []) == [invoice_a.invoice_key, invoice_b.invoice_key]

    full = service.build_report(
        db_session, 2026, 8, company=company,
        selected_keys=list(latest.selected_keys or []),
    )
    assert full["summary"]["invoiceCount"] == 2

    with pytest.raises(ValueError):
        service.save_corporate_payment_adjustment(
            db_session, company=company, year=2026, month=8,
            selected_keys=["不存在发票号"],
        )
    with pytest.raises(ValueError):
        service.save_corporate_payment_adjustment(
            db_session, company=company, year=2026, month=8,
            selected_keys=[],
        )



def test_duplicate_invoice_number_requires_stable_invoice_key_selection(db_session):
    """同号码不同 invoice_key 时，旧号码选择必须报歧义；canonical key 可以精确选择一张。"""
    token = uuid4().hex[:10]
    company = f"同号测试公司-{token}"
    shared_number = f"INV-SAME-{token}"
    invoice_a = TaxInvoice(
        invoice_key=f"CODE-A-{token}|{shared_number}",
        invoice_code=f"CODE-A-{token}",
        invoice_number=shared_number,
        direction="input",
        status="issued",
        issue_date=datetime(2026, 8, 18, tzinfo=timezone.utc),
        seller_name=f"同号供应商A-{token}",
        buyer_name=company,
        total_amount=Decimal("100.00"),
        source_system="tax_export",
        raw={},
    )
    invoice_b = TaxInvoice(
        invoice_key=f"CODE-B-{token}|{shared_number}",
        invoice_code=f"CODE-B-{token}",
        invoice_number=shared_number,
        direction="input",
        status="issued",
        issue_date=datetime(2026, 8, 19, tzinfo=timezone.utc),
        seller_name=f"同号供应商B-{token}",
        buyer_name=company,
        total_amount=Decimal("200.00"),
        source_system="tax_export",
        raw={},
    )
    db_session.add_all([invoice_a, invoice_b])
    db_session.commit()

    report = service.build_report(db_session, 2026, 8, company=company)
    assert {row["invoiceKey"] for row in report["invoiceRows"]} == {
        invoice_a.invoice_key,
        invoice_b.invoice_key,
    }
    assert {row["invoiceNumber"] for row in report["invoiceRows"]} == {shared_number}

    with pytest.raises(ValueError, match="对应多张发票"):
        service.save_corporate_payment_adjustment(
            db_session,
            company=company,
            year=2026,
            month=8,
            selected_keys=[shared_number],
            actor="pytest-adjuster",
        )

    saved = service.save_corporate_payment_adjustment(
        db_session,
        company=company,
        year=2026,
        month=8,
        selected_keys=[invoice_a.invoice_key],
        actor="pytest-adjuster",
    )
    assert saved["selectedKeys"] == [invoice_a.invoice_key]

    selected_report = service.build_report(
        db_session,
        2026,
        8,
        company=company,
        selected_keys=saved["selectedKeys"],
    )
    assert [row["invoiceId"] for row in selected_report["invoiceRows"]] == [invoice_a.id]
    assert Decimal(selected_report["summary"]["invoiceTotal"]) == Decimal("100.00")


def test_legacy_unique_invoice_number_selection_maps_to_invoice_key(db_session):
    """旧版本若保存的是唯一发票号码，读取/打包仍能无损映射到 canonical invoice_key。"""
    token = uuid4().hex[:10]
    company = f"旧选择兼容公司-{token}"
    invoice = _add_adjustment_invoice(db_session, token, 0, company, "123.00")
    db_session.commit()

    report = service.build_report(db_session, 2026, 8, company=company)
    canonical = service.canonical_report_selection_keys(report, [invoice.invoice_number])

    assert canonical == [invoice.invoice_key]
    selected = service.apply_invoice_selection(report, [invoice.invoice_number])
    assert [row["invoiceKey"] for row in selected["invoiceRows"]] == [invoice.invoice_key]


def test_legacy_null_bank_allocation_uses_full_invoice_amount_in_report(db_session):
    from app.models.bank import BankTransaction

    token = uuid4().hex[:10]
    invoice = TaxInvoice(
        invoice_key=f"pytest-corp-legacy-null-{token}",
        invoice_number=f"INV-NULL-{token}",
        direction="input",
        status="issued",
        issue_date=datetime(2026, 8, 20, tzinfo=timezone.utc),
        seller_name=f"历史分摊供应商-{token}",
        total_amount=Decimal("321.00"),
        source_system="tax_export",
        raw={},
    )
    txn = BankTransaction(
        txn_date=date(2026, 8, 21),
        direction="out",
        amount=Decimal("321.00"),
        counterparty_name=invoice.seller_name,
        fingerprint=f"pytest-corp-null-{token}",
        raw={},
    )
    db_session.add_all([invoice, txn])
    db_session.flush()
    db_session.add(TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="bank_transaction",
        target_id=txn.id,
        allocated_amount=None,
        match_method="manual",
        confirmed=True,
    ))
    db_session.commit()

    report = service.build_report(db_session, 2026, 8)
    row = next(item for item in report["invoiceRows"] if item["invoiceId"] == invoice.id)
    assert Decimal(row["invoiceCorporatePaidTotal"]) == Decimal("321.00")
    assert Decimal(row["invoiceOutstandingAmount"]) == Decimal("0")
    assert row["bankReconciliationStatus"] == "paid"
    assert Decimal(row["payments"][0]["allocatedAmount"]) == Decimal("321.00")


def test_corporate_report_uses_shanghai_invoice_month_boundaries(db_session):
    from zoneinfo import ZoneInfo
    from app.config import settings

    tz = ZoneInfo(settings.TZ)
    token = uuid4().hex[:10]
    august = TaxInvoice(
        invoice_key=f"pytest-corp-boundary-a-{token}",
        invoice_number=f"INV-BOUND-A-{token}",
        direction="input", status="issued",
        issue_date=datetime(2026, 8, 31, 23, 30, tzinfo=tz),
        seller_name=f"边界供应商A-{token}", total_amount=Decimal("10"),
    )
    september = TaxInvoice(
        invoice_key=f"pytest-corp-boundary-s-{token}",
        invoice_number=f"INV-BOUND-S-{token}",
        direction="input", status="issued",
        issue_date=datetime(2026, 9, 1, 0, 30, tzinfo=tz),
        seller_name=f"边界供应商S-{token}", total_amount=Decimal("20"),
    )
    db_session.add_all([august, september])
    db_session.commit()

    august_ids = {row["invoiceId"] for row in service.build_report(db_session, 2026, 8)["invoiceRows"]}
    september_ids = {row["invoiceId"] for row in service.build_report(db_session, 2026, 9)["invoiceRows"]}
    assert august.id in august_ids
    assert september.id not in august_ids
    assert september.id in september_ids


def test_report_exposes_manual_and_derived_payment_method_for_mixed_invoice(db_session):
    token = uuid4().hex[:10]
    account = BankAccount(
        account_no=f"ACCT-MIXED-{token}",
        account_name=f"混合付款账户-{token}",
        bank_name="测试银行",
        currency="CNY",
    )
    db_session.add(account)
    db_session.flush()
    txn = BankTransaction(
        account_id=account.id,
        txn_date=date(2026, 8, 18),
        direction="out",
        amount=Decimal("400.00"),
        counterparty_name=f"混合付款供应商-{token}",
        fingerprint=f"pytest-mixed-report-{token}",
        raw={},
    )
    invoice = TaxInvoice(
        invoice_key=f"pytest-mixed-report-{token}",
        invoice_number=f"INV-MIXED-{token}",
        direction="input",
        status="issued",
        issue_date=datetime(2026, 8, 15, tzinfo=timezone.utc),
        seller_name=f"混合付款供应商-{token}",
        total_amount=Decimal("1000.00"),
        payment_method="personal",
        source_system="tax_export",
        raw={},
    )
    db_session.add_all([txn, invoice])
    db_session.flush()
    db_session.add(TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="bank_transaction",
        target_id=txn.id,
        allocated_amount=Decimal("400.00"),
        match_method="manual",
        confirmed=True,
    ))
    db_session.commit()

    report = service.build_report(db_session, 2026, 8)
    row = next(item for item in report["invoiceRows"] if item["invoiceId"] == invoice.id)

    assert row["bankReconciliationStatus"] == "partial"
    assert row["manualPaymentMethod"] == "personal"
    assert row["paymentMethod"] == "mixed"
    assert row["paymentMethodLabel"] == "对公 + 个人垫付"


def test_report_exposes_platform_auto_debit_as_separate_payment_source(db_session):
    invoice = TaxInvoice(
        invoice_key=f"pytest-platform-auto-debit-{uuid4().hex[:10]}",
        invoice_number=f"INV-PLATFORM-{uuid4().hex[:10]}",
        direction="input",
        status="issued",
        total_amount=Decimal("1000.00"),
        payment_method="platform_auto_debit",
    )

    payment_source = service._payment_source(db_session, invoice, "unpaid")

    assert payment_source == (
        "platform_auto_debit",
        "平台自动扣款货款",
        "人工已确认平台自动扣款货款（不经过银行流水）",
    )
