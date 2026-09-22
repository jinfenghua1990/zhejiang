from decimal import Decimal
from io import BytesIO

import pytest

from app.adapters.tax_invoice_file import parse_tax_invoice_export
from app.services import tax_invoice_service as service
from uuid import uuid4


CSV = """税务系统发票清单
发票号码,发票代码,开票日期,销售方名称,销售方识别号,购买方名称,购买方识别号,金额,税额,价税合计,发票状态,进销项,关联订单号
INV-001, CODE-1,2026-09-01,供应商甲,91330000000000001A,本公司,91330000000000002B,100,13,113,正常,进项,PO-001
INV-002,CODE-2,2026-09-02,本公司,91330000000000002B,客户乙,91330000000000003C,200,26,226,正常,销项,
"""


def test_tax_export_auto_detects_fields_and_direction():
    parsed = parse_tax_invoice_export(CSV.encode(), "tax-list.csv", max_rows=100)

    assert parsed.sheet_name == "CSV"
    assert parsed.mapping["invoice_number"] == "发票号码"
    assert parsed.mapping["total_amount"] == "价税合计"
    assert parsed.direction_hint == "unknown"
    assert len(parsed.rows) == 2


def test_tax_xlsx_uses_base_sheet_and_validates_summary_only():
    from openpyxl import Workbook

    workbook = Workbook()
    summary = workbook.active
    summary.title = "信息汇总表"
    summary.append(["序号", "发票号码", "销方识别号", "销方名称", "开票日期", "金额", "税额", "价税合计"])
    summary.append([1, "INV-BASE-001", "91330000000000001A", "供应商甲", "2026-09-01", 100, 13, 113])
    summary.append(["合计行", None, None, None, None, 100, 13, 113])

    base = workbook.create_sheet("发票基础信息")
    base.append(["序号", "发票号码", "销方识别号", "销方名称", "开票日期", "金额", "税额", "价税合计", "发票状态"])
    base.append([1, "INV-BASE-001", "91330000000000001A", "供应商甲", "2026-09-01", 100, 13, 113, "正常"])
    base.append(["合计行", None, None, None, None, 100, 13, 113, None])

    output = BytesIO()
    workbook.save(output)
    workbook.close()

    parsed = parse_tax_invoice_export(output.getvalue(), "tax-base.xlsx", max_rows=100)

    assert parsed.sheet_name == "发票基础信息"
    assert len(parsed.rows) == 1
    assert parsed.rows[0]["发票号码"] == "INV-BASE-001"
    assert parsed.ignored_row_count == 1
    assert parsed.summary_sheet_name == "信息汇总表"
    assert parsed.summary_validation["status"] == "matched"
    assert parsed.summary_validation["baseInvoiceCount"] == 1
    assert parsed.summary_validation["summaryInvoiceCount"] == 1
    assert parsed.summary_validation["baseTotals"]["total_amount"] == "113.00"


def test_tax_xlsx_rejects_summary_without_base_sheet():
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "信息汇总表"
    workbook.active.append(["发票号码", "金额", "税额", "价税合计"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(ValueError, match="必须包含「发票基础信息」"):
        parse_tax_invoice_export(output.getvalue(), "summary-only.xlsx", max_rows=100)


def test_tax_import_normalizes_rows_is_idempotent_and_links_explicit_order(
    db_session, monkeypatch, tmp_path
):
    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice, TaxInvoiceImport, TaxInvoiceImportRecord, TaxInvoiceLink

    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    token = uuid4().hex[:10]
    order_no = f"PO-TEST-{token}"
    invoice_1 = f"INV-TEST-A-{token}"
    invoice_2 = f"INV-TEST-B-{token}"
    content = CSV.replace("PO-001", order_no).replace("INV-001", invoice_1).replace("INV-002", invoice_2).encode()
    baseline_invoices = db_session.query(TaxInvoice).count()
    po = ExternalPurchaseOrder(external_order_id=order_no, paid_amount="113")
    db_session.add(po)
    db_session.commit()

    batch, duplicate = service.import_export(
        db_session,
        content=content,
        original_name="tax-list.csv",
        actor="pytest",
        period_year=2026,
        period_month=9,
    )
    same_batch, same_duplicate = service.import_export(
        db_session,
        content=content,
        original_name="tax-list.csv",
        actor="pytest",
    )

    assert duplicate is False
    assert same_duplicate is True
    assert same_batch.id == batch.id
    assert batch.row_count == 2
    assert batch.recognized_row_count == 2
    assert batch.needs_review_count == 0
    assert db_session.query(TaxInvoice).count() == baseline_invoices + 2
    assert db_session.query(TaxInvoiceImport).filter_by(id=batch.id).one().status == "parsed"
    assert db_session.query(TaxInvoiceImportRecord).filter_by(import_id=batch.id).count() == 2
    invoice = db_session.query(TaxInvoice).filter_by(invoice_number=invoice_1).one()
    assert invoice.direction == "input"
    assert invoice.status == "issued"
    assert str(invoice.total_amount) == "113.0000"
    assert invoice.match_status == "matched"
    assert db_session.query(TaxInvoiceLink).filter_by(invoice_id=invoice.id).count() == 1

    db_session.query(TaxInvoiceLink).filter_by(invoice_id=invoice.id).delete()
    db_session.query(TaxInvoiceImportRecord).filter_by(import_id=batch.id).delete()
    db_session.query(TaxInvoice).filter(TaxInvoice.id.in_((
        row.id for row in db_session.query(TaxInvoice).filter_by(source_import_id=batch.id).all()
    ))).delete(synchronize_session=False)
    db_session.delete(batch)
    db_session.delete(po)
    db_session.commit()


def test_tax_import_keeps_unrecognized_row_for_review(db_session, monkeypatch, tmp_path):
    from app.models.tax import TaxInvoice, TaxInvoiceImportRecord

    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    content = "发票号码,开票日期,价税合计\nINV-003,,\n".encode()
    batch, _ = service.import_export(
        db_session, content=content, original_name="needs-review.csv", actor="pytest"
    )

    assert batch.status == "needs_review"
    assert batch.recognized_row_count == 0
    assert batch.needs_review_count == 1
    row = db_session.query(TaxInvoiceImportRecord).filter_by(import_id=batch.id).one()
    assert row.recognition_status == "needs_review"
    assert row.invoice_id is None
    assert db_session.query(TaxInvoice).filter_by(invoice_number="INV-003").count() == 0

    db_session.query(TaxInvoiceImportRecord).filter_by(import_id=batch.id).delete()
    db_session.delete(batch)
    db_session.commit()


def test_invoice_category_is_single_source_of_truth_for_processing_status(db_session):
    """category 是唯一事实：写 v2 分类后 processing_status 同步派生
    （运营成本三分类→required；报销两分类与 excluded→not_required；空→pending）。"""
    from app.models.tax import TaxInvoice, TaxInvoiceImport

    batch = TaxInvoiceImport(
        original_name=f"derive-{uuid4().hex}.xlsx",
        stored_path=f"/tmp/derive-{uuid4().hex}.xlsx",
        sha256=uuid4().hex + uuid4().hex,
        lifecycle="active",
    )
    db_session.add(batch)
    db_session.flush()

    def _invoice(prefix: str) -> TaxInvoice:
        row = TaxInvoice(
            invoice_key=f"pytest|{prefix}-{uuid4().hex[:8]}",
            invoice_number=f"{prefix}-{uuid4().hex[:8]}",
            direction="input",
            seller_name="供应商丙",
            total_amount=Decimal("113"),
            source_import_id=batch.id,
            source_row_index=1,
        )
        db_session.add(row)
        return row

    goods = _invoice("DERIVE-GOODS")
    reimb = _invoice("DERIVE-REIMB")
    platform = _invoice("DERIVE-PLAT")
    excluded = _invoice("DERIVE-EXCLUDED")
    empty = _invoice("DERIVE-EMPTY")
    db_session.flush()

    service.set_invoice_categories(db_session, [goods.id], "goods", actor="pytest")
    service.set_invoice_categories(db_session, [reimb.id], "reimburse_operating", actor="pytest")
    service.set_invoice_categories(db_session, [platform.id], "platform_fee", actor="pytest")
    service.set_invoice_categories(db_session, [excluded.id], "excluded", actor="pytest")
    # 旧四值 key 已废弃，必须拒绝。
    with pytest.raises(ValueError):
        service.set_invoice_categories(db_session, [goods.id], "goods_payment", actor="pytest")

    rows = {row.id: row for row in db_session.query(TaxInvoice).filter(TaxInvoice.id.in_(
        [goods.id, reimb.id, platform.id, excluded.id, empty.id]
    )).all()}
    assert rows[goods.id].processing_status == "required"
    assert rows[reimb.id].processing_status == "not_required"
    assert rows[platform.id].processing_status == "required"
    assert rows[excluded.id].processing_status == "not_required"
    assert rows[empty.id].category == ""
    assert rows[empty.id].processing_status == "pending"

    listed = {item["id"]: item for item in service.list_invoices(db_session, direction="input", limit=500)}
    assert listed[goods.id]["processingStatus"] == "required"
    assert listed[goods.id]["category"] == "goods"
    assert listed[goods.id]["categoryLabel"] == "运营成本：货款发票"
    assert listed[reimb.id]["processingStatus"] == "not_required"
    assert listed[platform.id]["processingStatus"] == "required"
    assert listed[excluded.id]["processingStatus"] == "not_required"
    assert listed[empty.id]["processingStatus"] == "pending"
    assert listed[empty.id]["categoryLabel"] == "待判断（未分类）"

    # 旧状态按钮接口行为等价：设状态时按映射写 v2 分类，处理结论保持一致。
    service.set_processing_status(db_session, reimb.id, "required", actor="pytest")
    assert db_session.get(TaxInvoice, reimb.id).category == "goods"
    assert service.serialize_invoice(db_session.get(TaxInvoice, reimb.id))["processingStatus"] == "required"
    service.set_processing_status(db_session, goods.id, "not_required", actor="pytest")
    assert db_session.get(TaxInvoice, goods.id).category == "reimburse_operating"
    assert service.serialize_invoice(db_session.get(TaxInvoice, goods.id))["processingStatus"] == "not_required"
    service.set_processing_status(db_session, excluded.id, "pending", actor="pytest")
    assert db_session.get(TaxInvoice, excluded.id).category == ""
    assert service.serialize_invoice(db_session.get(TaxInvoice, excluded.id))["processingStatus"] == "pending"


def test_output_invoice_category_enum_and_direction_guard(db_session):
    """销项分类独立枚举：buyer_sales/platform_service/空=待判断；进项传销项 key、
    销项传进项 key 都必须拒绝；销项不按明细兜底推断，空即待判断。"""
    from app.models.tax import TaxInvoice, TaxInvoiceImport

    batch = TaxInvoiceImport(
        original_name=f"output-{uuid4().hex}.xlsx",
        stored_path=f"/tmp/output-{uuid4().hex}.xlsx",
        sha256=uuid4().hex + uuid4().hex,
        lifecycle="active",
    )
    db_session.add(batch)
    db_session.flush()

    def _invoice(prefix: str, direction: str) -> TaxInvoice:
        row = TaxInvoice(
            invoice_key=f"pytest|{prefix}-{uuid4().hex[:8]}",
            invoice_number=f"{prefix}-{uuid4().hex[:8]}",
            direction=direction,
            seller_name="供应商丁",
            total_amount=Decimal("226"),
            source_import_id=batch.id,
            source_row_index=1,
        )
        db_session.add(row)
        return row

    output = _invoice("OUT-BUYER", "output")
    output2 = _invoice("OUT-PLAT", "output")
    output3 = _invoice("OUT-EMPTY", "output")
    input_row = _invoice("IN-GUARD", "input")
    db_session.flush()

    service.set_invoice_categories(db_session, [output.id], "buyer_sales", actor="pytest")
    service.set_invoice_categories(db_session, [output2.id], "platform_service", actor="pytest")
    service.set_invoice_categories(db_session, [output3.id], "", actor="pytest")
    # 销项传进项 key → 拒绝；进项传销项 key → 拒绝。
    with pytest.raises(ValueError):
        service.set_invoice_categories(db_session, [output.id], "goods", actor="pytest")
    with pytest.raises(ValueError):
        service.set_invoice_categories(db_session, [input_row.id], "buyer_sales", actor="pytest")

    rows = {row.id: row for row in db_session.query(TaxInvoice).filter(TaxInvoice.id.in_(
        [output.id, output2.id, output3.id, input_row.id]
    )).all()}
    assert rows[output.id].category == "buyer_sales"
    assert rows[output.id].processing_status == "pending"
    assert rows[output2.id].category == "platform_service"
    assert rows[output2.id].processing_status == "pending"
    assert rows[output3.id].category == ""
    assert rows[output3.id].processing_status == "pending"
    # 拒绝后进项行保持未分类，不被误写。
    assert rows[input_row.id].category == ""

    listed = {item["id"]: item for item in service.list_invoices(db_session, direction="output", limit=500)}
    assert listed[output.id]["category"] == "buyer_sales"
    assert listed[output.id]["categoryLabel"] == "给买家开发票"
    assert listed[output2.id]["categoryLabel"] == "给平台开服务费"
    assert listed[output3.id]["category"] == ""
    assert listed[output3.id]["categoryLabel"] == "待判断"

    # 置回空 = 待判断，幂等可逆。
    service.set_invoice_categories(db_session, [output.id], "", actor="pytest")
    assert db_session.get(TaxInvoice, output.id).category == ""
    assert service.serialize_invoice(db_session.get(TaxInvoice, output.id))["categoryLabel"] == "待判断"



def test_invoice_business_match_and_bank_payment_status_are_independent(db_session):
    """同一张进项发票的业务匹配与银行付款核对必须各算各的，互不污染。"""
    from datetime import datetime, timezone

    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice, TaxInvoiceLink

    purchase = ExternalPurchaseOrder(
        external_order_id=f"DOMAIN-PO-{uuid4().hex[:10]}",
        platform="other",
        supplier_name="独立域供应商",
        order_amount=Decimal("1000.00"),
        paid_amount=Decimal("1000.00"),
    )
    db_session.add(purchase)
    db_session.flush()

    invoice = TaxInvoice(
        invoice_key=f"pytest-domain-{uuid4().hex}",
        invoice_number=f"DOMAIN-{uuid4().hex[:10]}",
        direction="input",
        status="issued",
        issue_date=datetime(2026, 9, 20, tzinfo=timezone.utc),
        seller_name="独立域供应商",
        total_amount=Decimal("1000.00"),
        match_status="unmatched",  # 故意放旧缓存值，序列化必须以真实业务链接为准。
        match_note="",
        raw={},
    )
    db_session.add(invoice)
    db_session.flush()

    business_link = TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="external_purchase_order",
        target_id=purchase.id,
        allocated_amount=Decimal("600.00"),
        match_method="manual",
        confirmed=True,
        note="采购业务匹配",
    )
    bank_link = TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="bank_transaction",
        target_id=123456789,
        allocated_amount=Decimal("400.00"),
        match_method="manual",
        confirmed=True,
        note="银行付款核对",
    )
    db_session.add_all([business_link, bank_link])
    db_session.flush()

    payload = service.serialize_invoice(invoice, db=db_session)
    assert payload["businessMatchStatus"] == "partial"
    assert payload["matchStatus"] == "partial"
    assert payload["businessMatchedAmount"] == "600.0000"
    assert payload["businessRemainingAmount"] == "400.0000"
    assert payload["bankPaymentStatus"] == "partial"
    assert payload["bankPaidAmount"] == "400.00"
    assert payload["bankRemainingAmount"] == "600.00"

    # 银行核对改成付清，发票业务匹配仍然只能保持 600/1000 的 partial。
    bank_link.allocated_amount = Decimal("1000.00")
    db_session.flush()
    payload = service.serialize_invoice(invoice, db=db_session)
    assert payload["businessMatchStatus"] == "partial"
    assert payload["matchStatus"] == "partial"
    assert payload["businessMatchedAmount"] == "600.0000"
    assert payload["bankPaymentStatus"] == "matched"
    assert payload["bankPaidAmount"] == "1000.00"

    # 解除银行核对也不能改变发票业务匹配。
    bank_link.match_method = "rejected"
    bank_link.confirmed = False
    db_session.flush()
    payload = service.serialize_invoice(invoice, db=db_session)
    assert payload["businessMatchStatus"] == "partial"
    assert payload["businessMatchedAmount"] == "600.0000"
    assert payload["bankPaymentStatus"] == "unmatched"
    assert payload["bankPaidAmount"] == "0.00"

def test_list_invoices_match_filter_uses_live_business_domain_not_cached_status(db_session):
    """服务端筛选必须与页面实时 businessMatchStatus 完全一致，银行链接不得参与。"""
    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice, TaxInvoiceLink

    purchase = ExternalPurchaseOrder(
        external_order_id=f"FILTER-PO-{uuid4().hex[:10]}",
        platform="other",
        supplier_name="筛选口径供应商",
        order_amount=Decimal("1000.00"),
        paid_amount=Decimal("1000.00"),
    )
    db_session.add(purchase)
    db_session.flush()

    invoice = TaxInvoice(
        invoice_key=f"pytest-filter-domain-{uuid4().hex}",
        invoice_number=f"FILTER-{uuid4().hex[:10]}",
        direction="input",
        status="issued",
        seller_name="筛选口径供应商",
        total_amount=Decimal("1000.00"),
        match_status="unmatched",  # 故意保留旧缓存值
        raw={},
    )
    db_session.add(invoice)
    db_session.flush()
    db_session.add_all([
        TaxInvoiceLink(
            invoice_id=invoice.id,
            target_type="external_purchase_order",
            target_id=purchase.id,
            allocated_amount=Decimal("600.00"),
            match_method="manual",
            confirmed=True,
        ),
        TaxInvoiceLink(
            invoice_id=invoice.id,
            target_type="bank_transaction",
            target_id=777002,
            allocated_amount=Decimal("1000.00"),
            match_method="manual",
            confirmed=True,
        ),
    ])
    db_session.flush()

    partial_ids = {
        row["id"]
        for row in service.list_invoices(
            db_session, direction="input", match_status="partial", limit=500
        )
    }
    matched_ids = {
        row["id"]
        for row in service.list_invoices(
            db_session, direction="input", match_status="matched", limit=500
        )
    }
    unmatched_ids = {
        row["id"]
        for row in service.list_invoices(
            db_session, direction="input", match_status="unmatched", limit=500
        )
    }

    assert invoice.id in partial_ids
    assert invoice.id not in matched_ids
    assert invoice.id not in unmatched_ids

    # 未确认业务链接不计入业务匹配；银行已付清也不能改变业务筛选结果。
    business_link = db_session.query(TaxInvoiceLink).filter_by(
        invoice_id=invoice.id, target_type="external_purchase_order"
    ).one()
    business_link.confirmed = False
    db_session.flush()

    unmatched_ids = {
        row["id"]
        for row in service.list_invoices(
            db_session, direction="input", match_status="unmatched", limit=500
        )
    }
    assert invoice.id in unmatched_ids

def test_multiline_tax_invoice_links_only_after_final_amount(db_session, monkeypatch, tmp_path):
    """一票多行必须先汇总整票金额，再按订单号只关联一次。"""
    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice, TaxInvoiceLink

    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    token = uuid4().hex[:10]
    invoice_no = f"INV-MULTI-{token}"
    order_no = f"PO-MULTI-{token}"
    content = f"""发票号码,开票日期,销售方名称,购买方名称,项目名称,金额,税额,价税合计,发票状态,进销项,关联订单号
{invoice_no},2026-09-03,多行供应商,本公司,商品A,60,7.8,67.8,正常,进项,{order_no}
{invoice_no},2026-09-03,多行供应商,本公司,商品B,40,5.2,45.2,正常,进项,{order_no}
""".encode()
    db_session.add(ExternalPurchaseOrder(
        external_order_id=order_no,
        platform="other",
        paid_amount=Decimal("113"),
        order_amount=Decimal("113"),
    ))
    db_session.commit()

    batch, duplicate = service.import_export(
        db_session,
        content=content,
        original_name=f"multi-{token}.csv",
        actor="pytest",
        period_year=2026,
        period_month=9,
    )
    assert duplicate is False
    assert batch.recognized_row_count == 2
    assert batch.matched_row_count == 2
    assert batch.needs_review_count == 0

    invoice = db_session.query(TaxInvoice).filter_by(invoice_number=invoice_no).one()
    assert invoice.total_amount == Decimal("113.0000")
    assert invoice.match_status == "matched"
    links = db_session.query(TaxInvoiceLink).filter_by(invoice_id=invoice.id).all()
    assert len(links) == 1
    assert links[0].allocated_amount == Decimal("113.0000")
    assert links[0].match_method == "source_ref"
    assert links[0].confirmed is True


def test_multiline_tax_invoice_with_multiple_order_refs_requires_review(db_session, monkeypatch, tmp_path):
    """同一张发票的不同明细行出现不同订单号时不能逐行自动挂单。"""
    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice, TaxInvoiceImportRecord, TaxInvoiceLink

    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    token = uuid4().hex[:10]
    invoice_no = f"INV-MULTI-REF-{token}"
    order_a = f"PO-A-{token}"
    order_b = f"PO-B-{token}"
    content = f"""发票号码,开票日期,销售方名称,购买方名称,项目名称,金额,税额,价税合计,发票状态,进销项,关联订单号
{invoice_no},2026-09-04,多订单供应商,本公司,商品A,60,7.8,67.8,正常,进项,{order_a}
{invoice_no},2026-09-04,多订单供应商,本公司,商品B,40,5.2,45.2,正常,进项,{order_b}
""".encode()
    db_session.add_all([
        ExternalPurchaseOrder(external_order_id=order_a, platform="other", paid_amount=Decimal("67.8"), order_amount=Decimal("67.8")),
        ExternalPurchaseOrder(external_order_id=order_b, platform="other", paid_amount=Decimal("45.2"), order_amount=Decimal("45.2")),
    ])
    db_session.commit()

    batch, _ = service.import_export(
        db_session,
        content=content,
        original_name=f"multi-ref-{token}.csv",
        actor="pytest",
        period_year=2026,
        period_month=9,
    )
    assert batch.status == "needs_review"
    assert batch.recognized_row_count == 2
    assert batch.matched_row_count == 0
    assert batch.needs_review_count == 2

    invoice = db_session.query(TaxInvoice).filter_by(invoice_number=invoice_no).one()
    assert invoice.match_status == "needs_review"
    assert "多个不同关联订单号" in invoice.match_note
    assert db_session.query(TaxInvoiceLink).filter_by(invoice_id=invoice.id).count() == 0
    records = db_session.query(TaxInvoiceImportRecord).filter_by(import_id=batch.id).all()
    assert {row.recognition_status for row in records} == {"needs_review"}




def test_reimport_multiple_source_refs_preserves_closed_manual_business_match(
    db_session, monkeypatch, tmp_path
):
    """重复导入的来源提示变歧义时，已闭合人工关联优先于导入提示。"""
    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice, TaxInvoiceLink

    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    token = uuid4().hex[:10]
    invoice_no = f"INV-MANUAL-REIMPORT-{token}"
    first_content = f"""发票号码,开票日期,销售方名称,购买方名称,项目名称,金额,税额,价税合计,发票状态,进销项,关联订单号
{invoice_no},2026-09-04,人工关联供应商,本公司,商品A,100,13,113,正常,进项,
""".encode()
    first_batch, _ = service.import_export(
        db_session,
        content=first_content,
        original_name=f"manual-first-{token}.csv",
        actor="pytest",
        period_year=2026,
        period_month=9,
    )
    assert first_batch.status == "parsed"

    invoice = db_session.query(TaxInvoice).filter_by(invoice_number=invoice_no).one()
    manual_po = ExternalPurchaseOrder(
        external_order_id=f"PO-MANUAL-REIMPORT-{token}",
        platform="other",
        supplier_name="人工关联供应商",
        paid_amount=Decimal("113"),
        order_amount=Decimal("113"),
    )
    db_session.add(manual_po)
    db_session.commit()
    service.link_purchase_order(
        db_session,
        invoice_id=invoice.id,
        target_type="external_purchase_order",
        target_id=manual_po.id,
        allocated_amount=Decimal("113"),
        actor="pytest",
    )
    db_session.refresh(invoice)
    assert invoice.match_status == "matched"

    ref_a = f"PO-REF-A-{token}"
    ref_b = f"PO-REF-B-{token}"
    second_content = f"""发票号码,开票日期,销售方名称,购买方名称,项目名称,金额,税额,价税合计,发票状态,进销项,关联订单号
{invoice_no},2026-09-04,人工关联供应商,本公司,商品A,60,7.8,67.8,正常,进项,{ref_a}
{invoice_no},2026-09-04,人工关联供应商,本公司,商品B,40,5.2,45.2,正常,进项,{ref_b}
""".encode()
    second_batch, duplicate = service.import_export(
        db_session,
        content=second_content,
        original_name=f"manual-second-{token}.csv",
        actor="pytest",
        period_year=2026,
        period_month=9,
    )

    assert duplicate is False
    db_session.refresh(invoice)
    assert invoice.total_amount == Decimal("113.0000")
    assert invoice.match_status == "matched"
    assert second_batch.needs_review_count == 0
    assert second_batch.status == "parsed"
    manual_links = db_session.query(TaxInvoiceLink).filter_by(
        invoice_id=invoice.id,
        match_method="manual",
        confirmed=True,
    ).all()
    assert len(manual_links) == 1
    assert manual_links[0].target_id == manual_po.id


def test_explicit_missing_order_ref_marks_batch_for_review(db_session, monkeypatch, tmp_path):
    """税务清单明确给了订单号但系统找不到时，不能静默当作正常解析完成。"""
    from app.models.tax import TaxInvoice

    monkeypatch.setattr(service.settings, "DATA_DIR", str(tmp_path))
    token = uuid4().hex[:10]
    missing_order = f"PO-MISSING-{token}"
    invoice_no = f"INV-MISSING-{token}"
    content = CSV.replace("PO-001", missing_order).replace("INV-001", invoice_no).encode()

    batch, _ = service.import_export(
        db_session,
        content=content,
        original_name=f"missing-order-{token}.csv",
        actor="pytest",
        period_year=2026,
        period_month=9,
    )
    assert batch.status == "needs_review"
    assert batch.needs_review_count == 1
    assert batch.matched_row_count == 0
    invoice = db_session.query(TaxInvoice).filter_by(invoice_number=invoice_no).one()
    assert invoice.match_status == "needs_review"
    assert missing_order in invoice.match_note



def test_source_ref_ambiguity_deactivates_stale_auto_link(db_session):
    """来源订单号后来变成歧义时，历史自动关联不能继续假装已确认。"""
    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice, TaxInvoiceLink

    token = uuid4().hex[:10]
    order_no = f"PO-AMB-{token}"
    invoice = TaxInvoice(
        invoice_key=f"INV-AMB-{token}",
        invoice_number=f"INV-AMB-{token}",
        direction="input",
        status="issued",
        total_amount=Decimal("100"),
        match_status="matched",
    )
    po_a = ExternalPurchaseOrder(
        external_order_id=order_no, platform="1688",
        paid_amount=Decimal("100"), order_amount=Decimal("100"),
    )
    po_b = ExternalPurchaseOrder(
        external_order_id=order_no, platform="pdd",
        paid_amount=Decimal("100"), order_amount=Decimal("100"),
    )
    db_session.add_all([invoice, po_a, po_b])
    db_session.flush()
    stale = TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="external_purchase_order",
        target_id=po_a.id,
        allocated_amount=Decimal("100"),
        match_method="source_ref",
        confidence=Decimal("1"),
        confirmed=True,
    )
    db_session.add(stale)
    db_session.commit()

    outcome = service._auto_link(db_session, invoice, order_no, "input")
    db_session.flush()

    assert outcome is False
    assert stale.confirmed is False
    assert stale.confidence is None
    assert invoice.match_status == "needs_review"
    serialized = service.serialize_invoice(invoice, db=db_session)
    assert serialized["businessMatchStatus"] == "needs_review"




def test_source_ref_ambiguity_does_not_override_manual_confirmed_match(db_session):
    """导入来源变歧义时，只撤销自动证据；完整人工关联仍应保持 matched。"""
    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice, TaxInvoiceLink

    token = uuid4().hex[:10]
    ambiguous_no = f"PO-AMB-MANUAL-{token}"
    invoice = TaxInvoice(
        invoice_key=f"INV-AMB-MANUAL-{token}",
        invoice_number=f"INV-AMB-MANUAL-{token}",
        direction="input",
        status="issued",
        total_amount=Decimal("100"),
        match_status="matched",
    )
    manual_po = ExternalPurchaseOrder(
        external_order_id=f"PO-MANUAL-{token}",
        platform="other",
        paid_amount=Decimal("100"),
        order_amount=Decimal("100"),
    )
    ambiguous_a = ExternalPurchaseOrder(
        external_order_id=ambiguous_no,
        platform="1688",
        paid_amount=Decimal("100"),
        order_amount=Decimal("100"),
    )
    ambiguous_b = ExternalPurchaseOrder(
        external_order_id=ambiguous_no,
        platform="pdd",
        paid_amount=Decimal("100"),
        order_amount=Decimal("100"),
    )
    db_session.add_all([invoice, manual_po, ambiguous_a, ambiguous_b])
    db_session.flush()
    manual = TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="external_purchase_order",
        target_id=manual_po.id,
        allocated_amount=Decimal("100"),
        match_method="manual",
        confidence=Decimal("1"),
        confirmed=True,
        note="人工已确认",
    )
    stale_auto = TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="external_purchase_order",
        target_id=ambiguous_a.id,
        allocated_amount=Decimal("100"),
        match_method="source_ref",
        confidence=Decimal("1"),
        confirmed=True,
    )
    db_session.add_all([manual, stale_auto])
    db_session.commit()

    outcome = service._auto_link(db_session, invoice, ambiguous_no, "input")
    db_session.flush()

    assert outcome is False
    assert stale_auto.confirmed is False
    assert manual.confirmed is True
    assert invoice.match_status == "matched"
    assert "命中多个采购渠道" in invoice.match_note
    assert service.serialize_invoice(invoice, db=db_session)["businessMatchStatus"] == "matched"


def test_noneligible_invoice_with_manual_business_link_requires_review(db_session):
    """票据失效后若仍有人工业务分摊，必须暴露冲突，不能伪装成 unmatched。"""
    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice, TaxInvoiceLink

    token = uuid4().hex[:10]
    order_no = f"PO-RED-MANUAL-{token}"
    invoice = TaxInvoice(
        invoice_key=f"INV-RED-MANUAL-{token}",
        invoice_number=f"INV-RED-MANUAL-{token}",
        direction="input",
        status="red",
        total_amount=Decimal("-100"),
        match_status="matched",
    )
    po = ExternalPurchaseOrder(
        external_order_id=order_no,
        platform="other",
        paid_amount=Decimal("100"),
        order_amount=Decimal("100"),
    )
    db_session.add_all([invoice, po])
    db_session.flush()
    manual = TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="external_purchase_order",
        target_id=po.id,
        allocated_amount=Decimal("100"),
        match_method="manual",
        confidence=Decimal("1"),
        confirmed=True,
        note="人工已确认",
    )
    db_session.add(manual)
    db_session.commit()

    outcome = service._auto_link(db_session, invoice, order_no, "input")
    db_session.flush()

    assert outcome is False
    assert manual.confirmed is True
    assert invoice.match_status == "needs_review"
    assert "超额" in invoice.match_note or "非有效正数发票" in invoice.match_note


def test_noneligible_invoice_deactivates_historical_source_ref(db_session):
    """红冲/非正数票不能继续沿用历史 source_ref 自动匹配状态。"""
    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice, TaxInvoiceLink

    token = uuid4().hex[:10]
    order_no = f"PO-RED-{token}"
    invoice = TaxInvoice(
        invoice_key=f"INV-RED-{token}",
        invoice_number=f"INV-RED-{token}",
        direction="input",
        status="red",
        total_amount=Decimal("-100"),
        match_status="matched",
    )
    po = ExternalPurchaseOrder(
        external_order_id=order_no, platform="other",
        paid_amount=Decimal("100"), order_amount=Decimal("100"),
    )
    db_session.add_all([invoice, po])
    db_session.flush()
    stale = TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="external_purchase_order",
        target_id=po.id,
        allocated_amount=Decimal("100"),
        match_method="source_ref",
        confidence=Decimal("1"),
        confirmed=True,
    )
    db_session.add(stale)
    db_session.commit()

    outcome = service._auto_link(db_session, invoice, order_no, "input")
    db_session.flush()

    assert outcome is False
    assert stale.confirmed is False
    assert invoice.match_status == "unmatched"


def test_red_blue_pairing_partial_then_full_uses_accounting_net(db_session):
    """蓝字+红字分别留存；部分红冲只剩未冲余额，全额红冲余额归零，红字分类继承蓝字。"""
    from datetime import datetime, timezone

    from app.models.tax import TaxInvoice

    token = uuid4().hex[:10]
    blue_no = f"BLUE-{token}"
    blue = TaxInvoice(
        invoice_key=f"pytest-blue-{token}",
        invoice_number=blue_no,
        direction="input",
        status="issued",
        issue_date=datetime(2026, 7, 1, tzinfo=timezone.utc),
        seller_name=f"红冲测试供应商-{token}",
        seller_tax_id=f"TAX-{token}",
        buyer_name="测试公司",
        total_amount=Decimal("1000.00"),
        category="goods",
        processing_status="required",
        raw={"是否正数发票": "是", "发票状态": "正常"},
    )
    red_a = TaxInvoice(
        invoice_key=f"pytest-red-a-{token}",
        invoice_number=f"RED-A-{token}",
        direction="input",
        status="red",
        issue_date=datetime(2026, 7, 5, tzinfo=timezone.utc),
        seller_name=blue.seller_name,
        seller_tax_id=blue.seller_tax_id,
        buyer_name=blue.buyer_name,
        total_amount=Decimal("-300.00"),
        raw={
            "是否正数发票": "否",
            "备注": f"被红冲蓝字数电发票号码：{blue_no} 红字发票信息确认单编号：HZ-A-{token}",
        },
    )
    db_session.add_all([blue, red_a])
    db_session.commit()

    context = service.red_accounting_context(db_session, [blue, red_a])
    assert context[blue.id]["invoiceColor"] == "blue"
    assert context[blue.id]["redStatus"] == "partially_red_offset"
    assert context[blue.id]["invoiceStatusLabel"] == "蓝字发票（部分红冲）"
    assert Decimal(context[blue.id]["redOffsetAmount"]) == Decimal("300.00")
    assert Decimal(context[blue.id]["remainingAfterRedAmount"]) == Decimal("700.00")
    assert context[red_a.id]["invoiceColor"] == "red"
    assert context[red_a.id]["redStatus"] == "red_invoice"
    assert context[red_a.id]["redRelatedInvoiceId"] == blue.id
    assert service.effective_invoice_amount_after_red(db_session, blue) == Decimal("700.00")
    assert service.effective_invoice_amount_after_red(db_session, red_a) == Decimal("0")
    assert service.is_bank_payment_reconciliation_eligible(blue, db=db_session) is True

    red_payload = service.serialize_invoice(red_a, db=db_session)
    assert red_payload["category"] == "goods"
    assert red_payload["categoryInherited"] is True
    assert red_payload["categoryInheritedFromInvoiceId"] == blue.id
    with pytest.raises(ValueError, match="必须继承对应蓝字发票"):
        service.set_invoice_categories(db_session, [red_a.id], "reimburse_operating", actor="pytest")

    red_b = TaxInvoice(
        invoice_key=f"pytest-red-b-{token}",
        invoice_number=f"RED-B-{token}",
        direction="input",
        status="red",
        issue_date=datetime(2026, 7, 6, tzinfo=timezone.utc),
        seller_name=blue.seller_name,
        seller_tax_id=blue.seller_tax_id,
        buyer_name=blue.buyer_name,
        total_amount=Decimal("-700.00"),
        raw={
            "是否正数发票": "否",
            "备注": f"被红冲蓝字发票号码：{blue_no} 红字发票信息确认单编号：HZ-B-{token}",
        },
    )
    db_session.add(red_b)
    db_session.commit()

    context = service.red_accounting_context(db_session, [blue, red_a, red_b])
    assert context[blue.id]["redStatus"] == "fully_red_offset"
    assert context[blue.id]["invoiceStatusLabel"] == "蓝字发票（已全额红冲）"
    assert Decimal(context[blue.id]["redOffsetAmount"]) == Decimal("1000.00")
    assert Decimal(context[blue.id]["remainingAfterRedAmount"]) == Decimal("0.00")
    assert service.effective_invoice_amount_after_red(db_session, blue) == Decimal("0")
    assert service.is_bank_payment_reconciliation_eligible(blue, db=db_session) is False
    assert "已全额红冲" in service.bank_payment_reconciliation_ineligible_reason(blue, db=db_session)


def test_red_invoice_without_blue_is_kept_but_flagged(db_session):
    """红字票不能丢；找不到蓝字原票时保留真实负数，并明确标记待关联异常。"""
    from datetime import datetime, timezone

    from app.models.tax import TaxInvoice

    token = uuid4().hex[:10]
    red = TaxInvoice(
        invoice_key=f"pytest-red-unpaired-{token}",
        invoice_number=f"RED-U-{token}",
        direction="input",
        status="red",
        issue_date=datetime(2026, 7, 10, tzinfo=timezone.utc),
        seller_name=f"未配对供应商-{token}",
        total_amount=Decimal("-123.45"),
        raw={
            "是否正数发票": "否",
            "备注": f"被红冲蓝字数电发票号码：MISSING-{token}",
        },
    )
    db_session.add(red)
    db_session.commit()

    context = service.red_accounting_context(db_session, [red])[red.id]
    assert context["invoiceStatusLabel"] == "红字发票（待关联蓝字）"
    assert context["redPairStatus"] == "unpaired"
    assert context["accountingNetIncluded"] is True
    assert Decimal(context["accountingNetAmount"]) == Decimal("-123.45")
    assert "未识别到对应蓝字" in context["accountingException"]


def test_positive_blue_marked_red_offset_is_not_misparsed_as_red_document():
    """“已红冲”是蓝字票生命周期，不等于这张凭证本身就是红字票。"""
    assert service._status(
        "已红冲",
        Decimal("100.00"),
        None,
        "该蓝字发票已红冲",
    ) == "issued"
    assert service._status(
        "正常",
        Decimal("-100.00"),
        None,
        "被红冲蓝字数电发票号码：123456",
    ) == "red"


def test_purchase_reconciliation_uses_remaining_amount_after_partial_red(db_session):
    """采购对账只配蓝字票红冲后的剩余有效金额，不能继续按原票金额重复计入。"""
    from datetime import datetime, timezone

    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice
    from app.services import invoice_reconciliation

    token = uuid4().hex[:10]
    seller = f"净额供应商-{token}"
    blue_no = f"BLUE-PO-{token}"
    po = ExternalPurchaseOrder(
        external_order_id=f"PO-NET-{token}",
        platform="other",
        supplier_name=seller,
        ordered_at=datetime(2026, 7, 1, tzinfo=timezone.utc),
        order_amount=Decimal("700.00"),
        paid_amount=Decimal("700.00"),
        currency="CNY",
        raw={},
    )
    blue = TaxInvoice(
        invoice_key=f"pytest-blue-po-{token}",
        invoice_number=blue_no,
        direction="input",
        status="issued",
        issue_date=datetime(2026, 7, 5, tzinfo=timezone.utc),
        seller_name=seller,
        total_amount=Decimal("1000.00"),
        raw={"是否正数发票": "是"},
    )
    red = TaxInvoice(
        invoice_key=f"pytest-red-po-{token}",
        invoice_number=f"RED-PO-{token}",
        direction="input",
        status="red",
        issue_date=datetime(2026, 7, 6, tzinfo=timezone.utc),
        seller_name=seller,
        total_amount=Decimal("-300.00"),
        raw={"是否正数发票": "否", "备注": f"被红冲蓝字发票号码：{blue_no}"},
    )
    db_session.add_all([po, blue, red])
    db_session.commit()

    report = invoice_reconciliation.reconcile(db_session, supplier=seller)
    assert len(report["suppliers"]) == 1
    invoice_row = report["suppliers"][0]["months"][0]["invoices"][0]
    assert invoice_row["invoiceId"] == blue.id
    assert Decimal(str(invoice_row["originalAmount"])) == Decimal("1000.0")
    assert Decimal(str(invoice_row["redOffsetAmount"])) == Decimal("300.0")
    assert Decimal(str(invoice_row["amount"])) == Decimal("700.0")
    assert invoice_row["status"] == "matched"


def test_manual_red_blue_relation_overrides_missing_official_reference_and_is_audited(db_session):
    from datetime import datetime, timezone
    from app.models.tax import TaxInvoice

    token = uuid4().hex[:10]
    blue = TaxInvoice(
        invoice_key=f"manual-blue-{token}", invoice_number=f"B-{token}",
        direction="input", status="issued",
        issue_date=datetime(2026, 7, 1, tzinfo=timezone.utc),
        seller_name=f"人工红冲供应商-{token}", seller_tax_id=f"SELL-{token}",
        buyer_tax_id=f"BUY-{token}", total_amount=Decimal("500.00"),
        raw={"是否正数发票": "是"},
    )
    red = TaxInvoice(
        invoice_key=f"manual-red-{token}", invoice_number=f"R-{token}",
        direction="input", status="red",
        issue_date=datetime(2026, 8, 2, tzinfo=timezone.utc),
        seller_name=blue.seller_name, seller_tax_id=blue.seller_tax_id,
        buyer_tax_id=blue.buyer_tax_id, total_amount=Decimal("-120.00"),
        raw={"是否正数发票": "否", "备注": "原始导出未提供可识别蓝票号码"},
    )
    db_session.add_all([blue, red])
    db_session.commit()

    before = service.red_accounting_context(db_session, [red])[red.id]
    assert before["redPairStatus"] == "unpaired"

    service.set_manual_red_blue_relation(db_session, red.id, blue.id, actor="pytest")
    paired = service.red_accounting_context(db_session, [red])[red.id]
    assert paired["redPairStatus"] == "paired_manual"
    assert paired["redPairMethod"] == "manual"
    assert paired["redRelatedInvoiceId"] == blue.id
    assert paired["redCrossPeriod"] is True
    assert paired["redRelatedInvoicePeriod"] == "2026-07"

    service.clear_manual_red_blue_relation(db_session, red.id, actor="pytest")
    cleared = service.red_accounting_context(db_session, [red])[red.id]
    assert cleared["redPairStatus"] == "unpaired"


def test_red_settlement_bank_refund_cannot_be_overallocated(db_session):
    from datetime import date, datetime, timezone
    from app.models.bank import BankTransaction
    from app.models.tax import TaxInvoice

    token = uuid4().hex[:10]
    red = TaxInvoice(
        invoice_key=f"settle-red-{token}", invoice_number=f"R-{token}",
        direction="input", status="red",
        issue_date=datetime(2026, 8, 2, tzinfo=timezone.utc),
        seller_name=f"退款供应商-{token}", total_amount=Decimal("-300.00"),
        raw={"是否正数发票": "否"},
    )
    txn = BankTransaction(
        txn_date=date(2026, 8, 8), direction="in", amount=Decimal("300.00"),
        counterparty_name=red.seller_name, voucher_no=f"REF-{token}",
        fingerprint=f"pytest-refund-{token}", raw={},
    )
    db_session.add_all([red, txn])
    db_session.commit()

    result = service.add_red_settlement(
        db_session, red.id, service.RED_BANK_REFUND_TARGET_TYPE, Decimal("200.00"),
        target_id=txn.id, actor="pytest",
    )
    assert result["redSettlementStatus"] == "partial"
    assert Decimal(result["redSettlementRemainingAmount"]) == Decimal("100.00")

    with pytest.raises(ValueError, match="已经登记|可用金额不足"):
        service.add_red_settlement(
            db_session, red.id, service.RED_BANK_REFUND_TARGET_TYPE, Decimal("150.00"),
            target_id=txn.id, actor="pytest",
        )

    service.remove_red_settlement(db_session, red.id, result["redSettlements"][0]["linkId"], actor="pytest")
    after = service._red_settlement_context(db_session, [red])[red.id]
    assert after["redSettlementStatus"] == "unsettled"
    assert Decimal(after["redSettlementRemainingAmount"]) == Decimal("300.00")


def test_future_invoice_offset_respects_supplier_date_and_available_amount(db_session):
    from datetime import datetime, timezone
    from app.models.tax import TaxInvoice

    token = uuid4().hex[:10]
    seller_tax = f"SELL-{token}"
    red = TaxInvoice(
        invoice_key=f"future-red-{token}", invoice_number=f"R-{token}",
        direction="input", status="red", issue_date=datetime(2026, 8, 2, tzinfo=timezone.utc),
        seller_name=f"冲抵供应商-{token}", seller_tax_id=seller_tax,
        total_amount=Decimal("-300.00"), raw={"是否正数发票": "否"},
    )
    future = TaxInvoice(
        invoice_key=f"future-blue-{token}", invoice_number=f"B-{token}",
        direction="input", status="issued", issue_date=datetime(2026, 9, 2, tzinfo=timezone.utc),
        seller_name=red.seller_name, seller_tax_id=seller_tax,
        total_amount=Decimal("250.00"), raw={"是否正数发票": "是"},
    )
    db_session.add_all([red, future])
    db_session.commit()

    with pytest.raises(ValueError, match="可抵扣余额不足"):
        service.add_red_settlement(
            db_session, red.id, service.RED_FUTURE_OFFSET_TARGET_TYPE, Decimal("300.00"),
            target_id=future.id, actor="pytest",
        )

    result = service.add_red_settlement(
        db_session, red.id, service.RED_FUTURE_OFFSET_TARGET_TYPE, Decimal("250.00"),
        target_id=future.id, actor="pytest",
    )
    assert Decimal(result["redSettlementRemainingAmount"]) == Decimal("50.00")


def test_vat_context_marks_meal_nondeductible_and_verified_red_transfer_pending(db_session):
    from datetime import datetime, timezone
    from app.models.tax import TaxInvoice

    token = uuid4().hex[:10]
    meal = TaxInvoice(
        invoice_key=f"vat-meal-{token}", invoice_number=f"M-{token}",
        direction="input", status="issued", issue_date=datetime(2026, 8, 1, tzinfo=timezone.utc),
        total_amount=Decimal("106.00"), tax_amount=Decimal("6.00"),
        raw={"是否正数发票": "是", "货物或应税劳务、服务名称": "*餐饮服务*餐费"},
    )
    blue = TaxInvoice(
        invoice_key=f"vat-blue-{token}", invoice_number=f"VB-{token}",
        direction="input", status="issued", issue_date=datetime(2026, 7, 1, tzinfo=timezone.utc),
        seller_tax_id=f"S-{token}", buyer_tax_id=f"B-{token}",
        total_amount=Decimal("1130.00"), tax_amount=Decimal("130.00"),
        verified=True, raw={"是否正数发票": "是"},
    )
    red = TaxInvoice(
        invoice_key=f"vat-red-{token}", invoice_number=f"VR-{token}",
        direction="input", status="red", issue_date=datetime(2026, 8, 5, tzinfo=timezone.utc),
        seller_tax_id=blue.seller_tax_id, buyer_tax_id=blue.buyer_tax_id,
        total_amount=Decimal("-226.00"), tax_amount=Decimal("-26.00"),
        raw={"是否正数发票": "否", "备注": f"被红冲蓝字发票号码：{blue.invoice_number}"},
    )
    db_session.add_all([meal, blue, red])
    db_session.commit()

    meal_payload = service.serialize_invoice(meal, db=db_session)
    assert meal_payload["vatDeductibleStatus"] == "non_deductible"
    assert Decimal(meal_payload["vatDeductibleAmount"]) == Decimal("0.00")

    red_payload = service.serialize_invoice(red, db=db_session)
    assert red_payload["inputVatTransferStatus"] == "required_confirmation"
    assert Decimal(red_payload["inputVatTransferAmount"]) == Decimal("26.00")

    confirmed = service.set_vat_review(
        db_session,
        red.id,
        input_vat_transfer_status="completed",
        input_vat_transfer_amount=Decimal("26.00"),
        actor="pytest",
    )
    assert confirmed["inputVatTransferStatus"] == "completed"
    assert confirmed["inputVatTransferManual"] is True
    red_payload = service.serialize_invoice(red, db=db_session)
    assert red_payload["inputVatTransferStatus"] == "completed"
    assert Decimal(red_payload["inputVatTransferAmount"]) == Decimal("26.00")

    with pytest.raises(ValueError, match="不可抵扣"):
        service.set_vat_review(
            db_session,
            meal.id,
            vat_deductible_status="deductible",
            actor="pytest",
        )
    reviewed_meal = service.set_vat_review(
        db_session,
        meal.id,
        vat_deductible_status="non_deductible",
        actor="pytest",
    )
    assert reviewed_meal["vatDeductibleStatus"] == "non_deductible"
    assert reviewed_meal["vatDeductibleManual"] is True


def test_red_after_historical_business_link_marks_overmatched_for_review(db_session):
    from datetime import datetime, timezone
    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice, TaxInvoiceLink

    token = uuid4().hex[:10]
    blue_no = f"BO-{token}"
    blue = TaxInvoice(
        invoice_key=f"bo-blue-{token}", invoice_number=blue_no,
        direction="input", status="issued", issue_date=datetime(2026, 7, 1, tzinfo=timezone.utc),
        seller_name=f"超配供应商-{token}", total_amount=Decimal("1000.00"),
        raw={"是否正数发票": "是"},
    )
    red = TaxInvoice(
        invoice_key=f"bo-red-{token}", invoice_number=f"BR-{token}",
        direction="input", status="red", issue_date=datetime(2026, 8, 1, tzinfo=timezone.utc),
        seller_name=blue.seller_name, total_amount=Decimal("-300.00"),
        raw={"是否正数发票": "否", "备注": f"被红冲蓝字发票号码：{blue_no}"},
    )
    po = ExternalPurchaseOrder(
        external_order_id=f"PO-{token}", platform="other",
        supplier_name=blue.seller_name, order_amount=Decimal("1000.00"), paid_amount=Decimal("1000.00"),
    )
    db_session.add_all([blue, red, po])
    db_session.flush()
    db_session.add(TaxInvoiceLink(
        invoice_id=blue.id, target_type="external_purchase_order", target_id=po.id,
        allocated_amount=Decimal("1000.00"), match_method="manual", confirmed=True,
    ))
    db_session.commit()

    payload = service.serialize_invoice(blue, db=db_session)
    assert payload["businessMatchStatus"] == "needs_review"
    assert Decimal(payload["businessOvermatchedAmount"]) == Decimal("300.00")
    assert "超额" in payload["businessMatchException"]


def test_orphan_business_link_is_needs_review_in_serialization_and_filter(db_session):
    from app.models.tax import TaxInvoice, TaxInvoiceLink

    invoice = TaxInvoice(
        invoice_key=f"pytest-orphan-link-{uuid4().hex}",
        invoice_number=f"ORPHAN-{uuid4().hex[:10]}",
        direction="input",
        status="issued",
        seller_name="孤立关联供应商",
        total_amount=Decimal("100.00"),
        match_status="matched",
        raw={},
    )
    db_session.add(invoice)
    db_session.flush()
    db_session.add(TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="external_purchase_order",
        target_id=999999999,
        allocated_amount=Decimal("100.00"),
        match_method="manual",
        confirmed=True,
    ))
    db_session.flush()

    payload = service.serialize_invoice(invoice, db=db_session)
    assert payload["businessMatchStatus"] == "needs_review"
    assert payload["businessMatchedAmount"] == "0.0000"
    assert payload["invalidLinkCount"] == 1

    review_ids = {
        row["id"]
        for row in service.list_invoices(
            db_session, direction="input", match_status="needs_review", limit=500
        )
    }
    matched_ids = {
        row["id"]
        for row in service.list_invoices(
            db_session, direction="input", match_status="matched", limit=500
        )
    }
    assert invoice.id in review_ids
    assert invoice.id not in matched_ids


def test_1688_alias_candidate_shares_invoice_occupancy(db_session):
    """1688 原始单已挂票时，工作流副本候选必须显示同一占用事实。"""
    from app.models.alibaba1688_import import Alibaba1688Order
    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice, TaxInvoiceLink

    token = uuid4().hex[:10]
    order_no = f"ALIAS-OCC-{token}"
    invoice = TaxInvoice(
        invoice_key=f"alias-occ-{token}",
        invoice_number=f"INV-ALIAS-OCC-{token}",
        direction="input",
        status="issued",
        total_amount=Decimal("100"),
    )
    raw_order = Alibaba1688Order(
        external_order_id=order_no,
        seller_company_name="Alias供应商",
        actual_payment=Decimal("100"),
        import_id=1,
    )
    workflow_order = ExternalPurchaseOrder(
        external_order_id=order_no,
        platform="1688",
        supplier_name="Alias供应商",
        paid_amount=Decimal("100"),
        order_amount=Decimal("100"),
    )
    db_session.add_all([invoice, raw_order, workflow_order])
    db_session.flush()
    db_session.add(TaxInvoiceLink(
        invoice_id=invoice.id,
        target_type="alibaba1688_order",
        target_id=raw_order.id,
        allocated_amount=Decimal("100"),
        match_method="manual",
        confirmed=True,
    ))
    db_session.commit()

    candidates = service.purchase_link_candidates(
        db_session, invoice.id, keyword=order_no, limit=20
    )
    workflow = next(
        row for row in candidates
        if row["targetType"] == "external_purchase_order"
        and row["targetId"] == workflow_order.id
    )
    assert workflow["linkedInvoiceId"] == invoice.id
    assert workflow["linkedInvoiceNo"] == invoice.invoice_number


def test_1688_alias_blocks_same_invoice_duplicate_and_shares_order_capacity(db_session):
    """同一 1688 订单的双实体不能重复挂同一票，且不同票也必须共享订单总额度。"""
    from app.models.alibaba1688_import import Alibaba1688Order
    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice, TaxInvoiceLink

    token = uuid4().hex[:10]
    order_no = f"ALIAS-CAP-{token}"
    invoice_a = TaxInvoice(
        invoice_key=f"alias-cap-a-{token}",
        invoice_number=f"INV-ALIAS-CAP-A-{token}",
        direction="input",
        status="issued",
        total_amount=Decimal("60"),
    )
    invoice_b = TaxInvoice(
        invoice_key=f"alias-cap-b-{token}",
        invoice_number=f"INV-ALIAS-CAP-B-{token}",
        direction="input",
        status="issued",
        total_amount=Decimal("50"),
    )
    raw_order = Alibaba1688Order(
        external_order_id=order_no,
        seller_company_name="Alias额度供应商",
        actual_payment=Decimal("100"),
        import_id=1,
    )
    workflow_order = ExternalPurchaseOrder(
        external_order_id=order_no,
        platform="1688",
        supplier_name="Alias额度供应商",
        paid_amount=Decimal("100"),
        order_amount=Decimal("100"),
    )
    db_session.add_all([invoice_a, invoice_b, raw_order, workflow_order])
    db_session.flush()
    db_session.add(TaxInvoiceLink(
        invoice_id=invoice_a.id,
        target_type="alibaba1688_order",
        target_id=raw_order.id,
        allocated_amount=Decimal("60"),
        match_method="manual",
        confirmed=True,
    ))
    db_session.commit()

    with pytest.raises(ValueError, match="另一数据来源"):
        service.link_purchase_order(
            db_session,
            invoice_id=invoice_a.id,
            target_type="external_purchase_order",
            target_id=workflow_order.id,
            allocated_amount=Decimal("40"),
            actor="pytest",
        )

    with pytest.raises(ValueError, match="超过订单金额"):
        service.link_purchase_order(
            db_session,
            invoice_id=invoice_b.id,
            target_type="external_purchase_order",
            target_id=workflow_order.id,
            allocated_amount=Decimal("50"),
            actor="pytest",
        )


def test_duplicate_1688_alias_history_is_needs_review_in_cache_and_live_view(db_session):
    """历史双实体重复链接即使总额未超票面，也必须暴露为 needs_review。"""
    from app.models.alibaba1688_import import Alibaba1688Order
    from app.models.purchase import ExternalPurchaseOrder
    from app.models.tax import TaxInvoice, TaxInvoiceLink

    token = uuid4().hex[:10]
    order_no = f"ALIAS-DUP-{token}"
    invoice = TaxInvoice(
        invoice_key=f"alias-dup-{token}",
        invoice_number=f"INV-ALIAS-DUP-{token}",
        direction="input",
        status="issued",
        total_amount=Decimal("100"),
        match_status="partial",
    )
    raw_order = Alibaba1688Order(
        external_order_id=order_no,
        seller_company_name="Alias重复供应商",
        actual_payment=Decimal("100"),
        import_id=1,
    )
    workflow_order = ExternalPurchaseOrder(
        external_order_id=order_no,
        platform="1688",
        supplier_name="Alias重复供应商",
        paid_amount=Decimal("100"),
        order_amount=Decimal("100"),
    )
    db_session.add_all([invoice, raw_order, workflow_order])
    db_session.flush()
    db_session.add_all([
        TaxInvoiceLink(
            invoice_id=invoice.id,
            target_type="alibaba1688_order",
            target_id=raw_order.id,
            allocated_amount=Decimal("40"),
            match_method="manual",
            confirmed=True,
        ),
        TaxInvoiceLink(
            invoice_id=invoice.id,
            target_type="external_purchase_order",
            target_id=workflow_order.id,
            allocated_amount=Decimal("40"),
            match_method="manual",
            confirmed=True,
        ),
    ])
    db_session.flush()

    assert service.sync_business_match_status(
        db_session, invoice, "external_purchase_order"
    ) == "needs_review"
    payload = service.serialize_invoice(invoice, db=db_session)
    assert payload["businessMatchStatus"] == "needs_review"
    assert payload["duplicateLogicalLinkCount"] == 1
    assert "重复业务链接" in payload["businessMatchException"]
