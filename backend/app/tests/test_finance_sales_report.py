from datetime import datetime
from decimal import Decimal
from io import BytesIO
from uuid import uuid4
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.config import settings
from app.models.catalog import ProductSku
from app.models.jackyun import JackyunGoodsDocument, JackyunGoodsDocumentItem
from app.models.sales import SalesOrder, SalesOrderItem
from app.services import finance_sales_report_service as svc


def _order_no(prefix: str) -> str:
    return f"{prefix}-{uuid4().hex[:12]}"


def test_default_fields_are_configurable_and_stable():
    fields = svc.default_fields()
    keys = [row["key"] for row in fields]
    enabled = {row["key"] for row in fields if row["enabled"]}
    assert len(keys) == len(set(keys))
    assert enabled == set(svc.DEFAULT_FIELD_KEYS)
    assert {"period", "warehouse", "tax_code", "total_quantity", "total_sales", "total_cost"} <= set(keys)


def test_monthly_sales_report_summarizes_by_warehouse(db_session):
    tz = ZoneInfo(settings.TZ)
    order_no = _order_no("FIN")
    ignored_no = _order_no("UNPAID")

    paid = SalesOrder(
        order_no=order_no, platform="淘宝", order_status="已完成", pay_status="已支付",
        order_amount=Decimal("110"), paid_amount=Decimal("100"),
        ordered_at=datetime(2026, 8, 15, 12, 0, tzinfo=tz),
        paid_at=datetime(2026, 8, 15, 12, 1, tzinfo=tz),
        raw={"warehouseName": "常州-示范仓"},
    )
    other_wh = SalesOrder(
        order_no=_order_no("FIN2"), platform="拼多多", order_status="paid",
        paid_amount=Decimal("150"),
        ordered_at=datetime(2026, 8, 16, 12, 0, tzinfo=tz),
        raw={"warehouseName": "杭州仓"},
    )
    ignored = SalesOrder(
        order_no=ignored_no, platform="拼多多", order_status="待付款", pay_status="unpaid",
        ordered_at=datetime(2026, 8, 16, 12, 0, tzinfo=tz),
        raw={"warehouseName": "常州-示范仓"},
    )
    db_session.add_all([paid, other_wh, ignored])
    db_session.flush()
    db_session.add_all([
        SalesOrderItem(order_id=paid.id, sku_code="SKU-A", goods_name="商品A",
                       quantity=Decimal("2"), unit_price=Decimal("30"), amount=Decimal("60"), discount_amount=Decimal("10")),
        SalesOrderItem(order_id=paid.id, sku_code="SKU-B", goods_name="商品B",
                       quantity=Decimal("1"), unit_price=Decimal("50"), amount=Decimal("50"), discount_amount=Decimal("0")),
        SalesOrderItem(order_id=other_wh.id, sku_code="SKU-B", goods_name="商品B",
                       quantity=Decimal("3"), unit_price=Decimal("50"), amount=Decimal("150"), discount_amount=Decimal("0")),
        SalesOrderItem(order_id=ignored.id, sku_code="SKU-X", goods_name="未支付商品",
                       quantity=Decimal("99"), amount=Decimal("999"), discount_amount=Decimal("0")),
        ProductSku(jackyun_sku_id=f"sku-{uuid4().hex}", sku_code="SKU-A", default_cost=Decimal("12"),
                   tax_code="3040205000000000000"),
        ProductSku(jackyun_sku_id=f"sku-{uuid4().hex}", sku_code="SKU-B", default_cost=Decimal("8"),
                   tax_code="3040205000000000000"),
        ProductSku(jackyun_sku_id=f"sku-{uuid4().hex}", sku_code="SKU-X"),
    ])
    db_session.flush()
    # 成本只取采购入库加权成本（不再用 default_cost 兜底），补入库单让成本口径生效
    sku_a = db_session.query(ProductSku).filter_by(sku_code="SKU-A").one()
    sku_b = db_session.query(ProductSku).filter_by(sku_code="SKU-B").one()
    inbound = JackyunGoodsDocument(
        document_type="inbound", goodsdoc_no=f"RK-{uuid4().hex}",
        document_at=datetime(2026, 7, 1, 10, 0, tzinfo=tz),
    )
    db_session.add(inbound)
    db_session.flush()
    db_session.add_all([
        JackyunGoodsDocumentItem(
            document_id=inbound.id, line_no=1, goods_no="SKU-A", quantity=Decimal("100"),
            unit_price_tax=Decimal("12"), matched_sku_id=sku_a.id,
        ),
        JackyunGoodsDocumentItem(
            document_id=inbound.id, line_no=2, goods_no="SKU-B", quantity=Decimal("100"),
            unit_price_tax=Decimal("8"), matched_sku_id=sku_b.id,
        ),
    ])
    db_session.flush()

    company = f"pytest-{uuid4().hex}"
    template = svc.get_or_create_template(db_session, company)
    report = svc.build_report(db_session, 2026, 8, template)

    assert report["summary"]["orderCount"] == 2
    assert report["summary"]["warehouseCount"] == 2
    assert report["summary"]["totalQuantity"] == "6.0000"
    assert report["summary"]["salesAmount"] == "250.00"
    assert report["summary"]["costAmount"] == "56.00"  # 2*12 + 1*8 + 3*8（全部来自入库加权成本）
    assert report["summary"]["costIncomplete"] is False
    assert report["summary"]["costMissingDetail"] == []

    by_wh = {row["warehouse"]: row for row in report["rows"]}
    assert set(by_wh) == {"常州-示范仓", "杭州仓"}
    cz = by_wh["常州-示范仓"]
    assert cz["period"] == "2026-08"
    assert cz["tax_code"] == "3040205000000000000"
    assert cz["total_quantity"] == "3.0000"
    assert cz["total_sales"] == "100.00"
    assert cz["total_cost"] == "32.00"
    hz = by_wh["杭州仓"]
    assert hz["total_sales"] == "150.00"
    assert hz["total_cost"] == "24.00"

    workbook = load_workbook(BytesIO(svc.to_xlsx(report)), data_only=True)
    assert workbook.sheetnames == ["销售汇总"]
    sheet = workbook["销售汇总"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == ["月度时间", "仓库", "税务编号", "发货总数量", "销售总金额", "销售总成本"]
    # 最后一行是合计
    assert sheet.cell(row=sheet.max_row, column=2).value == "合计"
    assert sheet.cell(row=sheet.max_row, column=5).value == 250.0


def test_monthly_report_prefers_weighted_inbound_cost_before_default_cost(db_session):
    tz = ZoneInfo(settings.TZ)
    sku = ProductSku(
        jackyun_sku_id=f"sku-{uuid4().hex}", sku_code=f"INBOUND-COST-{uuid4().hex[:8]}",
        sku_name="入库成本商品", default_cost=Decimal("99"),
        tax_code="3040205000000000000",
    )
    db_session.add(sku)
    db_session.flush()
    old_doc = JackyunGoodsDocument(
        document_type="inbound", goodsdoc_no=f"RK-{uuid4().hex}",
        document_at=datetime(2026, 7, 20, 10, 0, tzinfo=tz),
    )
    current_doc = JackyunGoodsDocument(
        document_type="inbound", goodsdoc_no=f"RK-{uuid4().hex}",
        document_at=datetime(2026, 8, 10, 10, 0, tzinfo=tz),
    )
    future_doc = JackyunGoodsDocument(
        document_type="inbound", goodsdoc_no=f"RK-{uuid4().hex}",
        document_at=datetime(2026, 9, 1, 10, 0, tzinfo=tz),
    )
    db_session.add_all([old_doc, current_doc, future_doc])
    db_session.flush()
    db_session.add_all([
        JackyunGoodsDocumentItem(
            document_id=old_doc.id, line_no=1, goods_no=sku.sku_code,
            sku_barcode=sku.sku_code, goods_name=sku.sku_name,
            quantity=Decimal("2"), unit_price_tax=Decimal("10"), matched_sku_id=sku.id,
        ),
        JackyunGoodsDocumentItem(
            document_id=current_doc.id, line_no=1, goods_no=sku.sku_code,
            sku_barcode=sku.sku_code, goods_name=sku.sku_name,
            quantity=Decimal("3"), unit_price_tax=Decimal("20"), matched_sku_id=sku.id,
        ),
        JackyunGoodsDocumentItem(
            document_id=future_doc.id, line_no=1, goods_no=sku.sku_code,
            sku_barcode=sku.sku_code, goods_name=sku.sku_name,
            quantity=Decimal("100"), unit_price_tax=Decimal("100"), matched_sku_id=sku.id,
        ),
    ])
    order = SalesOrder(
        order_no=_order_no("WEIGHTED"), platform="淘宝", order_status="已完成", pay_status="已支付",
        ordered_at=datetime(2026, 8, 20, 12, 0, tzinfo=tz), raw={"warehouseName": "常州-示范仓"},
    )
    db_session.add(order)
    db_session.flush()
    db_session.add(SalesOrderItem(
        order_id=order.id, sku_code=sku.sku_code, goods_name=sku.sku_name,
        quantity=Decimal("1"), amount=Decimal("100"), discount_amount=Decimal("0"),
        sku_id=sku.id,
    ))
    db_session.flush()

    company = f"pytest-{uuid4().hex}"
    template = svc.get_or_create_template(db_session, company)
    report = svc.build_report(db_session, 2026, 8, template)
    assert report["summary"]["costAmount"] == "16.00"  # (2*10 + 3*20) / 5

    unbilled = svc.build_unbilled_income_report(db_session, 2026, 8, company=company)
    assert unbilled["details"][0]["cost"] == "16.00"


def test_multiple_tax_codes_joined_and_missing_cost_is_zero(db_session):
    tz = ZoneInfo(settings.TZ)
    order = SalesOrder(
        order_no=_order_no("FIN"), platform="淘宝", order_status="已完成", pay_status="已支付",
        paid_amount=Decimal("70"),
        ordered_at=datetime(2026, 8, 15, 12, 0, tzinfo=tz),
        raw={"warehouseName": "常州-示范仓"},
    )
    db_session.add(order)
    db_session.flush()
    db_session.add_all([
        SalesOrderItem(order_id=order.id, sku_code="SKU-A", goods_name="商品A",
                       quantity=Decimal("1"), amount=Decimal("30"), discount_amount=Decimal("0")),
        SalesOrderItem(order_id=order.id, sku_code="SKU-B", goods_name="商品B",
                       quantity=Decimal("1"), amount=Decimal("40"), discount_amount=Decimal("0")),
        ProductSku(jackyun_sku_id=f"sku-{uuid4().hex}", sku_code="SKU-A", default_cost=None,
                   tax_code="3040205000000000000"),
        ProductSku(jackyun_sku_id=f"sku-{uuid4().hex}", sku_code="SKU-B", default_cost=Decimal("5"),
                   tax_code="1090101020000000000"),
    ])
    db_session.flush()

    company = f"pytest-{uuid4().hex}"
    template = svc.get_or_create_template(db_session, company)
    template = svc.save_template(
        db_session, company=company, fields=svc.default_fields(), rules={},
        to_addrs=[], cc_addrs=[], auto_send=False, send_day=3, send_hour=10,
    )
    report = svc.build_report(db_session, 2026, 8, template)
    assert len(report["rows"]) == 1
    row = report["rows"][0]
    assert row["tax_code"] == "1090101020000000000 / 3040205000000000000"
    assert row["total_quantity"] == "2.0000"
    assert row["total_sales"] == "70.00"
    # 两个 SKU 都没有采购入库成本：不用 default_cost 兜底，成本为 0 并标记待补充
    assert row["total_cost"] == "0.00"
    assert report["summary"]["costIncomplete"] is True
    assert [item["skuCode"] for item in report["summary"]["costMissingDetail"]] == ["SKU-A", "SKU-B"]


def test_legacy_template_auto_upgrades_to_warehouse_fields(db_session):
    company = f"pytest-{uuid4().hex}"
    template = svc.get_or_create_template(db_session, company)
    # 模拟旧版订单明细口径模板：字段不在当前注册表时自动迁移为仓库汇总默认字段
    template.fields = [
        {"key": "platform", "label": "平台", "enabled": True},
        {"key": "order_no", "label": "财务订单号", "enabled": True},
        {"key": "refund_amount", "label": "退款", "enabled": True},
    ]
    db_session.commit()
    template = svc.get_or_create_template(db_session, company)
    enabled = {row["key"] for row in template.fields if row["enabled"]}
    assert enabled == set(svc.DEFAULT_FIELD_KEYS)


def test_unbilled_adjustment_persists_selected_details_and_version(db_session):
    tz = ZoneInfo(settings.TZ)
    order = SalesOrder(
        order_no=_order_no("UNBILLED"), platform="淘宝", order_status="已完成", pay_status="已支付",
        paid_amount=Decimal("50"),
        ordered_at=datetime(2098, 8, 15, 12, 0, tzinfo=tz),
        raw={"warehouseName": "常州-示范仓"},
    )
    db_session.add(order)
    db_session.flush()
    sku_a = f"UNBILLED-A-{uuid4().hex[:8]}"
    sku_b = f"UNBILLED-B-{uuid4().hex[:8]}"
    db_session.add_all([
        SalesOrderItem(order_id=order.id, sku_code=sku_a, goods_name="商品A", quantity=Decimal("2"),
                       amount=Decimal("20"), discount_amount=Decimal("0")),
        SalesOrderItem(order_id=order.id, sku_code=sku_b, goods_name="商品B", quantity=Decimal("3"),
                       amount=Decimal("30"), discount_amount=Decimal("0")),
        ProductSku(jackyun_sku_id=f"sku-{uuid4().hex}", sku_code=sku_a, sku_name="商品A",
                   default_cost=Decimal("4"), tax_code="3010101000000000000"),
        ProductSku(jackyun_sku_id=f"sku-{uuid4().hex}", sku_code=sku_b, sku_name="商品B",
                   default_cost=Decimal("5"), tax_code="3010102000000000000"),
    ])
    db_session.flush()

    company = f"pytest-{uuid4().hex}"
    base = svc.build_unbilled_income_report(db_session, 2098, 8, company=company)
    assert base["sourceCount"] == 2
    selected_key = svc.unbilled_detail_key(base["details"][0])

    saved = svc.save_unbilled_adjustment(
        db_session, company=company, year=2098, month=8,
        selected_keys=[selected_key], actor="pytest-adjuster",
    )
    current = svc.build_unbilled_income_report(db_session, 2098, 8, company=company)

    assert saved["version"] == 1
    assert current["adjusted"] is True
    assert current["version"] == saved["version"]
    assert current["sourceCount"] == 2
    assert current["selectedCount"] == 1
    assert current["selectedKeys"] == [selected_key]
    assert len(current["sourceDetails"]) == 2
    assert {svc.unbilled_detail_key(row) for row in current["sourceDetails"]} == {
        svc.unbilled_detail_key(row) for row in base["details"]
    }
    assert len(current["details"]) == 1
    assert svc.unbilled_detail_key(current["details"][0]) == selected_key


def test_output_red_discount_adjusts_sales_basis_and_does_not_create_unbilled_income(db_session):
    from datetime import timezone
    from app.models.tax import TaxInvoice, TaxInvoiceLink

    token = uuid4().hex[:10]
    sku_code = f"RED-SALE-{token}"
    sku = ProductSku(
        jackyun_sku_id=f"red-sale-sku-{token}",
        sku_code=sku_code,
        sku_name="红冲测试商品",
        tax_code="3040205000000000000",
    )
    order = SalesOrder(
        order_no=_order_no("RED-DISCOUNT"),
        platform="淘宝",
        order_status="已完成",
        pay_status="已支付",
        paid_amount=Decimal("1000.00"),
        ordered_at=datetime(2026, 8, 5, 12, 0, tzinfo=ZoneInfo(settings.TZ)),
        raw={"warehouseName": "测试仓"},
    )
    db_session.add_all([sku, order])
    db_session.flush()
    db_session.add(SalesOrderItem(
        order_id=order.id,
        sku_code=sku_code,
        goods_name="红冲测试商品",
        quantity=Decimal("1"),
        amount=Decimal("1000.00"),
        discount_amount=Decimal("0"),
        sku_id=sku.id,
    ))
    blue = TaxInvoice(
        invoice_key=f"out-blue-{token}",
        invoice_number=f"OUT-B-{token}",
        direction="output",
        status="issued",
        issue_date=datetime(2026, 8, 10, tzinfo=timezone.utc),
        seller_name="本公司",
        buyer_name=f"客户-{token}",
        total_amount=Decimal("1000.00"),
        source_system="tax_export",
        raw={"是否正数发票": "是"},
    )
    red = TaxInvoice(
        invoice_key=f"out-red-{token}",
        invoice_number=f"OUT-R-{token}",
        direction="output",
        status="red",
        issue_date=datetime(2026, 8, 20, tzinfo=timezone.utc),
        seller_name=blue.seller_name,
        buyer_name=blue.buyer_name,
        total_amount=Decimal("-300.00"),
        source_system="tax_export",
        raw={"是否正数发票": "否", "备注": f"被红冲蓝字发票号码：{blue.invoice_number}"},
    )
    db_session.add_all([blue, red])
    db_session.flush()
    db_session.add(TaxInvoiceLink(
        invoice_id=blue.id,
        target_type="sales_order",
        target_id=order.id,
        allocated_amount=Decimal("1000.00"),
        match_method="manual",
        confirmed=True,
    ))
    db_session.commit()

    company = f"red-discount-company-{token}"
    report = svc.build_unbilled_income_report(db_session, 2026, 8, company=company)

    assert Decimal(report["salesAmount"]) == Decimal("1000.00")
    assert Decimal(report["redSalesAdjustmentAmount"]) == Decimal("-300.00")
    assert Decimal(report["adjustedSalesAmount"]) == Decimal("700.00")
    assert Decimal(report["invoicedAmount"]) == Decimal("700.00")
    assert Decimal(report["unbilledAmount"]) == Decimal("0.00")
    detail = next(row for row in report["details"] if row["product"] == "红冲测试商品")
    assert Decimal(detail["sales"]) == Decimal("1000.00")
    assert Decimal(detail["redSalesAdjustment"]) == Decimal("-300.00")
    assert Decimal(detail["adjustedSales"]) == Decimal("700.00")
    assert Decimal(detail["invoiced"]) == Decimal("700.00")
    assert Decimal(detail["unbilled"]) == Decimal("0.00")


def test_unbilled_income_xlsx_matches_delivery_tabs():
    report = {
        "year": 2026,
        "month": 8,
        "period": "2026-08",
        "salesAmount": "100.00",
        "redSalesAdjustmentAmount": "-5.00",
        "adjustedSalesAmount": "95.00",
        "invoicedAmount": "5.00",
        "unbilledAmount": "90.00",
        "details": [
            {
                "period": "2026-08", "taxCode": "A", "taxName": "食品",
                "product": "商品A", "quantity": "2", "sales": "60.00",
                "redSalesAdjustment": "-5.00", "adjustedSales": "55.00",
                "invoiced": "5.00", "unbilled": "50.00", "cost": "30.00",
            },
            {
                "period": "2026-08", "taxCode": "A", "taxName": "食品",
                "product": "商品B", "quantity": "1", "sales": "40.00",
                "redSalesAdjustment": "0.00", "adjustedSales": "40.00",
                "invoiced": "0.00", "unbilled": "40.00", "cost": "20.00",
            },
        ],
    }

    workbook = load_workbook(BytesIO(svc.unbilled_income_xlsx(report)), data_only=True)

    assert workbook.sheetnames == ["无票收入-汇总", "无票收入-明细"]
    summary = workbook["无票收入-汇总"]
    assert summary.cell(row=10, column=1).value == "月度时间"
    assert summary.cell(row=10, column=4).value == "发货数量"
    assert summary.cell(row=11, column=3).value == "食品"
    assert summary.cell(row=11, column=4).value == 3.0
    assert summary.cell(row=11, column=5).value == 100.0
    assert summary.cell(row=11, column=7).value == 95.0
    assert summary.cell(row=11, column=2).number_format == "@"
    assert summary.cell(row=12, column=3).value == "合计"

    detail = workbook["无票收入-明细"]
    assert detail.cell(row=10, column=4).value == "产品"
    assert detail.cell(row=10, column=8).value == "调整后销售金额"
    assert detail.cell(row=11, column=4).value == "商品A"
    assert detail.cell(row=12, column=4).value == "商品B"
    assert detail.cell(row=13, column=4).value == "合计"
    assert detail.cell(row=13, column=8).value == 95.0


def test_unlinked_output_service_invoice_is_excluded_from_unbilled_income(db_session):
    from app.models.tax import TaxInvoice

    tz = ZoneInfo(settings.TZ)
    token = uuid4().hex[:10]
    sku = ProductSku(
        jackyun_sku_id=f"sku-{token}", sku_code=f"SALES-{token}",
        sku_name="销售商品", tax_code="3010101000000000000",
    )
    order = SalesOrder(
        order_no=_order_no("SALES"), platform="淘宝", order_status="已完成",
        pay_status="已支付", paid_amount=Decimal("100.00"),
        ordered_at=datetime(2026, 8, 10, 12, 0, tzinfo=tz),
        raw={"warehouseName": "示范仓"},
    )
    db_session.add_all([sku, order])
    db_session.flush()
    db_session.add(SalesOrderItem(
        order_id=order.id, sku_code=sku.sku_code, goods_name=sku.sku_name,
        quantity=Decimal("1"), amount=Decimal("100.00"), discount_amount=Decimal("0"),
        sku_id=sku.id,
    ))
    db_session.add(TaxInvoice(
        invoice_key=f"service-fee-{token}", invoice_number=f"SERVICE-{token}",
        direction="output", status="issued", issue_date=datetime(2026, 8, 20, tzinfo=tz),
        seller_name="本公司", buyer_name="平台服务方", total_amount=Decimal("15.53"),
        source_system="tax_export", raw={"发票内容": "服务费"},
    ))
    db_session.commit()

    report = svc.build_unbilled_income_report(
        db_session, 2026, 8, company=f"service-fee-{token}",
    )

    assert report["salesAmount"] == "100.00"
    assert report["redSalesAdjustmentAmount"] == "0.00"
    assert report["invoicedAmount"] == "0.00"
    assert report["unbilledAmount"] == "100.00"
