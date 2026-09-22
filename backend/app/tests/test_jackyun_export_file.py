"""吉客云客户端官方导出文件解析测试，不使用真实业务数据。"""
import io
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

import pytest

from app.adapters.jackyun_export_file import parse_jackyun_export
from app.config import settings
from app.models.jackyun_import import JackyunFileImport, JackyunFileImportRecord
from app.models.jackyun import JackyunGoodsDocument, JackyunGoodsDocumentItem
from app.models.catalog import ProductSku, Warehouse
from app.models.sales import SalesOrder
from app.services import jackyun_file_import_service as service
from app.services.inbound_document_view import list_inbound_documents


def _xlsx(rows: list[list[object]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parse_xlsx_finds_header_after_export_title():
    content = _xlsx([
        ["吉客云库存明细表"],
        ["导出时间", "2026-09-02"],
        ["货品编号", "货品名称", "可用库存", "仓库名称"],
        ["SKU-001", "测试商品", 12, "主仓"],
    ])
    parsed = parse_jackyun_export(content, "库存明细.xlsx", max_rows=100)
    assert parsed.report_type == "inventory"
    assert parsed.headers == ["货品编号", "货品名称", "可用库存", "仓库名称"]
    assert parsed.rows == [{"货品编号": "SKU-001", "货品名称": "测试商品", "可用库存": "12", "仓库名称": "主仓"}]


def test_parse_csv_classifies_sales_and_keeps_original_headers():
    content = "销售订单明细\n订单编号,订单状态,店铺名称,下单时间\nSO-001,已付款,测试店,2026-09-02\n".encode("utf-8")
    parsed = parse_jackyun_export(content, "销售订单.csv", max_rows=100)
    assert parsed.report_type == "sales"
    assert parsed.rows[0]["订单编号"] == "SO-001"


def test_parse_inbound_apply_products_as_inbound():
    content = _xlsx([
        ["申请单号", "入库类型", "货品编号", "货品名称", "入库数量"],
        ["RK-IMPORT-001", "采购入库", "SKU-001", "测试商品", 3],
    ])
    parsed = parse_jackyun_export(content, "入库申请单货品.xlsx", max_rows=100)
    assert parsed.report_type == "inbound"


def test_fill_inbound_item_uses_purchase_total_and_derives_unit_price():
    item = JackyunGoodsDocumentItem(quantity=Decimal("197"), raw={})

    filled = service._fill_inbound_item(
        item,
        {
            "入库数量": "197",
            "采购总金额": "1940",
            "1688采购订单": "PO-DERIVE-001",
        },
    )

    assert "amount_tax" in filled
    assert item.amount_tax == Decimal("1940")
    assert item.unit_price_tax == Decimal("9.8477")
    assert item.raw["_1688采购订单"] == "PO-DERIVE-001"


def test_inbound_mapping_calculates_document_total_from_purchase_totals(db_session):
    import_row = JackyunFileImport(
        original_name="采购入库金额.xlsx", stored_path="/tmp/采购入库金额.xlsx",
        sha256=f"inbound-amount-{uuid4().hex}", report_type="inbound", lifecycle="active",
        headers=["申请单号", "货品编号", "货品名称", "入库数量", "采购总金额"],
        row_count=2, staged_row_count=2,
    )
    db_session.add(import_row)
    db_session.flush()
    db_session.add_all([
        JackyunFileImportRecord(
            import_id=import_row.id, row_index=1,
            payload={"申请单号": "RK-AMOUNT-001", "货品编号": "SKU-AMOUNT-001", "货品名称": "金额测试1", "入库数量": "100", "采购总金额": "1950"},
        ),
        JackyunFileImportRecord(
            import_id=import_row.id, row_index=2,
            payload={"申请单号": "RK-AMOUNT-001", "货品编号": "SKU-AMOUNT-002", "货品名称": "金额测试2", "入库数量": "50", "采购总金额": "850"},
        ),
    ])
    db_session.commit()

    document = None
    try:
        result = service.map_inbound_items(db_session, import_row.id, actor="pytest")
        document = db_session.query(JackyunGoodsDocument).filter_by(
            document_type="inbound", goodsdoc_no="RK-AMOUNT-001",
        ).one()
        items = db_session.query(JackyunGoodsDocumentItem).filter_by(document_id=document.id).order_by(JackyunGoodsDocumentItem.line_no).all()
        assert result["documents"] == 1
        assert [item.amount_tax for item in items] == [Decimal("1950"), Decimal("850")]
        assert [item.unit_price_tax for item in items] == [Decimal("19.5000"), Decimal("17.0000")]
        assert document.total_amount == Decimal("2800.0000")
    finally:
        if document:
            db_session.query(JackyunGoodsDocumentItem).filter_by(document_id=document.id).delete(synchronize_session=False)
            db_session.delete(document)
        db_session.query(JackyunFileImportRecord).filter_by(import_id=import_row.id).delete(synchronize_session=False)
        db_session.delete(import_row)
        db_session.commit()


def test_inbound_match_does_not_accept_missing_amount_as_price_ok(db_session):
    from app.services.sku_matching_service import match_inbound_items

    sku = ProductSku(
        jackyun_sku_id=f"J-AMOUNT-MISSING-{uuid4().hex}",
        sku_code=f"SKU-AMOUNT-MISSING-{uuid4().hex[:8]}",
        sku_name="缺少金额测试 SKU",
        status="active",
    )
    document = JackyunGoodsDocument(
        document_type="inbound", goodsdoc_no=f"RK-AMOUNT-MISSING-{uuid4().hex[:8]}",
    )
    db_session.add_all([sku, document])
    db_session.flush()
    item = JackyunGoodsDocumentItem(
        document_id=document.id, line_no=1, goods_no=sku.sku_code,
        goods_name=sku.sku_name, quantity=Decimal("1"), amount_tax=Decimal("0"), raw={},
    )
    db_session.add(item)
    db_session.commit()

    try:
        match_inbound_items(db_session)
        assert item.match_status == "auto"
        assert "金额信息不全" in item.match_note
    finally:
        db_session.delete(item)
        db_session.delete(document)
        db_session.delete(sku)
        db_session.commit()


def test_inbound_import_normalizes_warehouse_code_to_master_name(db_session):
    warehouse = Warehouse(code="WH-IMPORT-001", name="导入中文仓", status="active")
    import_row = JackyunFileImport(
        original_name="采购入库仓库编号.xlsx", stored_path="/tmp/采购入库仓库编号.xlsx",
        sha256=f"warehouse-normalize-{uuid4().hex}", report_type="inbound", lifecycle="active",
        headers=["申请单号", "仓库编号", "货品编号", "货品名称", "入库数量"],
        row_count=1, staged_row_count=1,
    )
    db_session.add_all([warehouse, import_row])
    db_session.flush()
    db_session.add(JackyunFileImportRecord(
        import_id=import_row.id,
        row_index=1,
        payload={
            "申请单号": "RK-WAREHOUSE-NORMALIZE-001",
            "仓库编号": "WH-IMPORT-001",
            "货品编号": "SKU-WAREHOUSE-001",
            "货品名称": "仓库匹配测试货品",
            "入库数量": "2",
        },
    ))
    db_session.commit()

    document = None
    try:
        result = service.map_inbound_items(db_session, import_row.id, actor="pytest")
        document = db_session.query(JackyunGoodsDocument).filter_by(
            document_type="inbound", goodsdoc_no="RK-WAREHOUSE-NORMALIZE-001",
        ).one()
        assert result["documents"] == 1
        assert document.warehouse_code == "WH-IMPORT-001"
        assert document.warehouse_name == "导入中文仓"
    finally:
        if document:
            db_session.query(JackyunGoodsDocumentItem).filter_by(document_id=document.id).delete(synchronize_session=False)
            db_session.delete(document)
        db_session.query(JackyunFileImportRecord).filter_by(import_id=import_row.id).delete(synchronize_session=False)
        db_session.delete(import_row)
        db_session.delete(warehouse)
        db_session.commit()


def test_parser_rejects_unsupported_extension():
    with pytest.raises(ValueError, match="XLSX 或 CSV"):
        parse_jackyun_export(b"not a pdf", "导出.pdf", max_rows=100)


def test_file_import_archives_and_deduplicates(db_session, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "DATA_DIR", str(tmp_path))
    content = _xlsx([
        ["订单编号", "订单状态", "店铺名称", "下单时间"],
        ["SO-ARCHIVE-001", "已付款", "测试店", "2026-09-02"],
    ])
    row, duplicate = service.import_export(
        db_session, content=content, original_name="销售订单.xlsx", actor="pytest-admin"
    )
    assert duplicate is False
    assert row.report_type == "sales"
    assert row.row_count == 1
    assert Path(row.stored_path).is_file()
    assert db_session.query(JackyunFileImportRecord).filter_by(import_id=row.id).count() == 1
    # 首次本地文件只做真实表头/原始行入 staging；未经样本映射不得伪造业务订单。
    assert db_session.query(SalesOrder).filter_by(order_no="SO-ARCHIVE-001").count() == 0

    same, duplicate = service.import_export(
        db_session, content=content, original_name="重命名销售订单.xlsx", actor="pytest-admin"
    )
    assert duplicate is True
    assert same.id == row.id


def test_upload_endpoint_returns_parsed_metadata(client, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "DATA_DIR", str(tmp_path))
    content = _xlsx([
        ["货品编号", "货品名称", "库存数量"],
        ["SKU-UPLOAD-001", "上传测试商品", 8],
    ])
    response = client.post(
        "/api/v1/jackyun-files/imports",
        files={"file": ("库存.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["duplicate"] is False
    assert body["import"]["reportType"] == "inventory"
    assert body["import"]["stagedRowCount"] == 1


def test_upload_endpoint_enforces_its_own_smaller_limit(client, monkeypatch):
    monkeypatch.setattr(settings, "MAX_JACKYUN_IMPORT_BYTES", 3)
    response = client.post(
        "/api/v1/jackyun-files/imports",
        files={"file": ("too-large.csv", b"1234", "text/csv")},
    )
    assert response.status_code == 413


def test_outbound_import_maps_document_lines_and_sku(db_session):
    sku = ProductSku(
        jackyun_sku_id=f"pytest-outbound-{uuid4().hex}",
        sku_code=f"PYTEST-OUTBOUND-{uuid4().hex[:8]}",
        sku_name="出库导入测试 SKU",
        barcode="BAR-OUTBOUND",
        status="active",
    )
    db_session.add(sku)
    db_session.flush()
    import_row = service.JackyunFileImport(
        original_name="销售出库.xlsx",
        stored_path="/tmp/pytest-outbound.xlsx",
        sha256=uuid4().hex,
        report_type="outbound",
        status="parsed",
        lifecycle="active",
        headers=["出库单号", "出库日期", "仓库名称", "货品编号", "货品名称", "出库数量"],
        row_count=1,
        staged_row_count=1,
    )
    db_session.add(import_row)
    db_session.flush()
    db_session.add(JackyunFileImportRecord(
        import_id=import_row.id,
        row_index=1,
        payload={
            "出库单号": "OUT-IMPORT-001",
            "出库日期": "2026-09-09",
            "仓库名称": "常州-示范仓",
            "货品编号": sku.sku_code,
            "货品名称": sku.sku_name,
            "出库数量": "7",
        },
    ))
    db_session.commit()

    result = service.map_import(db_session, import_row.id, actor="pytest")
    document = db_session.query(JackyunGoodsDocument).filter_by(
        document_type="outbound", goodsdoc_no="OUT-IMPORT-001"
    ).one()
    item = db_session.query(JackyunGoodsDocumentItem).filter_by(document_id=document.id).one()

    assert result["mapper"] == "outbound"
    assert result["documents"] == 1
    assert item.quantity == 7
    assert item.matched_sku_id == sku.id
    assert item.match_status in {"auto", "price_ok"}

    second = service.map_import(db_session, import_row.id, actor="pytest")
    assert second["documents"] == 1
    assert db_session.query(JackyunGoodsDocument).filter_by(
        document_type="outbound", goodsdoc_no="OUT-IMPORT-001"
    ).count() == 1


def test_inbound_import_wins_over_purchase_column(db_session):
    """采购入库报表同时带采购单号时，仍必须走入库明细映射。"""
    import_row = service.JackyunFileImport(
        original_name="采购入库单.xlsx", stored_path="/tmp/采购入库单.xlsx",
        sha256=f"inbound-dispatch-{uuid4().hex}", report_type="inbound", lifecycle="active",
        headers=["入库单号", "采购单号", "货品编号", "入库数量", "含税金额"],
        row_count=1, staged_row_count=1,
    )
    document = JackyunGoodsDocument(
        document_type="inbound", goodsdoc_no="RK-DISPATCH-001", supplier_name="测试供应商",
    )
    db_session.add_all([import_row, document])
    db_session.flush()
    item = JackyunGoodsDocumentItem(
        document_id=document.id, line_no=1, goods_no="SKU-DISPATCH-001", quantity=2,
    )
    db_session.add(item)
    db_session.add(JackyunFileImportRecord(
        import_id=import_row.id,
        row_index=1,
        payload={
            "入库单号": "RK-DISPATCH-001", "采购单号": "PO-DISPATCH-001",
            "货品编号": "SKU-DISPATCH-001", "入库数量": "2", "含税金额": "20",
        },
    ))
    db_session.commit()

    try:
        result = service.map_import(db_session, import_row.id, actor="pytest")
        assert result["mapper"] == "inbound_items"
        assert result["documents"] == 1
        assert result["matched"] == 1
    finally:
        db_session.query(JackyunFileImportRecord).filter_by(import_id=import_row.id).delete(synchronize_session=False)
        db_session.delete(item)
        db_session.delete(document)
        db_session.delete(import_row)
        db_session.commit()


def test_inbound_import_creates_missing_document_and_items(db_session):
    """只有入库申请单货品文件、库内没有主单时，也要落成可查看的入库单明细。"""
    from app.models.jackyun import JackyunGoodsDocument, JackyunGoodsDocumentItem
    from app.models.procurement_chain import ProcurementChainLink
    from app.models.purchase import ExternalPurchaseOrder

    rk_no = "RK-MISSING-DOCUMENT-001"
    order_no = "PO-MISSING-DOCUMENT-001"
    imp = JackyunFileImport(
        original_name="入库申请单货品.xlsx", stored_path="/tmp/missing-document.xlsx",
        sha256=f"missing-document-{uuid4().hex}", report_type="products", lifecycle="active",
        headers=["申请单号", "往来单位", "创建时间", "货品编号", "货品名称", "入库数量", "1688采购订单"],
        row_count=1, staged_row_count=1,
    )
    db_session.add(imp)
    db_session.flush()
    db_session.add(JackyunFileImportRecord(
        import_id=imp.id,
        row_index=1,
        payload={
            "申请单号": rk_no, "往来单位": "1688测试供应商", "创建时间": "2026-09-15 10:00:00",
            "货品编号": "SKU-MISSING-001", "货品名称": "导入测试商品", "入库数量": "3",
            "1688采购订单": order_no,
        },
    ))
    db_session.commit()

    try:
        result = service.map_inbound_items(db_session, imp.id, actor="pytest")
        document = db_session.query(JackyunGoodsDocument).filter_by(
            document_type="inbound", goodsdoc_no=rk_no,
        ).one()
        item = db_session.query(JackyunGoodsDocumentItem).filter_by(document_id=document.id).one()
        po = db_session.query(ExternalPurchaseOrder).filter_by(external_order_id=order_no).one()

        assert result["createdDocuments"] == 1
        assert result["createdDocumentItems"] == 1
        assert result["documents"] == 1
        assert result["matched"] == 1
        assert document.supplier_name == "1688测试供应商"
        assert item.goods_no == "SKU-MISSING-001"
        assert item.quantity == 3
        assert item.raw["1688采购订单"] == order_no
        assert db_session.query(ProcurementChainLink).filter_by(
            external_po_id=po.id, target_id=document.id, target_type="inbound",
        ).count() == 1
    finally:
        po = db_session.query(ExternalPurchaseOrder).filter_by(external_order_id=order_no).first()
        document = db_session.query(JackyunGoodsDocument).filter_by(goodsdoc_no=rk_no).first()
        if document:
            db_session.query(ProcurementChainLink).filter_by(target_id=document.id).delete(synchronize_session=False)
            db_session.query(JackyunGoodsDocumentItem).filter_by(document_id=document.id).delete(synchronize_session=False)
            db_session.delete(document)
        db_session.query(JackyunFileImportRecord).filter_by(import_id=imp.id).delete(synchronize_session=False)
        if po:
            db_session.delete(po)
        db_session.delete(imp)
        db_session.commit()


def test_inbound_document_view_returns_master_and_detail_rows(db_session):
    """到仓入库单列表按真实入库主单聚合，并把数量差异带到明细。"""
    document = JackyunGoodsDocument(
        document_type="inbound", goodsdoc_no="RK-VIEW-001", supplier_name="视图测试供应商",
        warehouse_name="视图测试仓", document_at=datetime(2026, 9, 15, 10, 0),
    )
    db_session.add(document)
    db_session.flush()
    db_session.add(JackyunGoodsDocumentItem(
        document_id=document.id, line_no=1, goods_no="SKU-VIEW-001", goods_name="视图测试货品",
        quantity=4, apply_quantity=5, match_status="matched",
    ))
    db_session.commit()

    try:
        payload = list_inbound_documents(db_session, q="RK-VIEW-001", limit=10)
        assert payload["total"] == 1
        assert payload["stats"]["pending"] == 1
        row = payload["rows"][0]
        assert row["inboundNo"] == "RK-VIEW-001"
        assert row["productCount"] == 1
        assert row["arrivedQuantity"] == 5
        assert row["actualQuantity"] == 4
        assert row["status"] == "部分入库"
        assert row["items"][0]["difference"] == 1
        assert list_inbound_documents(db_session, q="SKU-VIEW-001", limit=10)["total"] == 1
    finally:
        db_session.query(JackyunGoodsDocumentItem).filter_by(document_id=document.id).delete(synchronize_session=False)
        db_session.delete(document)
        db_session.commit()
