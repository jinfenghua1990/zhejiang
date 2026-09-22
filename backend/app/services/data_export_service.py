"""业务数据导出。

导出只读取本地已落库事实，不触发外部同步、不改变业务数据。
采购订单导出同时带出采购明细和已确认的吉客云入库关联，方便人工核验一条链路。
"""

from __future__ import annotations

import json
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.models.jackyun import JackyunGoodsDocument, JackyunGoodsDocumentItem
from app.models.purchase import ExternalPurchaseOrder, Supplier
from app.models.sales import SalesOrder, SalesOrderItem
from app.models.tax import TaxInvoice
from app.services import dashboard
from app.services import production_service
from app.services import procurement_workbench_service as procurement
from app.services import tax_category_rule_service
from app.services.procurement_chain_service import _order_row, chain_snapshot
from app.services.production_purchase_view import list_production_purchase_rows
from app.services.warehouse_service import list_warehouses


DATASET_OPTIONS: tuple[dict[str, str], ...] = (
    {"key": "purchase_orders", "label": "采购订单", "description": "订单主档、状态、仓库、入库单和发票状态"},
    {"key": "catalog", "label": "货品档案", "description": "正品与耗材统一档案、品类、税务代码和关联关系"},
    {"key": "inventory", "label": "库存快照", "description": "SKU 当前库存、仓库分布和快照时间"},
    {"key": "warehouses", "label": "仓库档案", "description": "仓库编码、用途、吉客云仓库 ID 和启用状态"},
    {"key": "bundles", "label": "套装档案", "description": "套装与虚拟组合套装 SKU 主档"},
    {"key": "tax_rules", "label": "财务分类", "description": "财务大类规则、匹配方式和税务代码"},
    {"key": "suppliers", "label": "供应商档案", "description": "供应商名称、税号、联系人和渠道信息"},
    {"key": "external_orders", "label": "其他渠道订单号库", "description": "淘宝、拼多多、线下等非 1688 采购订单主档"},
    {"key": "production_orders", "label": "生产订单（内部）", "description": "仅手工/内部生产单、货品明细和耗材需求；正品采购链路请导出生产采购链路"},
    {"key": "material_flow", "label": "生产采购链路", "description": "正品采购在生产、在途、到货、入库阶段的归档视图"},
    {"key": "inbound_documents", "label": "吉客云入库单", "description": "入库单主档和入库货品明细"},
    {"key": "sales_items", "label": "销售明细", "description": "销售单查询导入的订单货品行，包含订单状态"},
    {"key": "tax_invoices", "label": "发票台账", "description": "税务发票、采购关联和认证状态"},
)


def _cell(value: Any) -> Any:
    """转换成 Excel 可保存且便于人工阅读的值。"""
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (date,)):
        return value.isoformat()
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return value


def _workbook(sheets: list[tuple[str, list[str], list[dict[str, Any]]]], *, template: bool = False) -> bytes:
    """生成核验 Excel。template=True 时只保留表头行，作为可回导的标准空模板。"""
    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="EAF0FA")
    header_font = Font(bold=True, color="1E3A5F")
    for index, (name, headers, rows) in enumerate(sheets):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = name[:31]
        sheet.freeze_panes = "A2"
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
        for row in ([] if template else rows):
            sheet.append([_cell(row.get(header)) for header in headers])
        if headers:
            sheet.auto_filter.ref = sheet.dimensions
        for column_index, header in enumerate(headers, start=1):
            values = [str(header)]
            values.extend(str(sheet.cell(row=row_index, column=column_index).value or "") for row_index in range(2, min(sheet.max_row, 80) + 1))
            width = min(max(max(len(value) for value in values) + 2, 10), 42)
            sheet.column_dimensions[get_column_letter(column_index)].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _purchase_export(
    db: Session,
    *,
    status: str,
    q: str,
    channel: str,
    kind: str,
    warehouse: str,
    start_date: date | None,
    end_date: date | None,
    template: bool = False,
) -> tuple[bytes, int]:
    payload = procurement.list_orders(
        db,
        status=status,
        q=q,
        page=1,
        page_size=20_000,
        start_date=start_date,
        end_date=end_date,
        channel=channel,
        kind=kind,
        warehouse=warehouse,
    )
    orders = [item for group in payload.get("groups", []) for item in group.get("items", [])]
    selected_ids = {
        int(item["orderId"])
        for item in orders
        if item.get("orderId") is not None
    }
    # 列表服务已经完成了筛选和排序；详情数据只再做一次批量预加载，避免逐单调用
    # workbench（每单会重复扫描整套链路，48 单导出会拖到近一分钟）。
    detail_by_id: dict[int, dict[str, Any]] = {}
    if selected_ids:
        prefetch_pairs, prefetch = chain_snapshot(db)
        for source_order, external in prefetch_pairs:
            order_id = source_order.id if source_order is not None else -external.id
            if order_id in selected_ids:
                detail_by_id[order_id] = _order_row(db, source_order, external, pf=prefetch)
    order_rows: list[dict[str, Any]] = []
    detail_rows: list[dict[str, Any]] = []
    inbound_rows: list[dict[str, Any]] = []

    for row in orders:
        detail_payload = detail_by_id.get(int(row["orderId"])) or {}
        inbound = detail_payload.get("inbound") or []
        inbound_by_id = {
            item.get("documentId"): item
            for item in inbound
            if item.get("documentId") is not None
        }
        order_rows.append({
            "订单号": row.get("orderNo"),
            "渠道": row.get("platform"),
            "采购类型": "耗材" if row.get("orderKind") == "consumable" else "正品",
            "供应商/工厂": row.get("supplier"),
            "订单金额": row.get("amount"),
            "实付金额": row.get("paidAmount"),
            "下单日期": row.get("orderDate"),
            "采购状态": row.get("purchaseStatus"),
            "当前状态": row.get("firstUndoneLabel"),
            "仓库": row.get("warehouseName"),
            "物流单号": row.get("logisticsNo"),
            "吉客云入库单": row.get("jackyunInboundNo"),
            "发票状态": row.get("invoiceStatus"),
            "待开票金额": row.get("invoiceOutstanding"),
            "备注": row.get("remark"),
        })

        allocations = detail_payload.get("allocations") or []
        for allocation in allocations:
            inbound_item = inbound_by_id.get(allocation.get("inboundDocumentId")) or {}
            detail_rows.append({
                "订单号": row.get("orderNo"),
                "SKU": allocation.get("skuCode"),
                "商品": allocation.get("goodsName"),
                "数量": allocation.get("quantity"),
                "单价": allocation.get("unitPrice"),
                "金额": allocation.get("amount"),
                "来源入库单": inbound_item.get("goodsdocNo") or "手工补录/待关联",
                "入库仓库": (
                    allocation.get("warehouseName")
                    or inbound_item.get("warehouseName")
                    or (inbound[0].get("warehouseName") if inbound else "")
                    or (detail_payload.get("consumable") or {}).get("warehouseName")
                ),
                "明细备注": allocation.get("note"),
            })
        if not allocations:
            for raw_item in detail_payload.get("orderItems") or []:
                detail_rows.append({
                    "订单号": row.get("orderNo"),
                    "SKU": raw_item.get("productNumber"),
                    "商品": raw_item.get("productName"),
                    "数量": raw_item.get("quantity"),
                    "单价": raw_item.get("unitPrice"),
                    "金额": raw_item.get("amount"),
                    "来源入库单": "1688原始明细",
                    "入库仓库": "",
                    "明细备注": raw_item.get("spec"),
                })
        for inbound_item in inbound:
            usage = inbound_item.get("consumableUsageDecided")
            inbound_rows.append({
                "订单号": row.get("orderNo"),
                "入库单号": inbound_item.get("goodsdocNo"),
                "入库日期": inbound_item.get("date"),
                "仓库": inbound_item.get("warehouseName"),
                "入库金额": inbound_item.get("amount"),
                "入库明细金额": inbound_item.get("itemAmount"),
                "明细行数": inbound_item.get("itemCount"),
                "关联方式": inbound_item.get("matchMethod"),
                "匹配置信度": inbound_item.get("confidence"),
                "耗材使用": "已自动扣减" if usage and inbound_item.get("consumableUsageEnabled") else "待维护耗材映射" if not inbound_item.get("consumableUsageEnabled") else "待处理",
                "备注": inbound_item.get("note"),
            })

    content = _workbook([
        ("采购订单", list(order_rows[0].keys()) if order_rows else [
            "订单号", "渠道", "采购类型", "供应商/工厂", "订单金额", "实付金额", "下单日期", "采购状态",
            "当前状态", "仓库", "物流单号", "吉客云入库单", "发票状态", "待开票金额", "备注",
        ], order_rows),
        ("采购明细", ["订单号", "SKU", "商品", "数量", "单价", "金额", "来源入库单", "入库仓库", "明细备注"], detail_rows),
        ("入库关联", ["订单号", "入库单号", "入库日期", "仓库", "入库金额", "入库明细金额", "明细行数", "关联方式", "匹配置信度", "耗材使用", "备注"], inbound_rows),
    ], template=template)
    return content, len(orders)


def _catalog_export(
    db: Session,
    *,
    kind: str,
    q: str,
    template: bool = False,
    scope: str = "catalog",
) -> tuple[bytes, int]:
    rows = dashboard.catalog_unified(db, kind=kind if kind in {"all", "goods", "consumable"} else "all", search=q, limit=2000)
    if scope == "catalog":
        rows = [row for row in rows if row.get("kind") == "consumable" or row.get("category") == "single"]
    elif scope == "bundles":
        rows = [row for row in rows if row.get("kind") == "goods" and row.get("category") in {"bundle", "virtual_bundle"}]
    else:
        raise ValueError("不支持的货品档案范围")
    export_rows = [{
        "类型": "耗材" if row.get("kind") == "consumable" else "正品",
        "编码": row.get("code"),
        "名称": row.get("name"),
        "商品名称": row.get("goodsName"),
        "条码": row.get("barcode"),
        "单位": row.get("unit"),
        "货品类型": row.get("category") if row.get("kind") == "goods" else "",
        "品类": row.get("goodsCategory") or row.get("category"),
        "采购单价": row.get("purchaseUnitCost"),
        "售价": row.get("salePrice"),
        "默认成本": row.get("defaultCost"),
        "成本方式": row.get("costMode"),
        "成本容差": row.get("costTolerancePct"),
        "当前库存": row.get("stockOwn"),
        "工厂库存": row.get("stockFactory"),
        "在途数量": row.get("stockTransit"),
        "安全库存": row.get("minStock"),
        "状态": row.get("status"),
        "财务分类ID": row.get("taxCategoryRuleId"),
        "财务分类": row.get("taxCategoryRuleName"),
        "税务代码": row.get("taxCode"),
        "关联正品SKU": "、".join(str(item.get("skuCode") or item.get("skuId")) for item in (row.get("linkedSkus") or []) if item.get("skuCode") or item.get("skuId")),
        "档案ID": row.get("id"),
        "吉客云SKU ID": row.get("jackyunSkuId"),
    } for row in rows]
    headers = list(export_rows[0].keys()) if export_rows else [
        "类型", "编码", "名称", "商品名称", "条码", "单位", "货品类型", "品类", "采购单价", "售价",
        "默认成本", "成本方式", "成本容差", "当前库存", "工厂库存", "在途数量", "安全库存", "状态",
        "财务分类ID", "财务分类", "税务代码", "关联正品SKU", "档案ID", "吉客云SKU ID",
    ]
    sheet_name = "套装档案" if scope == "bundles" else "货品档案"
    return _workbook([(sheet_name, headers, export_rows)], template=template), len(rows)


def _tax_rules_export(db: Session, *, q: str, template: bool = False) -> tuple[bytes, int]:
    rows = tax_category_rule_service.list_rules(db, include_disabled=True)
    term = q.strip().casefold()
    if term:
        rows = [
            row for row in rows
            if term in " ".join(str(row.get(key) or "") for key in (
                "pattern", "categoryName", "itemName", "taxCode", "matchKeyword", "note",
            )).casefold()
        ]
    export_rows = [{
        "分类规则": row.get("pattern"),
        "财务大类": row.get("categoryName"),
        "项目名称": row.get("itemName"),
        "税务代码": row.get("taxCode"),
        "匹配关键词": row.get("matchKeyword"),
        "匹配方式": {"contains": "包含", "exact": "精确", "prefix": "前缀"}.get(row.get("matchMode"), row.get("matchMode")),
        "优先级": row.get("priority"),
        "启用": row.get("enabled"),
        "备注": row.get("note"),
        "规则ID": row.get("id"),
    } for row in rows]
    headers = ["分类规则", "财务大类", "项目名称", "税务代码", "匹配关键词", "匹配方式", "优先级", "启用", "备注", "规则ID"]
    return _workbook([("财务分类", headers, export_rows)], template=template), len(rows)


def _inventory_export(db: Session, *, q: str, template: bool = False) -> tuple[bytes, int]:
    rows = dashboard.inventory_skus(db, search=q, limit=2000)
    summary = dashboard.inventory_summary(db)
    export_rows = [{
        "SKU": row.get("skuCode"),
        "商品": row.get("goodsName") or row.get("skuName"),
        "吉客云SKU ID": row.get("jackyunSkuId"),
        "条码": row.get("barcode"),
        "单位": row.get("unit"),
        "当前库存": row.get("quantity"),
        "最后单据时间": row.get("lastDocumentAt"),
        "仓库分布": row.get("warehouses"),
        "有出入库单据": row.get("hasMovement"),
        "状态": row.get("status"),
    } for row in rows]
    warehouse_rows = [{
        "仓库": row.get("warehouseName"),
        "仓库ID": row.get("warehouseId"),
        "库存数量": row.get("quantity"),
        "SKU数": row.get("skus"),
        "最后单据时间": summary.get("lastDocumentAt"),
    } for row in summary.get("byWarehouse", [])]
    return _workbook([
        ("SKU库存", ["SKU", "商品", "吉客云SKU ID", "条码", "单位", "当前库存", "最后单据时间", "仓库分布", "有出入库单据", "状态"], export_rows),
        ("仓库汇总", ["仓库", "仓库ID", "库存数量", "SKU数", "最后单据时间"], warehouse_rows),
    ], template=template), len(rows)


def _warehouses_export(db: Session, *, template: bool = False) -> tuple[bytes, int]:
    rows = list_warehouses(db, include_inactive=True)
    export_rows = [{
        "仓库编号": row.get("code"),
        "仓库名称": row.get("name"),
        "类型": row.get("warehouseType"),
        "用途": row.get("purpose"),
        "吉客云仓库ID": row.get("jackyunWarehouseId"),
        "参与可售": row.get("isSellable"),
        "状态": row.get("status"),
        "来源": row.get("source"),
        "备注": row.get("note"),
        "档案ID": row.get("id"),
    } for row in rows]
    return _workbook([("仓库档案", ["仓库编号", "仓库名称", "类型", "用途", "吉客云仓库ID", "参与可售", "状态", "来源", "备注", "档案ID"], export_rows)], template=template), len(rows)


def _suppliers_export(db: Session, *, q: str, status: str, template: bool = False) -> tuple[bytes, int]:
    query = db.query(Supplier).order_by(Supplier.name, Supplier.id)
    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(
            or_(Supplier.name.ilike(like), Supplier.tax_no.ilike(like), Supplier.contact.ilike(like), Supplier.phone.ilike(like))
        )
    if status == "normal":
        query = query.filter(Supplier.is_temp.is_(False))
    elif status == "temp":
        query = query.filter(Supplier.is_temp.is_(True))
    rows = query.limit(20_000).all()
    export_rows = [{
        "供应商名称": row.name,
        "税号": row.tax_no,
        "平台": row.platform,
        "店铺/供应商编码": row.external_shop_id,
        "联系人": row.contact,
        "电话": row.phone,
        "地址": row.address,
        "备注": row.notes,
        "临时供应商": row.is_temp,
        "供应商ID": row.id,
    } for row in rows]
    headers = ["供应商名称", "税号", "平台", "店铺/供应商编码", "联系人", "电话", "地址", "备注", "临时供应商", "供应商ID"]
    return _workbook([("供应商档案", headers, export_rows)], template=template), len(rows)


def _external_orders_export(db: Session, *, q: str, status: str, template: bool = False) -> tuple[bytes, int]:
    query = db.query(ExternalPurchaseOrder).order_by(ExternalPurchaseOrder.ordered_at.desc().nullslast(), ExternalPurchaseOrder.id.desc())
    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(
            or_(ExternalPurchaseOrder.external_order_id.ilike(like), ExternalPurchaseOrder.supplier_name.ilike(like), ExternalPurchaseOrder.title.ilike(like))
        )
    if status:
        query = query.filter(ExternalPurchaseOrder.purchase_status == status)
    rows = [row for row in query.limit(20_000).all() if not (row.raw or {}).get("referenceOnly")]
    export_rows = [{
        "订单号": row.external_order_id,
        "渠道": row.platform,
        "供应商/往来单位": row.supplier_name,
        "备注": row.title,
        "采购时间": row.ordered_at,
        "订单金额": row.order_amount,
        "实付金额": row.paid_amount,
        "买家账号": row.buyer_account,
        "订单状态": row.order_status,
        "采购状态": row.purchase_status,
        "订单ID": row.id,
    } for row in rows]
    headers = ["订单号", "渠道", "供应商/往来单位", "备注", "采购时间", "订单金额", "实付金额", "买家账号", "订单状态", "采购状态", "订单ID"]
    return _workbook([("其他渠道订单", headers, export_rows)], template=template), len(rows)


def _production_export(db: Session, *, status: str, template: bool = False) -> tuple[bytes, int]:
    rows = production_service.list_production_orders(db, status=status, limit=2000)
    orders: list[dict[str, Any]] = []
    items: list[dict[str, Any]] = []
    materials: list[dict[str, Any]] = []
    for row in rows:
        orders.append({
            "生产单号": row.get("orderNo"),
            "工厂": row.get("factoryName"),
            "状态": row.get("status"),
            "计划开始": row.get("plannedStartDate"),
            "预计交货": row.get("expectedDeliveryDate"),
            "来源": row.get("sourceType"),
            "备注": row.get("note"),
            "货品数": row.get("itemCount"),
            "耗材数": row.get("materialCount"),
            "缺料数": row.get("materialShortageCount"),
        })
        for item in row.get("items") or []:
            items.append({
                "生产单号": row.get("orderNo"),
                "SKU": item.get("skuCode"),
                "商品": item.get("skuName"),
                "计划数量": item.get("quantity"),
                "已完成": item.get("completedQty"),
                "待入库": item.get("pendingInboundQty"),
                "入库数量": item.get("inboundQty"),
            })
        for material in row.get("materials") or []:
            materials.append({
                "生产单号": row.get("orderNo"),
                "耗材编码": material.get("code"),
                "耗材名称": material.get("name"),
                "单位": material.get("unit"),
                "需求数量": material.get("requiredQty"),
                "预占数量": material.get("reservedQty"),
                "发出数量": material.get("dispatchedQty"),
                "工厂收到": material.get("factoryReceivedQty"),
                "已消耗": material.get("consumedQty"),
                "状态": material.get("state"),
            })
    return _workbook([
        ("生产订单", ["生产单号", "工厂", "状态", "计划开始", "预计交货", "来源", "备注", "货品数", "耗材数", "缺料数"], orders),
        ("生产货品", ["生产单号", "SKU", "商品", "计划数量", "已完成", "待入库", "入库数量"], items),
        ("生产耗材", ["生产单号", "耗材编码", "耗材名称", "单位", "需求数量", "预占数量", "发出数量", "工厂收到", "已消耗", "状态"], materials),
    ], template=template), len(rows)


def _material_flow_export(db: Session, *, group: str, stage: str, q: str, template: bool = False) -> tuple[bytes, int]:
    payload = list_production_purchase_rows(db, group=group, stage=stage, q=q, limit=2000)
    rows = payload.get("rows", [])
    export_rows = [{
        "来源订单": row.get("orderNo"),
        "渠道": row.get("platform"),
        "供应商/工厂": row.get("supplier"),
        "生产内容": row.get("items"),
        "数量": row.get("quantityTotal"),
        "采购状态": row.get("purchaseStatus"),
        "当前阶段": row.get("stageLabel"),
        "采购单号": row.get("jackyunPurchaseNos"),
        "入库单号": row.get("jackyunInboundNos"),
        "入库仓库": row.get("inboundWarehouses"),
        "物流状态": row.get("shipStatus"),
        "下单时间": row.get("orderDate"),
    } for row in rows]
    return _workbook([("生产采购链路", ["来源订单", "渠道", "供应商/工厂", "生产内容", "数量", "采购状态", "当前阶段", "采购单号", "入库单号", "入库仓库", "物流状态", "下单时间"], export_rows)], template=template), len(rows)


def _inbound_export(db: Session, *, q: str, start_date: date | None, end_date: date | None, template: bool = False) -> tuple[bytes, int]:
    query = db.query(JackyunGoodsDocument).filter(JackyunGoodsDocument.document_type == "inbound")
    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(or_(JackyunGoodsDocument.goodsdoc_no.ilike(like), JackyunGoodsDocument.warehouse_name.ilike(like), JackyunGoodsDocument.supplier_name.ilike(like)))
    if start_date:
        query = query.filter(JackyunGoodsDocument.document_at >= start_date)
    if end_date:
        query = query.filter(JackyunGoodsDocument.document_at < end_date + timedelta(days=1))
    documents = query.order_by(JackyunGoodsDocument.document_at.desc().nullslast(), JackyunGoodsDocument.id.desc()).limit(20_000).all()
    document_ids = [document.id for document in documents]
    item_rows = db.query(JackyunGoodsDocumentItem).filter(JackyunGoodsDocumentItem.document_id.in_(document_ids)).order_by(JackyunGoodsDocumentItem.document_id, JackyunGoodsDocumentItem.line_no).all() if document_ids else []
    items_by_document: dict[int, list[JackyunGoodsDocumentItem]] = {}
    for item in item_rows:
        items_by_document.setdefault(item.document_id, []).append(item)
    master = [{
        "入库单号": document.goodsdoc_no,
        "入库日期": document.document_at,
        "仓库": document.warehouse_name,
        "仓库编码": document.warehouse_code,
        "供应商": document.supplier_name or document.company_name,
        "总数量": document.total_quantity,
        "总金额": document.total_amount,
        "总费用": document.total_fee,
    } for document in documents]
    details = [{
        "入库单号": next((document.goodsdoc_no for document in documents if document.id == item.document_id), ""),
        "行号": item.line_no,
        "货品编号": item.goods_no,
        "SKU/条码": item.sku_barcode,
        "货品名称": item.goods_name,
        "规格": item.spec,
        "数量": item.quantity,
        "单位": item.unit_name,
        "含税单价": item.unit_price_tax,
        "含税金额": item.amount_tax,
        "匹配SKU": item.matched_sku_id,
        "匹配状态": item.match_status,
    } for item in item_rows]
    return _workbook([
        ("入库单", ["入库单号", "入库日期", "仓库", "仓库编码", "供应商", "总数量", "总金额", "总费用"], master),
        ("入库明细", ["入库单号", "行号", "货品编号", "SKU/条码", "货品名称", "规格", "数量", "单位", "含税单价", "含税金额", "匹配SKU", "匹配状态"], details),
    ], template=template), len(documents)


def _sales_export(db: Session, *, q: str, status: str, start_date: date | None, end_date: date | None, template: bool = False) -> tuple[bytes, int]:
    conditions = []
    if q.strip():
        like = f"%{q.strip()}%"
        conditions.append(or_(
            SalesOrder.order_no.ilike(like),
            SalesOrderItem.sku_code.ilike(like),
            SalesOrderItem.goods_name.ilike(like),
            SalesOrder.raw["netOrderNo"].astext.ilike(like),
        ))
    if status:
        conditions.append(SalesOrder.order_status == status)
    if start_date:
        conditions.append(SalesOrder.ordered_at >= start_date)
    if end_date:
        conditions.append(SalesOrder.ordered_at < end_date + timedelta(days=1))
    query = db.query(SalesOrderItem, SalesOrder).join(SalesOrder, SalesOrder.id == SalesOrderItem.order_id)
    if conditions:
        query = query.filter(*conditions)
    rows = query.order_by(SalesOrder.ordered_at.desc().nullslast(), SalesOrder.order_no.desc(), SalesOrderItem.id).limit(50_000).all()
    export_rows = []
    for item, order in rows:
        order_raw = order.raw or {}
        item_raw = item.raw or {}
        export_rows.append({
            "订单编号": order.order_no,
            "网店单号": order_raw.get("netOrderNo"),
            "下单时间": order.ordered_at,
            "平台": order.platform,
            "订单类型": order.order_type,
            "订单状态": order.order_status,
            "支付状态": order.pay_status,
            "仓库": order_raw.get("warehouseName") or order_raw.get("warehouse"),
            "SKU": item.sku_code,
            "商品": item.goods_name,
            "规格": item_raw.get("spec"),
            "数量": item.quantity,
            "单价": item.unit_price,
            "金额": item.amount,
            "订单金额": order.order_amount,
            "实付金额": order.paid_amount,
            "来源": order.source_provider,
        })
    return _workbook([("销售明细", ["订单编号", "网店单号", "下单时间", "平台", "订单类型", "订单状态", "支付状态", "仓库", "SKU", "商品", "规格", "数量", "单价", "金额", "订单金额", "实付金额", "来源"], export_rows)], template=template), len(rows)


def _tax_export(db: Session, *, q: str, status: str, template: bool = False) -> tuple[bytes, int]:
    query = db.query(TaxInvoice)
    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(or_(TaxInvoice.invoice_number.ilike(like), TaxInvoice.seller_name.ilike(like), TaxInvoice.buyer_name.ilike(like)))
    if status:
        query = query.filter(TaxInvoice.match_status == status)
    rows = query.order_by(TaxInvoice.issue_date.desc().nullslast(), TaxInvoice.id.desc()).limit(20_000).all()
    export_rows = [{
        "发票号码": row.invoice_number,
        "发票代码": row.invoice_code,
        "进销项": row.direction,
        "发票类型": row.invoice_type,
        "开票日期": row.issue_date,
        "销售方名称": row.seller_name,
        "销售方税号": row.seller_tax_id,
        "购买方名称": row.buyer_name,
        "购买方税号": row.buyer_tax_id,
        "不含税金额": row.amount_excl_tax,
        "税额": row.tax_amount,
        "价税合计": row.total_amount,
        "关联状态": row.match_status,
        "关联说明": row.match_note,
        "已认证": row.verified,
        "认证账期": row.verified_month,
    } for row in rows]
    return _workbook([("发票台账", ["发票号码", "发票代码", "进销项", "发票类型", "开票日期", "销售方名称", "销售方税号", "购买方名称", "购买方税号", "不含税金额", "税额", "价税合计", "关联状态", "关联说明", "已认证", "认证账期"], export_rows)], template=template), len(rows)


def build_export(
    db: Session,
    dataset: str,
    *,
    status: str = "",
    q: str = "",
    channel: str = "all",
    kind: str = "all",
    warehouse: str = "",
    start_date: date | None = None,
    end_date: date | None = None,
    group: str = "all",
    stage: str = "",
    template: bool = False,
) -> tuple[bytes, int, str]:
    """生成指定数据集的核验 Excel，返回内容、主记录数和文件名。

    ``template=True`` 时只输出表头（空模板），供用户填写后从数据接入重新导入。
    """
    labels = {item["key"]: item["label"] for item in DATASET_OPTIONS}
    if dataset not in labels:
        raise ValueError("不支持的数据导出类型")
    if start_date and end_date and start_date > end_date:
        raise ValueError("开始日期不能晚于结束日期")

    if dataset == "purchase_orders":
        content, count = _purchase_export(db, status=status, q=q, channel=channel, kind=kind, warehouse=warehouse, start_date=start_date, end_date=end_date, template=template)
    elif dataset == "catalog":
        content, count = _catalog_export(db, kind=kind, q=q, template=template, scope="catalog")
    elif dataset == "inventory":
        content, count = _inventory_export(db, q=q, template=template)
    elif dataset == "warehouses":
        content, count = _warehouses_export(db, template=template)
    elif dataset == "bundles":
        content, count = _catalog_export(db, kind="goods", q=q, template=template, scope="bundles")
    elif dataset == "tax_rules":
        content, count = _tax_rules_export(db, q=q, template=template)
    elif dataset == "suppliers":
        content, count = _suppliers_export(db, q=q, status=status, template=template)
    elif dataset == "external_orders":
        content, count = _external_orders_export(db, q=q, status=status, template=template)
    elif dataset == "production_orders":
        content, count = _production_export(db, status=status, template=template)
    elif dataset == "material_flow":
        content, count = _material_flow_export(db, group=group, stage=stage, q=q, template=template)
    elif dataset == "inbound_documents":
        content, count = _inbound_export(db, q=q, start_date=start_date, end_date=end_date, template=template)
    elif dataset == "sales_items":
        content, count = _sales_export(db, q=q, status=status, start_date=start_date, end_date=end_date, template=template)
    else:
        content, count = _tax_export(db, q=q, status=status, template=template)
    return content, count, labels[dataset]
