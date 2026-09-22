"""月结财务汇总：基于统一 FinanceEntry 生成外贸财务表。"""
from __future__ import annotations

from collections import defaultdict
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.finance import FinanceEntry, FinanceLegalEntity
from app.services.finance_center_service import CATEGORY_LABELS


def resolve_entity_by_name(db: Session, company: str) -> FinanceLegalEntity | None:
    if not company.strip():
        return None
    return db.scalar(
        select(FinanceLegalEntity).where(
            FinanceLegalEntity.name == company.strip(),
            FinanceLegalEntity.status != "archived",
        )
    )


def foreign_trade_rows(
    db: Session,
    *,
    legal_entity_id: int,
    year: int,
    month: int,
) -> list[dict[str, Any]]:
    rows = db.scalars(
        select(FinanceEntry)
        .where(
            FinanceEntry.legal_entity_id == legal_entity_id,
            FinanceEntry.business_scope == "foreign_trade",
            FinanceEntry.accounting_year == year,
            FinanceEntry.accounting_month == month,
        )
        .order_by(FinanceEntry.occurred_at.asc(), FinanceEntry.id.asc())
    ).all()
    return [
        {
            "sourceNo": row.source_no,
            "sourceType": row.source_type,
            "category": row.category,
            "categoryLabel": CATEGORY_LABELS.get(row.category, row.category),
            "direction": row.direction,
            "currency": row.currency,
            "amount": str(row.amount or 0),
            "taxAmount": str(row.tax_amount or 0),
            "valueType": row.value_type,
            "settlementStatus": row.settlement_status,
            "invoiceStatus": row.invoice_status,
            "cashEffect": bool(row.cash_effect),
            "profitEffect": bool(row.profit_effect),
            "occurredAt": row.occurred_at.isoformat() if row.occurred_at else "",
            "note": row.note or "",
        }
        for row in rows
    ]


def foreign_trade_summary(
    db: Session,
    *,
    legal_entity_id: int,
    year: int,
    month: int,
) -> dict[str, Any]:
    rows = foreign_trade_rows(
        db, legal_entity_id=legal_entity_id, year=year, month=month
    )
    totals: dict[str, dict[str, Decimal]] = defaultdict(
        lambda: {
            "income": Decimal("0"),
            "expense": Decimal("0"),
            "profitIncome": Decimal("0"),
            "profitExpense": Decimal("0"),
            "cashIn": Decimal("0"),
            "cashOut": Decimal("0"),
        }
    )
    for row in rows:
        amount = Decimal(str(row["amount"] or "0"))
        bucket = totals[row["currency"]]
        if row["direction"] == "income":
            bucket["income"] += amount
            if row["profitEffect"]:
                bucket["profitIncome"] += amount
            if row["cashEffect"]:
                bucket["cashIn"] += amount
        elif row["direction"] == "expense":
            bucket["expense"] += amount
            if row["profitEffect"]:
                bucket["profitExpense"] += amount
            if row["cashEffect"]:
                bucket["cashOut"] += amount
        # 其他方向（脏数据）跳过，避免被静默误分类为支出
    return {
        "rowCount": len(rows),
        "rows": rows,
        "totalsByCurrency": [
            {
                "currency": currency,
                **{key: float(value) for key, value in values.items()},
                "profit": float(values["profitIncome"] - values["profitExpense"]),
                "netCash": float(values["cashIn"] - values["cashOut"]),
            }
            for currency, values in sorted(totals.items())
        ],
    }


def foreign_trade_xlsx(
    db: Session,
    *,
    legal_entity_id: int,
    year: int,
    month: int,
) -> bytes | None:
    entity = db.get(FinanceLegalEntity, legal_entity_id)
    if entity is None:
        return None
    report = foreign_trade_summary(
        db, legal_entity_id=legal_entity_id, year=year, month=month
    )
    if not report["rows"]:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "外贸财务汇总"

    ws["A1"] = f"{entity.name} {year}年{month:02d}月 外贸财务汇总"
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["币种", "收入", "支出", "利润口径", "现金流入", "现金流出", "净现金"])
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for row in report["totalsByCurrency"]:
        ws.append([
            row["currency"], row["income"], row["expense"], row["profit"],
            row["cashIn"], row["cashOut"], row["netCash"],
        ])
    ws.append([])
    headers = [
        "来源单号", "来源类型", "财务类别", "方向", "币种", "金额", "税额",
        "预计/实际", "结算状态", "发票状态", "影响现金", "影响利润", "发生时间", "备注",
    ]
    ws.append(headers)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)

    for row in report["rows"]:
        ws.append([
            row["sourceNo"], row["sourceType"], row["categoryLabel"],
            "收入" if row["direction"] == "income" else "支出",
            row["currency"], float(row["amount"]), float(row["taxAmount"]),
            "预计" if row["valueType"] == "estimated" else "实际",
            row["settlementStatus"], row["invoiceStatus"],
            "是" if row["cashEffect"] else "否",
            "是" if row["profitEffect"] else "否",
            row["occurredAt"], row["note"],
        ])

    for col in (2, 3, 4, 5, 6, 7):
        for row_idx in range(4, min(header_row, ws.max_row) + 1):
            ws.cell(row=row_idx, column=col).number_format = "0.00"
    for col in (6, 7):
        for row_idx in range(header_row + 1, ws.max_row + 1):
            ws.cell(row=row_idx, column=col).number_format = "0.00"

    widths = [18, 22, 18, 10, 9, 14, 12, 10, 12, 12, 10, 10, 22, 34]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = f"A{header_row + 1}"

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
