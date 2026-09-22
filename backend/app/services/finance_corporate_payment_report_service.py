"""月度「已收票付款方式核对」清单（对公 / 个人 / 平台扣款 / 混合）。

以已经落库的银行付款↔进项发票关联为唯一付款事实，并按发票展示采购订单关联。
当月进项发票全部保留用于财务交付，但只有有效、正数金额发票参与银行付款核对；
红冲、作废、待确认或非正数金额发票明确标记为“无需核对银行付款”。
"""
from __future__ import annotations

from collections import Counter, defaultdict
from decimal import Decimal
from io import BytesIO
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.core.audit import audit
from app.models.bank import BankAccount, BankTransaction
from app.models.finance import FinanceCorporatePaymentAdjustment
from app.models.purchase import Supplier
from app.models.tax import TaxInvoice, TaxInvoiceLink
from app.services import invoice_reconciliation, tax_invoice_service
from app.services.monthly_core import month_bounds

TOLERANCE = Decimal("0.05")
TARGET_TYPE = "bank_transaction"

# 财务主视图使用“支付方式 + 费用性质”两个维度。
# 银行流水是对公付款的唯一证据；没有已确认银行流水的有效正数进项票，
# 在本月财务交付表中按业务约定直接归为“个人支付”，而不是继续显示成待判断。
_PAYMENT_SOURCE_LABELS = {
    "corporate": "对公支付",
    "personal": "个人支付",
    "platform_auto_debit": "平台自动扣款货款",
    "mixed": "混合支付",
    "not_applicable": "不适用",
}

_EXPENSE_NATURE_FALLBACK = {
    "goods": ("goods_payment", "货款"),
    "platform_fee": ("platform_service_fee", "平台/技术服务费"),
    "operating_other": ("operating_other", "其他经营费用"),
    "reimburse_advance": ("advance_payment", "代付款/垫付款"),
    "reimburse_operating": ("reimbursement", "费用报销"),
    "excluded": ("excluded", "不计入经营/报销"),
}

_EXPENSE_NATURE_RULES = (
    ("meal", "餐饮费", re.compile(r"餐饮|餐费|餐厅|饭店|快餐|餐饮服务")),
    ("lodging", "住宿费", re.compile(r"住宿|酒店|宾馆|旅店|客房")),
    ("travel_transport", "交通费", re.compile(r"机票|航空|高铁|铁路|客运|出租车|打车|网约车|停车|过路|通行费")),
    ("logistics", "物流运输费", re.compile(r"物流|快递|运输服务|货运|配送")),
    ("warehouse", "仓储费", re.compile(r"仓储|仓库服务")),
    ("advertising", "广告宣传费", re.compile(r"广告|推广|宣传|营销|投流")),
    ("office", "办公费", re.compile(r"办公用品|文具|打印|耗材")),
    ("rent", "租赁费", re.compile(r"房租|租金|租赁")),
    ("utilities", "水电物业费", re.compile(r"水费|电费|物业|燃气")),
    ("platform_service_fee", "平台/技术服务费", re.compile(r"平台服务|技术服务|信息服务|软件|云服务|服务器|网络服务|会员|商标|经纪代理")),
)


def _payment_source(
    db: Session,
    invoice: TaxInvoice,
    bank_status: str,
) -> tuple[str, str, str]:
    """财务交付只解释证据；支付分类统一复用发票服务的单一派生入口。"""
    canonical_bank_status = {
        "paid": "matched",
        "unpaid": "unmatched",
        "partial": "partial",
        "overpaid_after_red": "overpaid_after_red",
        "red_overpayment_settled": "red_overpayment_settled",
        "not_applicable": "not_applicable",
    }.get(bank_status, "not_applicable")
    payment = tax_invoice_service.payment_method_context(invoice, canonical_bank_status)
    key = str(payment.get("paymentMethod") or "not_applicable")

    if key == "corporate":
        basis = (
            "历史对公付款超过红冲后有效金额，超额部分待退款或冲抵"
            if bank_status == "overpaid_after_red"
            else "历史红冲超额付款已通过退款/后续冲抵完成处理"
            if bank_status == "red_overpayment_settled"
            else "已确认银行支出流水"
        )
    elif key == "mixed":
        basis = (
            "银行部分付款 + 已确认个人垫付"
            if invoice.payment_method == "personal"
            else "银行部分付款；剩余部分按无对公流水系统推定为个人支付"
        )
    elif key == "personal":
        basis = (
            "人工已确认个人垫付"
            if invoice.payment_method == "personal"
            else "未匹配到对公银行支出流水，按月结规则系统推定"
        )
    elif key == "platform_auto_debit":
        basis = "人工已确认平台自动扣款货款（不经过银行流水）"
    else:
        basis = tax_invoice_service.bank_payment_reconciliation_ineligible_reason(invoice, db=db)

    label = _PAYMENT_SOURCE_LABELS[key]
    if key == "personal" and invoice.payment_method != "personal":
        label = "个人支付（系统推定）"
    elif key == "personal":
        label = "个人支付（已确认）"
    elif key == "platform_auto_debit":
        label = "平台自动扣款货款"
    elif key == "mixed" and invoice.payment_method != "personal":
        label = "混合支付（系统推定）"
    return key, label, basis


def _expense_nature(
    invoice: TaxInvoice,
    covered_orders: list[dict[str, Any]],
    line_items: list[dict[str, Any]],
) -> tuple[str, str, str]:
    """给财务看的费用性质。

    优先级：
    1. 已明确关联采购订单 → 货款；
    2. 发票明细关键词 → 餐饮/住宿/交通/物流/服务费等；
    3. 现有 category → 财务大类兜底；
    4. 无法识别 → 待分类。
    """
    if covered_orders:
        return "goods_payment", "货款", "采购订单关联"

    goods_names = [
        str(item.get("goodsName") or item.get("itemName") or "").strip()
        for item in line_items
        if isinstance(item, dict)
    ]
    text = " ".join(name for name in goods_names if name)
    for key, label, pattern in _EXPENSE_NATURE_RULES:
        if pattern.search(text):
            return key, label, "发票明细"

    # 发票导入 raw 中有时直接保留“货物或应税劳务、服务名称”等字段；
    # 具体费用性质优先于宽泛 category（例如 reimburse_operating 还能继续识别为餐饮/住宿）。
    raw = invoice.raw if isinstance(invoice.raw, dict) else {}
    raw_text = " ".join(str(value or "") for value in raw.values())
    for key, label, pattern in _EXPENSE_NATURE_RULES:
        if pattern.search(raw_text):
            return key, label, "发票原始明细"

    fallback = _EXPENSE_NATURE_FALLBACK.get(invoice.category or "")
    if fallback:
        return fallback[0], fallback[1], "发票分类"

    return "unclassified", "待分类", "未识别"


def _dec(value: Decimal | float | int | str | None) -> Decimal:
    if value is None:
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    return Decimal(str(value))


def _month_range(year: int, month: int):
    """统一使用 Asia/Shanghai 业务自然月 [start, end)。"""
    return month_bounds(year, month)


def _link_amount(link: TaxInvoiceLink, invoice: TaxInvoice | None) -> Decimal:
    """历史 NULL 分摊按旧版本“整票关联”解释，避免升级后付款事实凭空消失。"""
    if link.allocated_amount is not None:
        return _dec(link.allocated_amount)
    if invoice is not None:
        return _dec(invoice.total_amount)
    return Decimal("0")


def _active_bank_links(db: Session, *, invoice_ids: list[int] | None = None) -> list[TaxInvoiceLink]:
    q = (
        db.query(TaxInvoiceLink)
        .filter(
            TaxInvoiceLink.target_type == TARGET_TYPE,
            TaxInvoiceLink.confirmed.is_(True),
            TaxInvoiceLink.match_method != "rejected",
        )
    )
    if invoice_ids:
        q = q.filter(TaxInvoiceLink.invoice_id.in_(invoice_ids))
    return q.all()


def _company_matches(invoice: TaxInvoice, company: str) -> bool:
    """买方名称有值时按公司主体隔离；历史票缺买方名称时不误删。"""
    if not company or not (invoice.buyer_name or "").strip():
        return True
    return (
        invoice_reconciliation.normalize_supplier(invoice.buyer_name)
        == invoice_reconciliation.normalize_supplier(company)
    )


def _bank_reconciliation_meta(
    db: Session,
    invoice: TaxInvoice,
    paid_total: Decimal,
    bank_meta: dict[str, Any] | None = None,
) -> tuple[str, Decimal, bool, str]:
    """月结银行核对复用发票台账的同一事实源，避免红冲结算后两处状态漂移。"""
    bank_meta = bank_meta or tax_invoice_service._invoice_bank_payment_context(db, [invoice]).get(invoice.id, {})
    source_status = str(bank_meta.get("bankPaymentStatus") or "")
    total = _dec(bank_meta.get("bankEffectiveInvoiceAmount") or tax_invoice_service.effective_invoice_amount_after_red(db, invoice))
    outstanding = _dec(bank_meta.get("bankRemainingAmount"))
    overpaid = _dec(bank_meta.get("bankOverpaidAmount"))
    unresolved_overpaid = _dec(bank_meta.get("bankOverpaidUnsettledAmount", overpaid))

    if source_status == "overpaid_after_red":
        return (
            "overpaid_after_red",
            Decimal("0"),
            True,
            f"历史对公付款 {paid_total} 超过红冲后有效金额 {total}；历史超额 {overpaid}，当前仍待退款/冲抵 {unresolved_overpaid}",
        )
    if source_status == "red_overpayment_settled":
        return (
            "red_overpayment_settled",
            Decimal("0"),
            True,
            f"历史对公付款曾超出红冲后有效金额 {overpaid}，现已通过退款/后续冲抵完成处理",
        )
    if source_status == "not_applicable":
        return (
            "not_applicable",
            Decimal("0"),
            False,
            tax_invoice_service.bank_payment_reconciliation_ineligible_reason(invoice, db=db),
        )
    if source_status == "unmatched":
        return "unpaid", total, True, ""
    if source_status == "partial":
        return "partial", outstanding, True, ""
    if source_status == "matched":
        return "paid", Decimal("0"), True, ""

    # 兼容极少量历史调用：没有派生上下文时按原规则保守计算。
    if not tax_invoice_service.is_bank_payment_reconciliation_eligible(invoice, db=db):
        return (
            "not_applicable", Decimal("0"), False,
            tax_invoice_service.bank_payment_reconciliation_ineligible_reason(invoice, db=db),
        )
    if paid_total <= 0:
        return "unpaid", total, True, ""
    outstanding = max(total - paid_total, Decimal("0"))
    return ("paid" if outstanding <= TOLERANCE else "partial"), outstanding, True, ""


def _invoice_purchase_map(db: Session, invoices: list[TaxInvoice]) -> dict[int, list[dict[str, Any]]]:
    """复用采购发票对账的人工优先 + 截止开票日 FIFO 结果。"""
    by_supplier: dict[str, list[TaxInvoice]] = defaultdict(list)
    for inv in invoices:
        key = invoice_reconciliation.normalize_supplier(inv.seller_name)
        if key:
            by_supplier[key].append(inv)

    result: dict[int, list[dict[str, Any]]] = {}
    for rows in by_supplier.values():
        seller = rows[0].seller_name
        reconciliation = invoice_reconciliation.reconcile(db, supplier=seller)
        for supplier_block in reconciliation.get("suppliers", []):
            for month_block in supplier_block.get("months", []):
                for inv_row in month_block.get("invoices", []):
                    result[int(inv_row["invoiceId"])] = list(inv_row.get("covered") or [])
    return result


def _period_invoices(db: Session, year: int, month: int, company: str) -> list[TaxInvoice]:
    """当月收到的进项发票（按公司主体隔离）；保存选择与构建报告共用同一口径。"""
    start, end = _month_range(year, month)
    invoices = (
        tax_invoice_service.filter_visible_invoices(
            db.query(TaxInvoice).filter(
                TaxInvoice.direction == "input",
                TaxInvoice.issue_date >= start,
                TaxInvoice.issue_date < end,
            )
        )
        .order_by(TaxInvoice.issue_date, TaxInvoice.id)
        .all()
    )
    return [inv for inv in invoices if _company_matches(inv, company)]



def _period_company_mismatch_summary(db: Session, year: int, month: int, company: str) -> dict[str, Any]:
    if not company:
        return {"count": 0, "amount": "0.00", "invoiceNos": []}
    start, end = _month_range(year, month)
    rows = (
        tax_invoice_service.filter_visible_invoices(
            db.query(TaxInvoice).filter(
                TaxInvoice.direction == "input",
                TaxInvoice.issue_date >= start,
                TaxInvoice.issue_date < end,
            )
        )
        .all()
    )
    mismatched = [
        row for row in rows
        if (row.buyer_name or "").strip() and not _company_matches(row, company)
    ]
    return {
        "count": len(mismatched),
        "amount": str(sum((_dec(row.total_amount) for row in mismatched), Decimal("0"))),
        "invoiceNos": [row.invoice_number or "" for row in mismatched[:50]],
    }


def build_report(
    db: Session,
    year: int,
    month: int,
    company: str = "",
    *,
    selected_keys: list[str] | None = None,
) -> dict[str, Any]:
    """按“当月收到的进项发票”组织财务核对清单。

    发票是月度交付的主维度；银行付款和采购订单关联都挂在发票下面。
    已收票但尚未匹配银行付款的发票也必须进入清单，不能因为未付款而被漏掉。
    selected_keys 以 invoice_key 为规范键；兼容历史唯一的发票号码，歧义号码拒绝猜测。
    """
    invoices = _period_invoices(db, year, month, company)
    invoice_map = {inv.id: inv for inv in invoices}
    invoice_ids = sorted(invoice_map)

    links = _active_bank_links(db, invoice_ids=invoice_ids) if invoice_ids else []
    txn_ids = sorted({link.target_id for link in links})
    txns = (
        db.query(BankTransaction)
        .filter(BankTransaction.id.in_(txn_ids), BankTransaction.direction == "out")
        .all()
        if txn_ids
        else []
    )
    txn_map = {row.id: row for row in txns}

    account_ids = sorted({txn.account_id for txn in txns if txn.account_id})
    accounts = (
        db.query(BankAccount).filter(BankAccount.id.in_(account_ids)).all()
        if account_ids
        else []
    )
    account_map = {row.id: row for row in accounts}

    invoice_bank_total: dict[int, Decimal] = defaultdict(Decimal)
    txn_bank_total: dict[int, Decimal] = defaultdict(Decimal)
    links_by_invoice: dict[int, list[TaxInvoiceLink]] = defaultdict(list)
    for link in links:
        amount = _link_amount(link, invoice_map.get(link.invoice_id))
        invoice_bank_total[link.invoice_id] += amount
        txn_bank_total[link.target_id] += amount
        links_by_invoice[link.invoice_id].append(link)

    supplier_tax: dict[str, str] = {}
    for supplier in db.query(Supplier).all():
        key = invoice_reconciliation.normalize_supplier(supplier.name)
        if key and supplier.tax_no:
            supplier_tax.setdefault(key, supplier.tax_no)

    red_context = tax_invoice_service.red_accounting_context(db, invoices) if invoices else {}
    bank_context = tax_invoice_service._invoice_bank_payment_context(db, invoices) if invoices else {}
    settlement_context = tax_invoice_service._red_settlement_context(db, invoices) if invoices else {}
    vat_context = tax_invoice_service._vat_context(db, invoices, red_context) if invoices else {}
    company_mismatch = _period_company_mismatch_summary(db, year, month, company)
    related_blue_ids = sorted({
        int(ctx["redRelatedInvoiceId"])
        for ctx in red_context.values()
        if ctx.get("invoiceColor") == "red" and ctx.get("redRelatedInvoiceId")
    })
    related_blue_rows = (
        db.query(TaxInvoice).filter(TaxInvoice.id.in_(related_blue_ids)).all()
        if related_blue_ids else []
    )
    related_blue_map = {row.id: row for row in related_blue_rows}
    classification_invoices = list(invoices) + [
        row for row in related_blue_rows if row.id not in invoice_map
    ]
    purchase_map = _invoice_purchase_map(db, classification_invoices) if classification_invoices else {}
    line_summary_map = (
        tax_invoice_service.invoice_line_summaries(db, classification_invoices)
        if classification_invoices else {}
    )
    invoice_rows: list[dict[str, Any]] = []
    flat_rows: list[dict[str, Any]] = []
    for inv in invoices:
        inv_total = _dec(inv.total_amount)
        paid_total = invoice_bank_total.get(inv.id, Decimal("0"))
        invoice_bank_meta = bank_context.get(inv.id, {})
        status, outstanding, reconciliation_applicable, reconciliation_reason = _bank_reconciliation_meta(
            db, inv, paid_total, invoice_bank_meta
        )
        effective_total = _dec(invoice_bank_meta.get("bankEffectiveInvoiceAmount") or tax_invoice_service.effective_invoice_amount_after_red(db, inv))
        overpaid_amount = _dec(invoice_bank_meta.get("bankOverpaidAmount"))
        overpaid_settled_amount = _dec(invoice_bank_meta.get("bankOverpaidSettledAmount"))
        overpaid_unsettled_amount = _dec(invoice_bank_meta.get("bankOverpaidUnsettledAmount", overpaid_amount))
        red_meta = red_context.get(inv.id, {})
        settlement_meta = settlement_context.get(inv.id, {})
        vat_meta = vat_context.get(inv.id, {})
        source_invoice = inv
        if red_meta.get("invoiceColor") == "red" and red_meta.get("redRelatedInvoiceId"):
            source_invoice = related_blue_map.get(int(red_meta["redRelatedInvoiceId"]), inv)
        covered = purchase_map.get(source_invoice.id, [])
        order_nos = [str(x.get("orderNo") or "") for x in covered if x.get("orderNo")]
        line_items = list((line_summary_map.get(source_invoice.id) or {}).get("lineItems") or [])
        payment_source, payment_source_label, payment_source_basis = _payment_source(db, inv, status)
        canonical_bank_status = {
            "paid": "matched",
            "unpaid": "unmatched",
            "partial": "partial",
            "overpaid_after_red": "overpaid_after_red",
            "red_overpayment_settled": "red_overpayment_settled",
            "not_applicable": "not_applicable",
        }.get(status, "not_applicable")
        payment_context = tax_invoice_service.payment_method_context(inv, canonical_bank_status)
        payment_method = str(payment_context.get("paymentMethod") or "")
        payment_method_label = tax_invoice_service.PAYMENT_METHOD_LABELS.get(payment_method, "未设置")
        expense_nature, expense_nature_label, expense_nature_basis = _expense_nature(
            source_invoice, covered, line_items
        )
        if source_invoice.id != inv.id:
            expense_nature_basis = f"继承对应蓝字发票 · {expense_nature_basis}"
        payments: list[dict[str, Any]] = []

        for link in sorted(links_by_invoice.get(inv.id, []), key=lambda row: row.id):
            txn = txn_map.get(link.target_id)
            if txn is None:
                continue
            account = account_map.get(txn.account_id)
            allocated = _link_amount(link, inv)
            txn_total = _dec(txn.amount)
            txn_allocated = min(
                max(txn_bank_total.get(txn.id, Decimal("0")), Decimal("0")),
                max(txn_total, Decimal("0")),
            )
            payment = {
                "linkId": link.id,
                "paymentId": txn.id,
                "paymentDate": txn.txn_date.isoformat(),
                "paymentAccount": account.account_no if account else "",
                "paymentAccountName": account.account_name if account else "",
                "counterpartyAccount": txn.counterparty_account or "",
                "serialNo": txn.serial_no or "",
                "voucherNo": txn.voucher_no or "",
                "summary": txn.summary or "",
                "paymentAmount": str(txn_total),
                "allocatedAmount": str(allocated),
                "paymentMatchedTotal": str(txn_allocated),
                "paymentStatus": "matched" if txn_total - txn_allocated <= TOLERANCE else "partial",
            }
            payments.append(payment)
            flat_rows.append({
                **payment,
                "supplierName": inv.seller_name or txn.counterparty_name,
                "supplierTaxId": inv.seller_tax_id or supplier_tax.get(
                    invoice_reconciliation.normalize_supplier(inv.seller_name), ""
                ),
                "paymentAllocatedAmount": str(allocated),
                "invoiceId": inv.id,
                "invoiceKey": inv.invoice_key,
                "invoiceNumber": inv.invoice_number,
                "invoiceDate": inv.issue_date.strftime("%Y-%m-%d") if inv.issue_date else "",
                "invoiceType": inv.invoice_type,
                "invoiceAmountExclTax": str(_dec(inv.amount_excl_tax)),
                "invoiceTaxAmount": str(_dec(inv.tax_amount)),
                "invoiceTotalAmount": str(inv_total),
                "invoiceCorporatePaidTotal": str(paid_total),
                "invoiceOutstandingAmount": str(outstanding),
                "invoiceOverpaidAmount": str(overpaid_amount),
                "invoiceOverpaidSettledAmount": str(overpaid_settled_amount),
                "invoiceOverpaidUnsettledAmount": str(overpaid_unsettled_amount),
                "invoiceEffectiveAmount": str(effective_total),
                "invoiceStatus": status,
                "bankReconciliationStatus": status,
                "bankReconciliationApplicable": reconciliation_applicable,
                "bankReconciliationReason": reconciliation_reason,
                "invoiceColor": red_meta.get("invoiceColor", "unknown"),
                "invoiceStatusLabel": red_meta.get("invoiceStatusLabel", ""),
                "redStatus": red_meta.get("redStatus", "none"),
                "redPairStatus": red_meta.get("redPairStatus", "none"),
                "redRelatedInvoiceNo": red_meta.get("redRelatedInvoiceNo", ""),
                "redRelatedInvoicePeriod": red_meta.get("redRelatedInvoicePeriod", ""),
                "redCrossPeriod": bool(red_meta.get("redCrossPeriod")),
                "redOffsetAmount": red_meta.get("redOffsetAmount", "0.00"),
                "remainingAfterRedAmount": red_meta.get("remainingAfterRedAmount", "0.00"),
                "accountingNetAmount": red_meta.get("accountingNetAmount", "0.00"),
                "accountingException": red_meta.get("accountingException", ""),
                "redPairMethod": red_meta.get("redPairMethod", ""),
                **settlement_meta,
                **vat_meta,
                "manualPaymentMethod": payment_context["manualPaymentMethod"],
                "paymentMethod": payment_method,
                "paymentMethodLabel": payment_method_label,
                "paymentSource": payment_source,
                "paymentSourceLabel": payment_source_label,
                "paymentSourceBasis": payment_source_basis,
                "expenseNature": expense_nature,
                "expenseNatureLabel": expense_nature_label,
                "expenseNatureBasis": expense_nature_basis,
                "purchaseOrderNos": order_nos,
            })

        invoice_rows.append({
            "invoiceId": inv.id,
            "invoiceKey": inv.invoice_key,
            "invoiceNumber": inv.invoice_number or "",
            "invoiceDate": inv.issue_date.strftime("%Y-%m-%d") if inv.issue_date else "",
            "invoiceType": inv.invoice_type or "",
            "supplierName": inv.seller_name or "",
            "supplierTaxId": inv.seller_tax_id or supplier_tax.get(
                invoice_reconciliation.normalize_supplier(inv.seller_name), ""
            ),
            "invoiceAmountExclTax": str(_dec(inv.amount_excl_tax)),
            "invoiceTaxAmount": str(_dec(inv.tax_amount)),
            "invoiceTotalAmount": str(inv_total),
            "invoiceCorporatePaidTotal": str(paid_total),
            "invoiceOutstandingAmount": str(outstanding),
            "invoiceOverpaidAmount": str(overpaid_amount),
            "invoiceOverpaidSettledAmount": str(overpaid_settled_amount),
            "invoiceOverpaidUnsettledAmount": str(overpaid_unsettled_amount),
            "invoiceEffectiveAmount": str(effective_total),
            "invoiceStatus": status,
            "bankReconciliationStatus": status,
            "bankReconciliationApplicable": reconciliation_applicable,
            "bankReconciliationReason": reconciliation_reason,
            "invoiceColor": red_meta.get("invoiceColor", "unknown"),
            "invoiceStatusLabel": red_meta.get("invoiceStatusLabel", ""),
            "redStatus": red_meta.get("redStatus", "none"),
            "redPairStatus": red_meta.get("redPairStatus", "none"),
            "redRelatedInvoiceNo": red_meta.get("redRelatedInvoiceNo", ""),
            "redRelatedInvoicePeriod": red_meta.get("redRelatedInvoicePeriod", ""),
            "redCrossPeriod": bool(red_meta.get("redCrossPeriod")),
            "redOffsetAmount": red_meta.get("redOffsetAmount", "0.00"),
            "remainingAfterRedAmount": red_meta.get("remainingAfterRedAmount", "0.00"),
            "accountingNetAmount": red_meta.get("accountingNetAmount", "0.00"),
            "accountingException": red_meta.get("accountingException", ""),
            "redPairMethod": red_meta.get("redPairMethod", ""),
            **settlement_meta,
            **vat_meta,
            "manualPaymentMethod": payment_context["manualPaymentMethod"],
            "paymentMethod": payment_method,
            "paymentMethodLabel": payment_method_label,
            "paymentSource": payment_source,
            "paymentSourceLabel": payment_source_label,
            "paymentSourceBasis": payment_source_basis,
            "expenseNature": expense_nature,
            "expenseNatureLabel": expense_nature_label,
            "expenseNatureBasis": expense_nature_basis,
            "purchaseOrderNos": order_nos,
            "payments": payments,
        })

    unique_txns = {payment["paymentId"] for row in invoice_rows for payment in row["payments"]}
    payment_total = sum((_dec(txn_map[i].amount) for i in unique_txns if i in txn_map), Decimal("0"))
    invoice_total = sum((_dec(row["invoiceTotalAmount"]) for row in invoice_rows), Decimal("0"))
    accounting_net_total = sum((_dec(row.get("accountingNetAmount")) for row in invoice_rows), Decimal("0"))
    allocated_total = sum((_dec(row["invoiceCorporatePaidTotal"]) for row in invoice_rows), Decimal("0"))
    outstanding_total = sum((_dec(row["invoiceOutstandingAmount"]) for row in invoice_rows), Decimal("0"))

    report = {
        "year": year,
        "month": month,
        "company": company,
        "summary": {
            "paymentCount": len(unique_txns),
            "invoiceCount": len(invoice_rows),
            "paymentTotal": str(payment_total),
            "allocatedTotal": str(allocated_total),
            "invoiceTotal": str(invoice_total),
            "accountingNetTotal": str(accounting_net_total),
            "outstandingTotal": str(outstanding_total),
            "overpaidAfterRedCount": sum(1 for row in invoice_rows if row.get("bankReconciliationStatus") == "overpaid_after_red"),
            "overpaidAfterRedAmount": str(sum((_dec(row.get("invoiceOverpaidUnsettledAmount")) for row in invoice_rows), Decimal("0"))),
            "historicalOverpaidAfterRedAmount": str(sum((_dec(row.get("invoiceOverpaidAmount")) for row in invoice_rows), Decimal("0"))),
            "settledOverpaidAfterRedAmount": str(sum((_dec(row.get("invoiceOverpaidSettledAmount")) for row in invoice_rows), Decimal("0"))),
            "resolvedRedOverpaymentCount": sum(1 for row in invoice_rows if row.get("bankReconciliationStatus") == "red_overpayment_settled"),
            "redSettlementRemainingAmount": str(sum((_dec(row.get("redSettlementRemainingAmount")) for row in invoice_rows), Decimal("0"))),
            "companyMismatchInvoiceCount": company_mismatch["count"],
            "companyMismatchInvoiceAmount": company_mismatch["amount"],
            "companyMismatchInvoiceNos": company_mismatch["invoiceNos"],
            "redInvoiceCount": sum(1 for row in invoice_rows if row.get("invoiceColor") == "red"),
            "partiallyRedOffsetBlueCount": sum(
                1 for row in invoice_rows if row.get("redStatus") == "partially_red_offset"
            ),
            "fullyRedOffsetBlueCount": sum(
                1 for row in invoice_rows if row.get("redStatus") == "fully_red_offset"
            ),
            "redExceptionCount": sum(1 for row in invoice_rows if row.get("accountingException")),
            "crossPeriodRedCount": sum(1 for row in invoice_rows if row.get("redCrossPeriod")),
            "crossPeriodRedAmount": str(sum(
                (abs(_dec(row.get("invoiceTotalAmount"))) for row in invoice_rows if row.get("redCrossPeriod")),
                Decimal("0"),
            )),
            "paidInvoiceCount": sum(1 for row in invoice_rows if row["bankReconciliationStatus"] == "paid"),
            "partialInvoiceCount": sum(1 for row in invoice_rows if row["bankReconciliationStatus"] == "partial"),
            "unpaidInvoiceCount": sum(1 for row in invoice_rows if row["bankReconciliationStatus"] == "unpaid"),
            "notApplicableInvoiceCount": sum(
                1 for row in invoice_rows if row["bankReconciliationStatus"] == "not_applicable"
            ),
            "corporateInvoiceCount": sum(1 for row in invoice_rows if row["paymentSource"] == "corporate"),
            "personalInvoiceCount": sum(1 for row in invoice_rows if row["paymentSource"] == "personal"),
            "platformAutoDebitInvoiceCount": sum(1 for row in invoice_rows if row["paymentSource"] == "platform_auto_debit"),
            "mixedInvoiceCount": sum(1 for row in invoice_rows if row["paymentSource"] == "mixed"),
            "notApplicablePaymentCount": sum(
                1 for row in invoice_rows if row["paymentSource"] == "not_applicable"
            ),
            "personalInferredAmount": str(sum(
                (
                    _dec(row["invoiceOutstandingAmount"])
                    for row in invoice_rows
                    if row["paymentSource"] in {"personal", "mixed"}
                ),
                Decimal("0"),
            )),
            "expenseNatureCounts": dict(
                Counter(row["expenseNatureLabel"] for row in invoice_rows)
            ),
        },
        "invoiceRows": invoice_rows,
        "rows": flat_rows,
    }
    if selected_keys is not None:
        return apply_invoice_selection(report, selected_keys)
    return report


def latest_adjustment(
    db: Session, *, company: str, year: int, month: int
) -> FinanceCorporatePaymentAdjustment | None:
    return (
        db.query(FinanceCorporatePaymentAdjustment)
        .filter_by(company=company, period_year=year, period_month=month)
        .order_by(FinanceCorporatePaymentAdjustment.version.desc())
        .first()
    )


def _canonicalize_invoice_selection(
    available_rows: list[tuple[str, str]],
    selected_keys: list[str],
) -> list[str]:
    """把选择统一转换为 invoice_key；仅对唯一旧发票号码做兼容映射。"""
    canonical_keys: set[str] = set()
    keys_by_number: dict[str, list[str]] = defaultdict(list)
    for invoice_key, invoice_number in available_rows:
        canonical = str(invoice_key or "").strip()
        number = str(invoice_number or "").strip()
        if not canonical:
            continue
        canonical_keys.add(canonical)
        if number and canonical not in keys_by_number[number]:
            keys_by_number[number].append(canonical)

    requested = list(dict.fromkeys(
        str(key).strip() for key in selected_keys if str(key).strip()
    ))
    result: list[str] = []
    for key in requested:
        if key in canonical_keys:
            canonical = key
        else:
            legacy_matches = keys_by_number.get(key, [])
            if len(legacy_matches) == 1:
                canonical = legacy_matches[0]
            elif len(legacy_matches) > 1:
                raise ValueError(
                    f"历史选择使用的发票号码 {key} 当前对应多张发票，请刷新页面后重新选择"
                )
            else:
                raise ValueError("发票清单已发生变化，请刷新后重新选择")
        if canonical not in result:
            result.append(canonical)
    return result


def canonical_report_selection_keys(
    report: dict[str, Any], selected_keys: list[str]
) -> list[str]:
    """将报告选择转换为稳定 invoice_key，供页面、历史版本兼容与打包共同复用。"""
    available_rows = [
        (
            str(row.get("invoiceKey") or "").strip(),
            str(row.get("invoiceNumber") or "").strip(),
        )
        for row in report.get("invoiceRows", [])
    ]
    return _canonicalize_invoice_selection(available_rows, selected_keys)


def save_corporate_payment_adjustment(
    db: Session,
    *,
    company: str,
    year: int,
    month: int,
    selected_keys: list[str],
    actor: str = "system",
    note: str = "",
) -> dict[str, Any]:
    """保存本月保留的进项发票选择，并以新版本记录调整；历史版本不改写。"""
    if not 1 <= month <= 12:
        raise ValueError("非法月份")
    source_invoices = _period_invoices(db, year, month, company)
    normalized = _canonicalize_invoice_selection(
        [
            ((inv.invoice_key or "").strip(), (inv.invoice_number or "").strip())
            for inv in source_invoices
        ],
        selected_keys,
    )
    if source_invoices and not normalized:
        raise ValueError("至少保留一张发票")

    previous = (
        db.query(func.max(FinanceCorporatePaymentAdjustment.version))
        .filter_by(company=company, period_year=year, period_month=month)
        .scalar()
    )
    version = int(previous or 0) + 1
    row = FinanceCorporatePaymentAdjustment(
        company=company,
        period_year=year,
        period_month=month,
        version=version,
        selected_keys=normalized,
        actor=(actor or "system")[:64],
        note=(note or "")[:500],
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    audit(
        db,
        actor,
        "finance.corporate_payment.adjustment",
        "finance_corporate_payment_adjustments",
        row.id,
        {"company": company, "year": year, "month": month, "version": version, "selectedCount": len(normalized)},
    )
    return {
        "id": row.id,
        "version": row.version,
        "selectedKeys": normalized,
        "selectedCount": len(normalized),
        "sourceCount": len(source_invoices),
        "updatedAt": row.updated_at.isoformat() if row.updated_at else None,
    }


def apply_invoice_selection(report: dict[str, Any], selected_keys: list[str]) -> dict[str, Any]:
    """按稳定 invoice_key 过滤报告；兼容唯一旧发票号码，不对歧义号码做猜测。"""
    selected = set(canonical_report_selection_keys(report, selected_keys))
    invoice_rows = [
        row for row in report.get("invoiceRows", [])
        if (row.get("invoiceKey") or "").strip() in selected
    ]
    kept_ids = {row["invoiceId"] for row in invoice_rows}
    flat_rows = [row for row in report.get("rows", []) if row.get("invoiceId") in kept_ids]

    unique_txns = {payment["paymentId"] for row in invoice_rows for payment in row.get("payments", [])}
    payment_by_id = {
        payment["paymentId"]: _dec(payment.get("paymentAmount"))
        for row in report.get("invoiceRows", [])
        for payment in row.get("payments", [])
    }
    status_of = lambda row: row.get("bankReconciliationStatus") or row.get("invoiceStatus")
    summary = {
        "paymentCount": len(unique_txns),
        "invoiceCount": len(invoice_rows),
        "paymentTotal": str(sum((payment_by_id[i] for i in unique_txns), Decimal("0"))),
        "allocatedTotal": str(sum((_dec(row.get("invoiceCorporatePaidTotal")) for row in invoice_rows), Decimal("0"))),
        "invoiceTotal": str(sum((_dec(row.get("invoiceTotalAmount")) for row in invoice_rows), Decimal("0"))),
        "accountingNetTotal": str(sum((_dec(row.get("accountingNetAmount")) for row in invoice_rows), Decimal("0"))),
        "outstandingTotal": str(sum((_dec(row.get("invoiceOutstandingAmount")) for row in invoice_rows), Decimal("0"))),
        "overpaidAfterRedCount": sum(1 for row in invoice_rows if status_of(row) == "overpaid_after_red"),
        "overpaidAfterRedAmount": str(sum((_dec(row.get("invoiceOverpaidUnsettledAmount")) for row in invoice_rows), Decimal("0"))),
        "historicalOverpaidAfterRedAmount": str(sum((_dec(row.get("invoiceOverpaidAmount")) for row in invoice_rows), Decimal("0"))),
        "settledOverpaidAfterRedAmount": str(sum((_dec(row.get("invoiceOverpaidSettledAmount")) for row in invoice_rows), Decimal("0"))),
        "resolvedRedOverpaymentCount": sum(1 for row in invoice_rows if status_of(row) == "red_overpayment_settled"),
        "redSettlementRemainingAmount": str(sum((_dec(row.get("redSettlementRemainingAmount")) for row in invoice_rows), Decimal("0"))),
        "companyMismatchInvoiceCount": report.get("summary", {}).get("companyMismatchInvoiceCount", 0),
        "companyMismatchInvoiceAmount": report.get("summary", {}).get("companyMismatchInvoiceAmount", "0.00"),
        "companyMismatchInvoiceNos": report.get("summary", {}).get("companyMismatchInvoiceNos", []),
        "redInvoiceCount": sum(1 for row in invoice_rows if row.get("invoiceColor") == "red"),
        "partiallyRedOffsetBlueCount": sum(
            1 for row in invoice_rows if row.get("redStatus") == "partially_red_offset"
        ),
        "fullyRedOffsetBlueCount": sum(
            1 for row in invoice_rows if row.get("redStatus") == "fully_red_offset"
        ),
        "redExceptionCount": sum(1 for row in invoice_rows if row.get("accountingException")),
        "crossPeriodRedCount": sum(1 for row in invoice_rows if row.get("redCrossPeriod")),
        "crossPeriodRedAmount": str(sum(
            (abs(_dec(row.get("invoiceTotalAmount"))) for row in invoice_rows if row.get("redCrossPeriod")),
            Decimal("0"),
        )),
        "paidInvoiceCount": sum(1 for row in invoice_rows if status_of(row) == "paid"),
        "partialInvoiceCount": sum(1 for row in invoice_rows if status_of(row) == "partial"),
        "unpaidInvoiceCount": sum(1 for row in invoice_rows if status_of(row) == "unpaid"),
        "notApplicableInvoiceCount": sum(1 for row in invoice_rows if status_of(row) == "not_applicable"),
        "corporateInvoiceCount": sum(1 for row in invoice_rows if row.get("paymentSource") == "corporate"),
        "personalInvoiceCount": sum(1 for row in invoice_rows if row.get("paymentSource") == "personal"),
        "platformAutoDebitInvoiceCount": sum(1 for row in invoice_rows if row.get("paymentSource") == "platform_auto_debit"),
        "mixedInvoiceCount": sum(1 for row in invoice_rows if row.get("paymentSource") == "mixed"),
        "notApplicablePaymentCount": sum(
            1 for row in invoice_rows if row.get("paymentSource") == "not_applicable"
        ),
        "personalInferredAmount": str(sum(
            (
                _dec(row.get("invoiceOutstandingAmount"))
                for row in invoice_rows
                if row.get("paymentSource") in {"personal", "mixed"}
            ),
            Decimal("0"),
        )),
        "expenseNatureCounts": dict(
            Counter(str(row.get("expenseNatureLabel") or "待分类") for row in invoice_rows)
        ),
    }
    return {
        **report,
        "summary": summary,
        "invoiceRows": invoice_rows,
        "rows": flat_rows,
    }

def _style_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        width = max(
            [len(str(ws.cell(row=r, column=col).value or "")) for r in range(1, min(ws.max_row, 80) + 1)]
            or [10]
        )
        ws.column_dimensions[letter].width = min(max(width + 2, 10), 34)


def corporate_payment_xlsx(report: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    # 保留工作表名兼容历史交付包；列结构升级为财务主视图。
    ws.title = "已收票对公核对"
    headers = [
        "发票日期", "供应商", "供应商税号", "发票号码", "发票类型", "发票属性",
        "票据状态", "对应红/蓝发票", "原蓝票账期", "红冲金额", "财务净额",
        "红冲结算状态", "红冲待处理金额", "进项抵扣状态", "进项税转出状态", "进项税转出金额",
        "费用性质", "支付方式", "支付判断依据",
        "不含税金额", "税额", "价税合计",
        "对公匹配金额", "个人支付/未对公匹配金额",
        "银行付款日期", "我方付款账号", "银行流水/凭证号", "对公分摊金额",
        "关联采购订单", "银行匹配状态",
    ]
    ws.append(headers)

    for row in report.get("invoiceRows", []):
        payments = row.get("payments") or []
        bank_status = row.get("bankReconciliationStatus") or row.get("invoiceStatus")
        payment_source = row.get("paymentSourceLabel") or _PAYMENT_SOURCE_LABELS.get(
            row.get("paymentSource") or "", ""
        )
        ws.append([
            row["invoiceDate"],
            row["supplierName"],
            row["supplierTaxId"],
            row["invoiceNumber"],
            row["invoiceType"],
            (
                "红字发票"
                if row.get("invoiceColor") == "red"
                else "蓝字发票"
                if row.get("invoiceColor") == "blue"
                else "票据待确认"
            ),
            row.get("invoiceStatusLabel") or "",
            row.get("redRelatedInvoiceNo") or "",
            row.get("redRelatedInvoicePeriod") or "",
            float(_dec(row.get("redOffsetAmount"))),
            float(_dec(row.get("accountingNetAmount"))),
            row.get("redSettlementStatus") or "",
            float(_dec(row.get("redSettlementRemainingAmount"))),
            row.get("vatDeductibleStatus") or "",
            row.get("inputVatTransferStatus") or "",
            float(_dec(row.get("inputVatTransferAmount"))),
            row.get("expenseNatureLabel") or "待分类",
            payment_source,
            row.get("paymentSourceBasis") or "",
            float(_dec(row["invoiceAmountExclTax"])),
            float(_dec(row["invoiceTaxAmount"])),
            float(_dec(row["invoiceTotalAmount"])),
            float(_dec(row["invoiceCorporatePaidTotal"])),
            (
                None
                if bank_status == "not_applicable"
                else 0
                if row.get("paymentSource") == "platform_auto_debit"
                else float(_dec(row["invoiceOutstandingAmount"]))
            ),
            "、".join(str(payment.get("paymentDate") or "") for payment in payments),
            "、".join(
                str(payment.get("paymentAccount") or payment.get("paymentAccountName") or "")
                for payment in payments
            ),
            "、".join(str(payment.get("voucherNo") or "") for payment in payments),
            "、".join(str(_dec(payment.get("allocatedAmount"))) for payment in payments),
            "、".join(row.get("purchaseOrderNos") or []),
            (
                "已匹配"
                if bank_status == "paid"
                else "红冲后超额付款"
                if bank_status == "overpaid_after_red"
                else "部分匹配"
                if bank_status == "partial"
                else "不适用"
                if bank_status == "not_applicable"
                else "平台自动扣款货款"
                if row.get("paymentSource") == "platform_auto_debit"
                else "未匹配（按个人支付）"
            ),
        ])
    _style_sheet(ws)

    summary = wb.create_sheet("月度汇总", 0)
    s = report.get("summary", {})
    summary.append(["指标", "数值"])
    summary.append(["进项发票张数", s.get("invoiceCount", 0)])
    summary.append(["发票票面净合计", float(_dec(s.get("invoiceTotal")))])
    summary.append(["财务有效净额", float(_dec(s.get("accountingNetTotal", s.get("invoiceTotal"))))])
    summary.append(["红字发票张数", s.get("redInvoiceCount", 0)])
    summary.append(["部分红冲蓝字发票张数", s.get("partiallyRedOffsetBlueCount", 0)])
    summary.append(["已全额红冲蓝字发票张数", s.get("fullyRedOffsetBlueCount", 0)])
    summary.append(["红冲关联异常张数", s.get("redExceptionCount", 0)])
    summary.append(["跨期红字发票张数", s.get("crossPeriodRedCount", 0)])
    summary.append(["跨期红字冲减金额", float(_dec(s.get("crossPeriodRedAmount")))])
    summary.append(["红冲后仍待退款/冲抵金额", float(_dec(s.get("overpaidAfterRedAmount")))])
    summary.append(["历史红冲超额付款金额", float(_dec(s.get("historicalOverpaidAfterRedAmount")))])
    summary.append(["已处理红冲超额付款金额", float(_dec(s.get("settledOverpaidAfterRedAmount")))])
    summary.append(["红冲待退款/冲抵金额", float(_dec(s.get("redSettlementRemainingAmount")))])
    summary.append(["主体不符发票张数", s.get("companyMismatchInvoiceCount", 0)])
    summary.append(["主体不符发票金额", float(_dec(s.get("companyMismatchInvoiceAmount")))])
    summary.append(["对公支付已匹配金额", float(_dec(s.get("allocatedTotal")))])
    summary.append(["个人支付/未对公匹配金额", float(_dec(s.get("personalInferredAmount", s.get("outstandingTotal"))))])
    summary.append(["对公支付发票张数", s.get("corporateInvoiceCount", 0)])
    summary.append(["个人支付发票张数", s.get("personalInvoiceCount", 0)])
    summary.append(["平台自动扣款货款发票张数", s.get("platformAutoDebitInvoiceCount", 0)])
    summary.append(["混合支付发票张数", s.get("mixedInvoiceCount", 0)])
    summary.append(["支付方式不适用发票张数", s.get("notApplicablePaymentCount", 0)])
    summary.append(["关联银行付款笔数", s.get("paymentCount", 0)])
    for label, count in sorted((s.get("expenseNatureCounts") or {}).items()):
        summary.append([f"费用性质：{label}", count])
    _style_sheet(summary)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
