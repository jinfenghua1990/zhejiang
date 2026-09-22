from __future__ import annotations

from datetime import datetime
from decimal import Decimal

from sqlalchemy import select
from sqlalchemy.orm import Session
from io import BytesIO

from app.models.catalog import ProductSku, Warehouse
from app.models.consumable import Consumable, ConsumableSkuMapping, ConsumableTransaction, InboundConsumableUsage
from app.models.jackyun import JackyunGoodsDocument, JackyunGoodsDocumentItem
from app.models.production import ProductionOrder, ProductionOrderItem
from app.models.procurement_chain import ProcurementChainLink
from app.models.purchase import PurchaseAllocationItem
from app.models.tax import TaxAccountingCategoryRule
from app.services.consumable_ledger_service import record_transaction
from app.services.tax_category_rule_service import normalize_tax_code, unique_enabled_rule_for_tax_code
from app.utils.money import to_decimal


def _money(value: Decimal | None) -> str | None:
    return f"{to_decimal(value):.4f}" if value is not None else None


def _qty(value: Decimal | None) -> str:
    return f"{to_decimal(value):.4f}"


def _linked_skus(db: Session, consumable_id: int) -> list[dict]:
    rows = db.query(ConsumableSkuMapping, ProductSku).join(
        ProductSku, ProductSku.id == ConsumableSkuMapping.sku_id
    ).filter(ConsumableSkuMapping.consumable_id == consumable_id).order_by(ProductSku.sku_code).all()
    return [{"skuId": sku.id, "skuCode": sku.sku_code, "skuName": sku.sku_name} for _, sku in rows]


_OPEN_PRODUCTION_STATUSES = {
    "planned", "confirmed", "producing", "produced", "shipped", "arrived", "inbound",
}


def _coverage_metrics(available: Decimal, pending_consumable: Decimal, usage_per_unit: Decimal) -> tuple[Decimal | None, Decimal | None, Decimal]:
    """返回（可支撑正品数, 覆盖率%, 耗材缺口）；覆盖率始终使用同一耗材数量单位。"""
    available_nonnegative = max(available, Decimal("0"))
    support = available_nonnegative / usage_per_unit if usage_per_unit > 0 else None
    coverage = (
        available_nonnegative / pending_consumable * Decimal("100")
        if pending_consumable > 0 else None
    )
    gap = max(pending_consumable - available_nonnegative, Decimal("0"))
    return support, coverage, gap


def _inventory_context(db: Session, rows: list[Consumable]) -> dict[int, dict]:
    """批量准备耗材经营视图所需的映射、流水和待生产数据。"""
    row_ids = [row.id for row in rows]
    if not row_ids:
        return {}

    mapping_rows = db.query(ConsumableSkuMapping, ProductSku).join(
        ProductSku, ProductSku.id == ConsumableSkuMapping.sku_id
    ).filter(ConsumableSkuMapping.consumable_id.in_(row_ids)).order_by(
        ConsumableSkuMapping.consumable_id, ProductSku.sku_code,
    ).all()
    mappings_by_consumable: dict[int, list[dict]] = {row_id: [] for row_id in row_ids}
    usage_by_consumable: dict[int, list[Decimal]] = {row_id: [] for row_id in row_ids}
    for mapping, sku in mapping_rows:
        usage = to_decimal(mapping.usage_per_unit)
        mappings_by_consumable[mapping.consumable_id].append({
            "skuId": sku.id,
            "skuCode": sku.sku_code,
            "skuName": sku.sku_name,
            "usagePerUnit": _qty(usage),
        })
        if usage > 0:
            usage_by_consumable[mapping.consumable_id].append(usage)

    loss_by_consumable: dict[int, Decimal] = {row_id: Decimal("0") for row_id in row_ids}
    adjustment_by_consumable: dict[int, Decimal] = {row_id: Decimal("0") for row_id in row_ids}
    transactions = db.query(ConsumableTransaction).filter(
        ConsumableTransaction.consumable_id.in_(row_ids),
    ).all()
    for tx in transactions:
        qty = to_decimal(tx.quantity)
        if tx.transaction_type == "loss":
            loss_by_consumable[tx.consumable_id] += qty
        elif tx.transaction_type in {"stocktake", "adjustment", "manual"}:
            adjustment_by_consumable[tx.consumable_id] += qty

    pending_by_sku: dict[int, Decimal] = {}
    production_rows = db.query(ProductionOrderItem).join(
        ProductionOrder, ProductionOrder.id == ProductionOrderItem.production_order_id
    ).filter(
        ProductionOrder.status.in_(_OPEN_PRODUCTION_STATUSES),
        ProductionOrderItem.quantity > ProductionOrderItem.inbound_qty,
    ).all()
    for item in production_rows:
        remaining = max(to_decimal(item.quantity) - to_decimal(item.inbound_qty), Decimal("0"))
        if remaining > 0:
            pending_by_sku[item.sku_id] = pending_by_sku.get(item.sku_id, Decimal("0")) + remaining

    pending_by_consumable: dict[int, Decimal] = {row_id: Decimal("0") for row_id in row_ids}
    for mapping, _ in mapping_rows:
        pending_by_consumable[mapping.consumable_id] += (
            pending_by_sku.get(mapping.sku_id, Decimal("0")) * to_decimal(mapping.usage_per_unit)
        )

    context: dict[int, dict] = {}
    for row in rows:
        available = to_decimal(row.stock_qty) + to_decimal(row.factory_qty)
        usage_values = usage_by_consumable[row.id]
        usage_per_unit = max(usage_values) if usage_values else Decimal("0")
        pending = pending_by_consumable[row.id]
        support, coverage, gap = _coverage_metrics(available, pending, usage_per_unit)
        if coverage is None or coverage >= Decimal("100"):
            inventory_status = "正常"
        elif coverage >= Decimal("60"):
            inventory_status = "偏低"
        else:
            inventory_status = "缺货"
        context[row.id] = {
            "mappingDetails": mappings_by_consumable[row.id],
            "lossQty": loss_by_consumable[row.id],
            "adjustmentQty": adjustment_by_consumable[row.id],
            "pendingProductionQty": pending,
            "usagePerUnit": usage_per_unit if usage_per_unit > 0 else None,
            "supportQty": support,
            "coveragePct": coverage,
            "gapQty": gap,
            "inventoryStatus": inventory_status,
        }
    return context


def serialize_consumable(row: Consumable, db: Session, context: dict | None = None) -> dict:
    inventory = (context or {}).get(row.id, {})
    mapping_details = inventory.get("mappingDetails")
    if mapping_details is None:
        mapping_details = [
            {**item, "usagePerUnit": None}
            for item in _linked_skus(db, row.id)
        ]
    mapping_count = len(mapping_details)
    available = to_decimal(row.stock_qty) + to_decimal(row.factory_qty)
    min_qty = to_decimal(row.min_stock_qty)
    return {
        "id": row.id,
        "code": row.code,
        "name": row.name,
        "barcode": row.barcode or "",
        "category": row.category,
        "unit": row.unit,
        "purchaseUnitCost": _money(row.purchase_unit_cost),
        "purchasedQty": _qty(row.purchased_qty),
        "usedQty": _qty(row.used_qty),
        "stockQty": _qty(row.stock_qty),
        "factoryQty": _qty(row.factory_qty),
        "transitQty": _qty(row.transit_qty),
        "availableQty": _qty(available),
        "minStockQty": _qty(row.min_stock_qty),
        "taxCode": row.tax_code or "",
        "taxCategoryRuleId": row.tax_category_rule_id,
        "taxCategoryRuleName": _tax_rule_label(row.tax_category_rule_id, db),
        "usageRate": _qty((to_decimal(row.used_qty) / to_decimal(row.purchased_qty)) if to_decimal(row.purchased_qty) > 0 else Decimal("0")),
        "status": row.status,
        "mappingCount": mapping_count,
        "linkedSkus": [
            {key: item[key] for key in ("skuId", "skuCode", "skuName")}
            for item in mapping_details
        ],
        "mappingDetails": mapping_details,
        "totalStockQty": _qty(to_decimal(row.purchased_qty)),
        "totalUsedQty": _qty(to_decimal(row.used_qty)),
        "currentStockQty": _qty(available),
        "lossQty": _qty(inventory.get("lossQty", Decimal("0"))),
        "adjustmentQty": _qty(inventory.get("adjustmentQty", Decimal("0"))),
        "pendingProductionQty": _qty(inventory.get("pendingProductionQty", Decimal("0"))),
        "usagePerUnit": _qty(inventory["usagePerUnit"]) if inventory.get("usagePerUnit") is not None else None,
        "supportQty": _qty(inventory["supportQty"]) if inventory.get("supportQty") is not None else None,
        "coveragePct": _qty(inventory["coveragePct"]) if inventory.get("coveragePct") is not None else None,
        "gapQty": _qty(inventory["gapQty"]) if inventory.get("gapQty") is not None else None,
        "inventoryStatus": inventory.get("inventoryStatus", "正常"),
        # 预警口径：可用（各耗材仓汇总）≤ 安全库存，或任一库存口径为负
        "lowStock": (min_qty > 0 and available <= min_qty)
        or to_decimal(row.stock_qty) < 0
        or to_decimal(row.factory_qty) < 0,
    }


def _tax_rule_label(rule_id: int | None, db: Session) -> str:
    if rule_id is None:
        return ""
    row = db.get(TaxAccountingCategoryRule, rule_id)
    return f"{row.category_name} · {row.item_name}" if row else ""


def list_consumables(db: Session, search: str = "", status: str | None = None) -> list[dict]:
    query = db.query(Consumable).order_by(Consumable.code, Consumable.id)
    if status:
        query = query.filter(Consumable.status == status)
    term = search.strip()
    if term:
        pattern = f"%{term}%"
        query = query.filter((Consumable.code.ilike(pattern)) | (Consumable.name.ilike(pattern)))
    rows = query.limit(500).all()
    context = _inventory_context(db, rows)
    return [serialize_consumable(row, db, context) for row in rows]


def upsert_consumable(
    db: Session,
    *,
    consumable_id: int | None,
    code: str,
    name: str,
    category: str = "",
    unit: str = "个",
    purchase_unit_cost: str | None = None,
    stock_qty: str | None = None,
    min_stock_qty: str | None = None,
    barcode: str | None = None,
    tax_code: str | None = None,
    tax_category_rule_id: int | None = None,
    sku_ids: list[int] | None = None,
) -> Consumable:
    code = code.strip()
    if not code:
        raise ValueError("耗材编码不能为空")
    if stock_qty is not None:
        raise ValueError("库存请通过采购收货、领用或盘点流水调整")
    row = db.get(Consumable, consumable_id) if consumable_id else None
    if row is None:
        row = db.query(Consumable).filter_by(code=code).first()
    if row is None:
        row = Consumable(code=code)
        db.add(row)
    elif row.code != code and db.query(Consumable).filter_by(code=code).first():
        raise ValueError("耗材编码已存在，请换一个编码")
    row.code = code
    row.name = name.strip()
    row.category = category.strip()
    row.unit = unit.strip() or "个"
    affected_sku_ids: set[int] = set()
    if barcode is not None:
        row.barcode = barcode.strip()
    if purchase_unit_cost is not None:
        row.purchase_unit_cost = to_decimal(purchase_unit_cost)
    if min_stock_qty is not None:
        row.min_stock_qty = to_decimal(min_stock_qty)
    if tax_category_rule_id is not None:
        rule = db.get(TaxAccountingCategoryRule, tax_category_rule_id)
        if rule is None:
            raise ValueError("财务分类规则不存在")
        row.tax_category_rule_id = rule.id
        row.tax_code = rule.tax_code or ""
    elif tax_code is not None:
        row.tax_code = normalize_tax_code(tax_code)
        matching_rule = unique_enabled_rule_for_tax_code(db, row.tax_code)
        row.tax_category_rule_id = matching_rule.id if matching_rule is not None else None
    # 关联正品（多对多）：以本次提交为准同步映射；新建映射默认每个正品使用 1 个耗材。
    if sku_ids is not None:
        wanted: dict[int, None] = {}
        for sku_id in sku_ids:
            try:
                wanted[int(sku_id)] = None
            except (TypeError, ValueError):
                raise ValueError("关联正品格式不正确")
        for sku_id in wanted:
            if db.get(ProductSku, sku_id) is None:
                raise ValueError(f"正品 SKU {sku_id} 不存在")
        existing_mappings = {m.sku_id: m for m in db.query(ConsumableSkuMapping).filter_by(consumable_id=row.id).all()}
        affected_sku_ids.update(existing_mappings)
        affected_sku_ids.update(wanted)
        for sku_id in wanted:
            if sku_id not in existing_mappings:
                db.add(ConsumableSkuMapping(consumable_id=row.id, sku_id=sku_id, usage_per_unit=Decimal("1"), note="货品档案关联"))
        for sku_id, mapping in existing_mappings.items():
            if sku_id not in wanted:
                db.delete(mapping)
    db.commit()
    db.refresh(row)
    if affected_sku_ids:
        auto_apply_confirmed_inbound_usage(db, sku_ids=affected_sku_ids)
        db.refresh(row)
    return row


def suggest_inbound_usage(db: Session, link_id: int) -> dict:
    """按实际入库明细和 SKU 映射生成一条入库关联的耗材使用建议。

    优先使用 ``PurchaseAllocationItem.source_item_id``，因为同一张吉客云入库单
    可能被多个采购单拆分关联；只有单据没有拆分关联时，才允许直接使用入库明细。
    采购时间只作查询线索，不参与耗材扣减判断。
    """
    from app.services.inbound_allocation_seed import _po_from_link

    link = db.get(ProcurementChainLink, link_id)
    if link is None or link.target_type != "inbound":
        return {"ready": False, "items": [], "reason": "入库关联不存在"}

    item_rows = db.query(JackyunGoodsDocumentItem).filter(
        JackyunGoodsDocumentItem.document_id == link.target_id,
    ).order_by(JackyunGoodsDocumentItem.line_no).all()
    item_ids = {row.id for row in item_rows}
    if not item_ids:
        return {"ready": True, "items": [], "reason": "入库单没有明细"}

    po = _po_from_link(db, link)
    allocation_rows = []
    if po is not None:
        allocation_rows = db.query(PurchaseAllocationItem).filter(
            PurchaseAllocationItem.po_id == po.id,
            PurchaseAllocationItem.source_item_id.in_(item_ids),
        ).all()

    source_rows: list[tuple[int, Decimal, int | None]] = []
    basis = "allocation"
    if allocation_rows:
        source_rows = [
            (int(row.sku_id), to_decimal(row.quantity), row.source_item_id)
            for row in allocation_rows
            if row.sku_id is not None and row.quantity is not None and to_decimal(row.quantity) > 0
        ]
    else:
        active_link_count = db.query(ProcurementChainLink).filter(
            ProcurementChainLink.target_type == "inbound",
            ProcurementChainLink.target_id == link.target_id,
            ProcurementChainLink.match_method != "rejected",
            ProcurementChainLink.confirmed.is_(True),
        ).count()
        if active_link_count != 1:
            return {
                "ready": False,
                "items": [],
                "reason": "同一入库单关联多个采购单，尚无明细分摊",
            }
        basis = "single_inbound"
        source_rows = [
            (int(row.matched_sku_id), to_decimal(row.quantity), row.id)
            for row in item_rows
            if row.matched_sku_id is not None and row.quantity is not None and to_decimal(row.quantity) > 0
        ]

    sku_ids = {sku_id for sku_id, _, _ in source_rows}
    mappings = db.query(ConsumableSkuMapping).filter(
        ConsumableSkuMapping.sku_id.in_(sku_ids),
    ).all() if sku_ids else []
    mappings_by_sku: dict[int, list[ConsumableSkuMapping]] = {}
    for mapping in mappings:
        mappings_by_sku.setdefault(mapping.sku_id, []).append(mapping)

    totals: dict[int, Decimal] = {}
    source_item_ids: set[int] = set()
    for sku_id, quantity, source_item_id in source_rows:
        if source_item_id is not None:
            source_item_ids.add(source_item_id)
        for mapping in mappings_by_sku.get(sku_id, []):
            usage_quantity = quantity * to_decimal(mapping.usage_per_unit)
            if usage_quantity > 0:
                totals[mapping.consumable_id] = totals.get(mapping.consumable_id, Decimal("0")) + usage_quantity

    return {
        "ready": True,
        "basis": basis,
        "skuIds": sorted(sku_ids),
        "sourceItemIds": sorted(source_item_ids),
        "items": [
            {"consumable_id": consumable_id, "quantity": _qty(quantity)}
            for consumable_id, quantity in sorted(totals.items())
        ],
        "reason": "按入库明细和 SKU 耗材映射生成" if totals else "入库 SKU 尚未配置耗材映射",
    }


def auto_apply_inbound_usage(db: Session, link_id: int, note: str = "") -> dict:
    """把可确定的 SKU→耗材建议自动写入入库关联并扣减库存。

    正常采购入库不需要人工选择是否使用耗材：数量来自真实入库明细，耗材来自
    SKU 映射。只有历史上已经写入实际耗材流水的人工修正才保留；没有映射时
    保留为待处理，不能伪造“本次不使用”。
    """
    link = db.get(ProcurementChainLink, link_id)
    if link is None or link.target_type != "inbound":
        return {"status": "skipped", "reason": "入库关联不存在", "items": []}

    automatic_link = (
        link.match_method in {"auto", "file_import"}
        or "自动确认" in (link.note or "")
        or "自动按 SKU 耗材映射关联" in (link.note or "")
        or "按 SKU 耗材映射自动关联" in (link.note or "")
    )
    has_existing_usage = db.query(InboundConsumableUsage.id).filter_by(link_id=link_id).first() is not None
    manual_decided = link.match_method == "manual" and not automatic_link and has_existing_usage
    if manual_decided:
        return {"status": "manual_preserved", "reason": "保留人工明确的耗材登记结果", "items": []}

    suggestion = suggest_inbound_usage(db, link_id)
    if not suggestion["ready"]:
        return {"status": "pending", **suggestion}

    items = suggestion["items"]
    if items:
        usage_note = note.strip() or "按入库明细和 SKU 耗材映射自动关联"
        usage = set_inbound_usage(db, link_id=link_id, enabled=True, items=items, note=usage_note)
        link = db.get(ProcurementChainLink, link_id)
        if link is not None and "按 SKU 耗材映射自动关联" not in (link.note or ""):
            link.note = f"{link.note}；按 SKU 耗材映射自动关联" if link.note else "按 SKU 耗材映射自动关联"
            db.commit()
        return {"status": "applied", "basis": suggestion.get("basis"), "items": usage}

    if automatic_link or not link.consumable_usage_decided:
        # 未找到映射不是“不使用耗材”。保持未决状态，补齐货品档案映射后由
        # upsert_mapping -> auto_apply_confirmed_inbound_usage 自动补扣。
        link.consumable_usage_decided = False
        link.consumable_usage_enabled = None
        if "待维护耗材映射" not in (link.note or ""):
            link.note = f"{link.note}；待维护耗材映射" if link.note else "待维护耗材映射"
        db.commit()
    return {"status": "pending_mapping", "basis": suggestion.get("basis"), "items": []}


def auto_apply_confirmed_inbound_usage(
    db: Session,
    *,
    sku_ids: set[int] | None = None,
    note: str = "",
) -> dict:
    """在产品-耗材映射保存后，回溯补齐已有的真实入库关联。

    只扫描已确认的入库链；``suggest_inbound_usage`` 仍以入库明细分配行作为
    唯一数量来源，因此同一入库单拆给多个采购单时不会重复扣减。明确人工登记
    的结果不会被回溯覆盖。
    """
    link_ids = [link_id for (link_id,) in db.query(ProcurementChainLink.id).filter(
        ProcurementChainLink.target_type == "inbound",
        ProcurementChainLink.confirmed.is_(True),
        ProcurementChainLink.match_method != "rejected",
    ).all()]
    stats = {"applied": 0, "pending": 0, "noMapping": 0, "preserved": 0, "skipped": 0}
    for link_id in link_ids:
        try:
            suggestion = suggest_inbound_usage(db, link_id)
            if sku_ids is not None and not (set(suggestion.get("skuIds", [])) & sku_ids):
                continue
            result = auto_apply_inbound_usage(db, link_id, note=note)
            status = result.get("status")
            if status == "applied":
                stats["applied"] += 1
            elif status == "pending":
                stats["pending"] += 1
            elif status in {"no_mapping", "pending_mapping"}:
                stats["noMapping"] += 1
            elif status == "manual_preserved":
                stats["preserved"] += 1
            else:
                stats["skipped"] += 1
        except Exception:
            db.rollback()
            stats["skipped"] += 1
    return stats


def inbound_usage_rows(db: Session, link_id: int) -> list[dict]:
    """返回一条入库关联当前确认的耗材使用明细。"""
    rows = db.query(InboundConsumableUsage, Consumable).join(
        Consumable, Consumable.id == InboundConsumableUsage.consumable_id
    ).filter(InboundConsumableUsage.link_id == link_id).order_by(Consumable.code).all()
    return [
        {
            "id": usage.id,
            "consumableId": material.id,
            "consumableCode": material.code,
            "consumableName": material.name,
            "unit": material.unit,
            "quantity": _qty(usage.quantity),
            "note": usage.note,
        }
        for usage, material in rows
    ]


def _reverse_inbound_usage(db: Session, link_id: int) -> None:
    """撤销一条入库关联的耗材扣减，但保留原库存流水作为审计事实。"""
    from app.models.procurement_chain import ProcurementChainLink

    db.scalar(select(ProcurementChainLink).where(ProcurementChainLink.id == link_id).with_for_update())
    usages = db.query(InboundConsumableUsage).filter_by(
        link_id=link_id
    ).order_by(InboundConsumableUsage.consumable_id).all()
    for usage in usages:
        material = db.scalar(
            select(Consumable)
            .where(Consumable.id == usage.consumable_id)
            .with_for_update()
            .execution_options(populate_existing=True)
        )
        # 新流水以 usage.id 作为幂等来源；同时兼容历史 inbound_link/link_id 流水。
        tx = db.query(ConsumableTransaction).filter_by(
            source_type="inbound_usage",
            source_id=usage.id,
            consumable_id=usage.consumable_id,
        ).first()
        if tx is None:
            tx = db.query(ConsumableTransaction).filter_by(
                source_type="inbound_link",
                source_id=link_id,
                consumable_id=usage.consumable_id,
            ).first()

        tx_warehouse = db.get(Warehouse, tx.warehouse_id) if tx is not None and tx.warehouse_id else None
        is_factory = tx is not None and (
            tx.location == "factory"
            or (tx_warehouse is not None and tx_warehouse.warehouse_type == "factory")
        )
        quantity = to_decimal(usage.quantity)

        # 真实库存流水必须 append-only：保留原 consume，新增等量正向冲销流水。
        record_transaction(
            db,
            consumable_id=usage.consumable_id,
            transaction_type="manual",
            quantity=str(quantity),
            unit_cost=str(tx.unit_cost) if tx is not None and tx.unit_cost is not None else None,
            source_type="inbound_usage_reversal",
            source_id=usage.id,
            note=f"撤销入库关联 #{link_id} 耗材使用"
            + (f"；原流水 #{tx.id}" if tx is not None else ""),
            location="factory" if is_factory else "own",
            warehouse_id=tx.warehouse_id if tx is not None else None,
            commit=False,
        )
        if material is not None:
            material.used_qty = max(
                Decimal("0"),
                to_decimal(material.used_qty) - quantity,
            )
        db.delete(usage)


def set_inbound_usage(
    db: Session,
    *,
    link_id: int,
    enabled: bool,
    items: list[dict] | None = None,
    note: str = "",
) -> list[dict]:
    """写入一条入库关联的耗材流水，并按明细扣减库存。

    正常流程由 ``auto_apply_inbound_usage`` 调用；该接口仍保留给历史数据修正。
    ``enabled=False`` 仅兼容历史人工结果，不应作为正常正品入库的业务选择。
    enabled=True 必须至少有一条数量大于 0 的耗材明细。重复提交会先冲销旧明细
    再重算，保持幂等。
    """
    from app.models.procurement_chain import ProcurementChainLink
    from app.models.jackyun import JackyunGoodsDocument

    link = db.scalar(select(ProcurementChainLink).where(ProcurementChainLink.id == link_id).with_for_update().execution_options(populate_existing=True))
    if link is None or link.target_type != "inbound":
        raise ValueError("入库关联不存在")
    if db.get(JackyunGoodsDocument, link.target_id) is None:
        raise ValueError("关联的入库单不存在")
    raw_items = items or []
    if not isinstance(enabled, bool):
        raise ValueError("必须明确选择是否添加耗材使用")
    if not enabled and raw_items:
        raise ValueError("选择不使用耗材时不能填写耗材明细")
    normalized: dict[int, Decimal] = {}
    for item in raw_items:
        try:
            material_id = int(item.get("consumable_id"))
            quantity = to_decimal(item.get("quantity"))
        except (TypeError, ValueError, AttributeError) as exc:
            raise ValueError("耗材明细格式不正确") from exc
        if not quantity.is_finite() or quantity <= 0:
            raise ValueError("耗材使用数量必须大于 0")
        if db.get(Consumable, material_id) is None:
            raise ValueError(f"耗材 {material_id} 不存在")
        normalized[material_id] = normalized.get(material_id, Decimal("0")) + quantity
    if enabled and not normalized:
        raise ValueError("选择使用耗材时至少填写一条耗材及使用数量")

    old_items = db.query(InboundConsumableUsage).filter_by(link_id=link_id).all()
    old_quantities = {item.consumable_id: to_decimal(item.quantity) for item in old_items}
    if link.consumable_usage_decided and link.consumable_usage_enabled == enabled and old_quantities == normalized:
        return inbound_usage_rows(db, link_id)
    all_ids = sorted(set(old_quantities) | set(normalized))
    db.scalars(select(Consumable).where(Consumable.id.in_(all_ids)).order_by(Consumable.id).with_for_update().execution_options(populate_existing=True)).all()
    _reverse_inbound_usage(db, link_id)
    # 先落下旧明细删除，避免同一唯一键的新增先于删除触发冲突。
    db.flush()
    if enabled:
        for material_id, quantity in normalized.items():
            material = db.get(Consumable, material_id)
            assert material is not None
            # 优先使用唯一的启用耗材仓；只有仓库配置尚未形成唯一事实时，
            # 才回退到该耗材最近一次有仓库的实际收货。
            from app.models.consumable_purchase import ConsumablePurchaseItem, ConsumableReceipt
            from app.services import warehouse_service
            usage_warehouse = warehouse_service.default_for_consumable(db)
            # 测试库或尚未执行历史归一化的旧数据可能仍只有 stock_qty；在
            # 唯一工厂仓尚无余额时保留一次 own 兼容回退，避免把这批旧余额
            # 误判成工厂仓库存不足。归一化后的正式数据会优先走工厂仓。
            if (
                usage_warehouse is not None
                and usage_warehouse.warehouse_type == "factory"
                and to_decimal(material.factory_qty) == 0
                and to_decimal(material.stock_qty) > 0
            ):
                usage_warehouse = None
            if usage_warehouse is None:
                receipt_match = (
                    db.query(ConsumableReceipt, Warehouse)
                    .join(ConsumablePurchaseItem, ConsumablePurchaseItem.purchase_id == ConsumableReceipt.purchase_id)
                    .join(Warehouse, Warehouse.id == ConsumableReceipt.warehouse_id)
                    .filter(
                        ConsumablePurchaseItem.consumable_id == material_id,
                        ConsumableReceipt.warehouse_id.isnot(None),
                        Warehouse.status == "active",
                    )
                    .order_by(ConsumableReceipt.received_on.desc(), ConsumableReceipt.id.desc())
                    .first()
                )
                usage_warehouse = receipt_match[1] if receipt_match else None
            usage_location = "factory" if usage_warehouse is not None and usage_warehouse.warehouse_type == "factory" else "own"
            usage = InboundConsumableUsage(
                link_id=link_id,
                inbound_document_id=link.target_id,
                consumable_id=material_id,
                quantity=quantity,
                note=note.strip(),
            )
            db.add(usage)
            db.flush()
            record_transaction(
                db,
                consumable_id=material_id,
                transaction_type="consume",
                quantity=str(quantity),
                unit_cost=str(material.purchase_unit_cost) if material.purchase_unit_cost is not None else None,
                source_type="inbound_usage",
                source_id=usage.id,
                note=note.strip() or f"入库单关联 #{link.target_id} 确认耗材出库",
                location=usage_location,
                warehouse_id=usage_warehouse.id if usage_warehouse is not None else None,
                commit=False,
            )
    link.consumable_usage_decided = True
    link.consumable_usage_enabled = enabled
    db.commit()
    return inbound_usage_rows(db, link_id)


def list_transactions(db: Session, consumable_id: int, limit: int = 100) -> list[dict]:
    from app.models.consumable_purchase import ConsumableReceipt
    from app.models.procurement_chain import ProcurementChainLink
    from app.services import warehouse_service
    rows = db.query(ConsumableTransaction).filter_by(consumable_id=consumable_id).order_by(
        ConsumableTransaction.occurred_at.desc().nullslast(), ConsumableTransaction.id.desc()
    ).limit(limit).all()
    receipt_ids = [row.source_id for row in rows if row.source_type == "consumable_receipt"]
    receipts = {r.id: r for r in db.query(ConsumableReceipt).filter(ConsumableReceipt.id.in_(receipt_ids)).all()} if receipt_ids else {}
    legacy_link_ids = [row.source_id for row in rows if row.source_type == "inbound_link"]
    usage_ids = [row.source_id for row in rows if row.source_type == "inbound_usage"]
    usages = {
        usage.id: usage
        for usage in db.query(InboundConsumableUsage).filter(InboundConsumableUsage.id.in_(usage_ids)).all()
    } if usage_ids else {}
    usage_link_ids = [usage.link_id for usage in usages.values()]
    link_ids = sorted(set(legacy_link_ids + usage_link_ids))
    links = {
        link.id: link
        for link in db.query(ProcurementChainLink).filter(ProcurementChainLink.id.in_(link_ids)).all()
    } if link_ids else {}
    inbound_document_ids = {link.target_id for link in links.values() if link.target_id is not None}
    inbound_documents = {
        document.id: document
        for document in db.query(JackyunGoodsDocument).filter(JackyunGoodsDocument.id.in_(inbound_document_ids)).all()
    } if inbound_document_ids else {}
    warehouse_ids = {row.warehouse_id for row in rows if row.warehouse_id is not None}
    warehouses = {
        warehouse.id: warehouse
        for warehouse in db.query(Warehouse).filter(Warehouse.id.in_(warehouse_ids)).all()
    } if warehouse_ids else {}

    def _source_link(row: ConsumableTransaction):
        if row.source_type == "inbound_link" and row.source_id in links:
            return links[row.source_id]
        if row.source_type == "inbound_usage" and row.source_id in usages:
            return links.get(usages[row.source_id].link_id)
        return None

    return [
        {
            "id": row.id,
            "transactionType": row.transaction_type,
            "quantity": _qty(row.quantity),
            "unitCost": _money(row.unit_cost),
            "location": row.location,
            "warehouseId": row.warehouse_id,
            "warehouseCode": warehouse_service.display_code(warehouses[row.warehouse_id]) if row.warehouse_id in warehouses else "",
            "warehouseName": warehouses[row.warehouse_id].name if row.warehouse_id in warehouses else "",
            "stockBefore": _qty(row.stock_before) if row.stock_before is not None else None,
            "stockAfter": _qty(row.stock_after) if row.stock_after is not None else None,
            "factoryBefore": _qty(row.factory_before) if row.factory_before is not None else None,
            "factoryAfter": _qty(row.factory_after) if row.factory_after is not None else None,
            "sourceType": row.source_type,
            "sourceId": row.source_id,
            "note": row.note,
            "occurredAt": row.occurred_at.isoformat() if row.occurred_at else None,
            "purchaseId": receipts[row.source_id].purchase_id if row.source_type == "consumable_receipt" and row.source_id in receipts else None,
            "orderId": _source_link(row).workbench_order_id if _source_link(row) is not None else None,
            "inboundDocumentId": (
                inbound_documents[_source_link(row).target_id].id
                if _source_link(row) is not None
                and _source_link(row).target_id in inbound_documents
                else None
            ),
            "inboundDocumentNo": (
                inbound_documents[_source_link(row).target_id].goodsdoc_no
                if _source_link(row) is not None
                and _source_link(row).target_id in inbound_documents
                else ""
            ),
            "inboundDocumentAt": (
                inbound_documents[_source_link(row).target_id].document_at.isoformat()
                if _source_link(row) is not None
                and _source_link(row).target_id in inbound_documents
                and inbound_documents[_source_link(row).target_id].document_at is not None
                else None
            ),
            "inboundWarehouseName": (
                inbound_documents[_source_link(row).target_id].warehouse_name
                if _source_link(row) is not None
                and _source_link(row).target_id in inbound_documents
                else ""
            ),
        }
        for row in rows
    ]


def list_mappings(db: Session, consumable_id: int | None = None, sku_id: int | None = None) -> list[dict]:
    query = db.query(ConsumableSkuMapping, Consumable, ProductSku).join(
        Consumable, Consumable.id == ConsumableSkuMapping.consumable_id
    ).join(ProductSku, ProductSku.id == ConsumableSkuMapping.sku_id)
    if consumable_id is not None:
        query = query.filter(ConsumableSkuMapping.consumable_id == consumable_id)
    if sku_id is not None:
        query = query.filter(ConsumableSkuMapping.sku_id == sku_id)
    return [
        {
            "id": mapping.id,
            "skuId": sku.id,
            "skuCode": sku.sku_code,
            "skuName": sku.sku_name,
            "consumableId": consumable.id,
            "consumableCode": consumable.code,
            "consumableName": consumable.name,
            "usagePerUnit": _qty(mapping.usage_per_unit),
            "note": mapping.note,
        }
        for mapping, consumable, sku in query.order_by(Consumable.code, ProductSku.sku_code).all()
    ]


def upsert_mapping(db: Session, *, mapping_id: int | None, sku_id: int, consumable_id: int, usage_per_unit: str, note: str = "") -> ConsumableSkuMapping:
    if db.get(ProductSku, sku_id) is None or db.get(Consumable, consumable_id) is None:
        raise ValueError("货品或耗材不存在")
    qty = to_decimal(usage_per_unit)
    if qty <= 0:
        raise ValueError("单位消耗量必须大于 0")
    row = db.get(ConsumableSkuMapping, mapping_id) if mapping_id else db.query(ConsumableSkuMapping).filter_by(
        sku_id=sku_id, consumable_id=consumable_id
    ).first()
    if row is None:
        row = ConsumableSkuMapping(sku_id=sku_id, consumable_id=consumable_id)
        db.add(row)
    row.usage_per_unit = qty
    row.note = note.strip()
    db.commit()
    db.refresh(row)
    auto_apply_confirmed_inbound_usage(db, sku_ids={sku_id})
    db.refresh(row)
    return row


def remove_mapping(db: Session, mapping_id: int) -> None:
    row = db.get(ConsumableSkuMapping, mapping_id)
    if row is None:
        raise ValueError("耗材映射不存在")
    sku_id = row.sku_id
    db.delete(row)
    db.commit()
    auto_apply_confirmed_inbound_usage(db, sku_ids={sku_id})


def import_xlsx(db: Session, content: bytes, filename: str = "consumables.xlsx") -> dict[str, int | str]:
    """读取历史表格的“耗材使用情况”页；只导入可识别的值，不执行原表公式。"""
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("耗材清单仅支持 .xlsx 或 .xlsm")
    try:
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"耗材清单无法读取：{exc}") from exc
    sheet = book["耗材使用情况"] if "耗材使用情况" in book.sheetnames else None
    if sheet is None:
        raise ValueError("未找到“耗材使用情况”工作表")
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("耗材工作表为空")
    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    required = {"条形码", "产品名", "外包装采购数量", "采购单价"}
    if not required.issubset(headers):
        raise ValueError(f"耗材表缺少字段：{', '.join(sorted(required - set(headers)))}")
    idx = {name: headers.index(name) for name in headers if name}
    created = updated = skipped = mappings = 0
    mapping_cache: dict[tuple[int, int], ConsumableSkuMapping] = {}
    affected_sku_ids: set[int] = set()
    def json_value(value):
        if isinstance(value, (datetime,)):
            return value.isoformat()
        if isinstance(value, (Decimal, float, int, str, bool)) or value is None:
            return str(value) if isinstance(value, Decimal) else value
        return str(value)
    for values in rows[1:]:
        code = str(values[idx["条形码"]] or "").strip()
        name = str(values[idx["产品名"]] or "").strip()
        if not code or not name:
            skipped += 1
            continue
        def cell(name: str):
            return values[idx[name]] if name in idx and idx[name] < len(values) else None
        cost = cell("采购单价")
        purchased = cell("外包装采购数量")
        used = cell("使用耗材")
        stock = cell("剩余数量")
        category = str(cell("分类") or "外包装").strip()
        existing = db.query(Consumable).filter_by(code=code).first()
        row = existing or Consumable(code=code)
        row.name = name
        row.category = category
        row.unit = "个"
        if cost is not None and str(cost).strip() not in {"", "None"}:
            row.purchase_unit_cost = to_decimal(str(cost))
        if purchased is not None and isinstance(purchased, (int, float, Decimal)):
            row.purchased_qty = to_decimal(str(purchased))
        if used is not None and isinstance(used, (int, float, Decimal)):
            row.used_qty = to_decimal(str(used))
        if stock is not None and isinstance(stock, (int, float, Decimal)):
            row.stock_qty = to_decimal(str(stock))
        row.raw = {"source": filename, "sheet": "耗材使用情况", "row": {headers[i]: json_value(values[i]) for i in range(min(len(headers), len(values)))}}
        db.add(row)
        if existing: updated += 1
        else: created += 1
    db.flush()
    # “订货”页中条形码、耗材代码、耗材使用量和订购件数-盒形成明确映射；
    # 只有能命中现有吉客云 SKU 的行才自动建立，其他行留给界面人工维护。
    order_sheet = book["订货"] if "订货" in book.sheetnames else None
    if order_sheet is not None:
        order_rows = list(order_sheet.iter_rows(values_only=True))
        if order_rows:
            order_headers = [str(v).strip() if v is not None else "" for v in order_rows[0]]
            oi = {name: order_headers.index(name) for name in order_headers if name}
            for values in order_rows[1:]:
                def oval(name: str):
                    return values[oi[name]] if name in oi and oi[name] < len(values) else None
                sku_code = str(oval("条形码") or "").strip()
                consumable_code = str(oval("耗材代码") or "").strip()
                order_qty = oval("订购件数-盒")
                usage_qty = oval("耗材使用量")
                # 组合装行描述的是一次性拆包关系，不代表单品的固定耗材映射，
                # 不把它推断成单品单位消耗量，避免同一 SKU 出现 1/10 等冲突。
                if "组合装" in str(oval("备注") or ""):
                    continue
                if not sku_code or not consumable_code or order_qty is None or usage_qty is None:
                    continue
                try:
                    ratio = to_decimal(str(usage_qty)) / to_decimal(str(order_qty))
                except (TypeError, ValueError, ArithmeticError):
                    continue
                if ratio <= 0:
                    continue
                sku = db.query(ProductSku).filter((ProductSku.sku_code == sku_code) | (ProductSku.barcode == sku_code)).first()
                material = db.query(Consumable).filter_by(code=consumable_code).first()
                if sku is None or material is None:
                    continue
                key = (sku.id, material.id)
                mapping = mapping_cache.get(key) or db.query(ConsumableSkuMapping).filter_by(sku_id=sku.id, consumable_id=material.id).first()
                if mapping is None:
                    mapping = ConsumableSkuMapping(sku_id=sku.id, consumable_id=material.id, usage_per_unit=ratio, note="由历史订货表自动映射")
                    db.add(mapping)
                    mappings += 1
                    affected_sku_ids.add(sku.id)
                elif to_decimal(mapping.usage_per_unit) != ratio:
                    mapping.usage_per_unit = ratio
                    mapping.note = "由历史订货表更新映射"
                    mappings += 1
                    affected_sku_ids.add(sku.id)
                mapping_cache[key] = mapping
    db.commit()
    if affected_sku_ids:
        auto_apply_confirmed_inbound_usage(db, sku_ids=affected_sku_ids)
    return {"created": created, "updated": updated, "skipped": skipped, "mappings": mappings, "sheet": "耗材使用情况"}
