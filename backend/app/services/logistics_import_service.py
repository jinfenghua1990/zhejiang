"""物流 Excel 账单导入：识别仓储/快递账单并生成可核销账单。

当前支持用户提供的“汇总 / 发货明细 / 自提 / 增值 / 报价 / 顺丰报价”格式。
导入过程先预览，确认后再写入 logistics_bills；运单级明细保存在 LogisticsBill.raw，
避免为半年一次的外部账单额外引入高频业务表和迁移风险。
"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from hashlib import sha256
from io import BytesIO
import math
import statistics
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy.orm import Session

from app.models.jackyun import JackyunGoodsDocument, JackyunShopOrder
from app.models.logistics import LogisticsBill
from app.services import logistics_service

_REQUIRED_SHEETS = {"汇总", "发货明细"}
_MONEY = Decimal("0.01")


def _d(value: Any, default: Decimal = Decimal("0")) -> Decimal:
    if value in (None, ""):
        return default
    try:
        return Decimal(str(value)).quantize(_MONEY, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError, TypeError):
        return default


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _dt(value: Any, epoch) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch)
            if isinstance(converted, datetime):
                return converted
            if isinstance(converted, date):
                return datetime.combine(converted, time.min)
        except Exception:
            return None
    if isinstance(value, str):
        raw = value.strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d"):
            try:
                return datetime.strptime(raw, fmt)
            except ValueError:
                pass
    return None


def _money_str(value: Decimal) -> str:
    return str(value.quantize(_MONEY, rounding=ROUND_HALF_UP))


def _weight_band(weight: Decimal) -> str:
    if weight <= Decimal("0.5"):
        return "0-0.5kg"
    if weight <= Decimal("1"):
        return "0.51-1kg"
    if weight <= Decimal("2"):
        return "1.01-2kg"
    if weight <= Decimal("3"):
        return "2.01-3kg"
    return "3kg以上"


def _sheet_headers(ws, row: int = 1) -> dict[str, int]:
    out: dict[str, int] = {}
    for idx, cell in enumerate(ws[row], start=1):
        name = _text(cell.value)
        if name:
            out[name] = idx
    return out


def _parse_regular_quote(ws) -> dict[str, list[Any]]:
    result: dict[str, list[Any]] = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        province = _text(row[1] if len(row) > 1 else None)
        if not province:
            continue
        if province.startswith(("1、", "2、", "3、", "4、", "5、", "6、", "7、")):
            continue
        cells = list(row[2:7])
        if len(cells) < 5 or cells[0] is None:
            continue
        result.setdefault(province, cells)
    return result


def _parse_sf_quote(ws) -> dict[str, tuple[Decimal, Decimal, Decimal]]:
    result: dict[str, tuple[Decimal, Decimal, Decimal]] = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if len(row) < 14:
            continue
        province = _text(row[10])
        if not province:
            continue
        first, mid, high = _d(row[11]), _d(row[12]), _d(row[13])
        if first > 0:
            result[province] = (first, mid, high)
    return result


def _regular_fee(province: str, weight: Decimal, quote: dict[str, list[Any]]) -> Decimal | None:
    row = quote.get(province)
    if not row:
        return None
    if weight <= Decimal("0.5"):
        return _d(row[0])
    if weight <= Decimal("1"):
        return _d(row[1])
    if weight <= Decimal("2"):
        return _d(row[2])
    if weight <= Decimal("3"):
        return _d(row[3])

    expr = _text(row[4]).replace(" ", "").upper()
    if expr.startswith("5+W*"):
        try:
            rate = Decimal(expr.split("*", 1)[1])
            return (Decimal("5") + weight * rate).quantize(_MONEY, rounding=ROUND_HALF_UP)
        except Exception:
            return None
    if expr.endswith("*W"):
        try:
            rate = Decimal(expr[:-2])
            return (weight * rate).quantize(_MONEY, rounding=ROUND_HALF_UP)
        except Exception:
            return None
    fallback = _d(row[4], default=Decimal("-1")) if row[4] is not None else Decimal("-1")
    return fallback if fallback > 0 else None


def _sf_fee(province: str, weight: Decimal, quote: dict[str, tuple[Decimal, Decimal, Decimal]]) -> Decimal | None:
    row = quote.get(province)
    if not row:
        return None
    first, mid, high = row
    billed = max(Decimal("1"), Decimal(math.ceil(float(weight))))
    if billed <= 1:
        return first
    if billed <= 3:
        return (first + (billed - 1) * mid).quantize(_MONEY, rounding=ROUND_HALF_UP)
    return (first + Decimal("2") * mid + (billed - 3) * high).quantize(_MONEY, rounding=ROUND_HALF_UP)


def _charge_for(carrier: str, province: str, billed_weight: Decimal, regular_quote, sf_quote) -> Decimal | None:
    if "顺丰" in carrier:
        return _sf_fee(province, billed_weight, sf_quote)
    return _regular_fee(province, billed_weight, regular_quote)


def _parse_shipments(ws, regular_quote, sf_quote, epoch) -> list[dict[str, Any]]:
    headers = _sheet_headers(ws)
    required = ["包裹号", "物流公司", "物流单号", "发货单号", "仓库", "省份", "完成时间", "计费重量"]
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError("发货明细缺少字段：" + "、".join(missing))

    def get(row, name):
        idx = headers.get(name)
        return row[idx - 1] if idx and idx <= len(row) else None

    out: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        package_no = _text(get(row, "包裹号"))
        if not package_no:
            continue
        carrier = _text(get(row, "物流公司"))
        province = _text(get(row, "省份"))
        billed_weight = _d(get(row, "计费重量"))
        source_fee = get(row, "运费")
        fee = _d(source_fee, default=Decimal("-1"))
        if fee < 0:
            fee = _charge_for(carrier, province, billed_weight, regular_quote, sf_quote) or Decimal("0")
        completed_at = _dt(get(row, "完成时间"), epoch)
        out.append(
            {
                "kind": "shipment",
                "packageNo": package_no,
                "carrier": carrier,
                "trackingNo": _text(get(row, "物流单号")),
                "goodsdocNo": _text(get(row, "发货单号")),
                "warehouse": _text(get(row, "仓库")),
                "goodsSummary": _text(get(row, "货品摘要")),
                "quantity": str(_d(get(row, "数量合计"))),
                "province": province,
                "channel": _text(get(row, "销售渠道")),
                "sourceOrderNo": _text(get(row, "原始单号")),
                "completedAt": completed_at.isoformat() if completed_at else None,
                "weight": str(_d(get(row, "快递重量"))),
                "billedWeight": str(billed_weight),
                "weightBand": _weight_band(billed_weight),
                "fee": _money_str(fee),
            }
        )
    return out


def _parse_pickup(ws, epoch) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for row in ws.iter_rows(values_only=True):
        package_no = _text(row[0] if len(row) > 0 else None)
        if not package_no or package_no == "包裹号":
            continue
        if len(row) < 14:
            continue
        completed_at = _dt(row[12], epoch)
        fee = _d(row[13])
        out.append(
            {
                "kind": "pickup",
                "packageNo": package_no,
                "carrier": _text(row[1]),
                "trackingNo": _text(row[2]),
                "goodsdocNo": _text(row[3]),
                "warehouse": _text(row[5]),
                "goodsSummary": _text(row[7]),
                "quantity": str(_d(row[8])),
                "province": _text(row[9]),
                "channel": _text(row[10]),
                "sourceOrderNo": _text(row[11]),
                "completedAt": completed_at.isoformat() if completed_at else None,
                "weight": None,
                "billedWeight": None,
                "weightBand": "客户账号/自提",
                "fee": _money_str(fee),
            }
        )
    return out


def _parse_value_added(ws) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    month_label = ""
    for row in ws.iter_rows(values_only=True):
        first = _text(row[0] if len(row) > 0 else None)
        if "月份" in first and "增值服务" in first:
            month_label = first
            continue
        description = _text(row[3] if len(row) > 3 else None)
        amount = _d(row[8] if len(row) > 8 else None)
        if description and amount != 0:
            out.append({"monthLabel": month_label, "description": description, "amount": _money_str(amount)})
    return out


def _summary_amounts(ws) -> tuple[Decimal, list[dict[str, Any]], Decimal | None, str]:
    gross = _d(ws["E2"].value)
    warehouse = _text(ws["A4"].value)
    adjustments: list[dict[str, Any]] = []
    for row in range(11, min(ws.max_row, 30) + 1):
        label = _text(ws.cell(row, 4).value)
        amount = _d(ws.cell(row, 5).value)
        if label and amount != 0:
            adjustments.append({"label": label, "amount": _money_str(-amount)})
    final_value = ws.cell(ws.max_row, 5).value
    final_amount = _d(final_value, default=Decimal("-1"))
    if final_amount < 0:
        final_amount = None
    return gross, adjustments, final_amount, warehouse


def _match_shipments(db: Session, rows: list[dict[str, Any]]) -> dict[str, int]:
    goodsdoc_nos = {row["goodsdocNo"] for row in rows if row.get("goodsdocNo")}
    tracking_nos = {row["trackingNo"] for row in rows if row.get("trackingNo")}
    source_nos = {row["sourceOrderNo"] for row in rows if row.get("sourceOrderNo")}

    matched_goodsdocs = set()
    if goodsdoc_nos:
        matched_goodsdocs = {
            no for (no,) in db.query(JackyunGoodsDocument.goodsdoc_no)
            .filter(
                JackyunGoodsDocument.document_type == "outbound",
                JackyunGoodsDocument.goodsdoc_no.in_(goodsdoc_nos),
            ).all()
        }

    matched_tracking = set()
    if tracking_nos:
        matched_tracking = {
            no for (no,) in db.query(JackyunShopOrder.logistic_no)
            .filter(JackyunShopOrder.logistic_no.in_(tracking_nos)).all()
            if no
        }

    matched_source = set()
    if source_nos:
        matched_source = {
            no for (no,) in db.query(JackyunShopOrder.source_trade_no)
            .filter(JackyunShopOrder.source_trade_no.in_(source_nos)).all()
            if no
        }

    for row in rows:
        match_type = ""
        if row.get("goodsdocNo") in matched_goodsdocs:
            match_type = "goodsdoc"
        elif row.get("trackingNo") in matched_tracking:
            match_type = "tracking"
        elif row.get("sourceOrderNo") in matched_source:
            match_type = "source_order"
        row["matchType"] = match_type
        row["matched"] = bool(match_type)

    tracking_counter = Counter(row.get("trackingNo") for row in rows if row.get("trackingNo"))
    duplicates = sum(count - 1 for count in tracking_counter.values() if count > 1)
    matched = sum(1 for row in rows if row.get("matched"))
    return {
        "matched": matched,
        "unmatched": len(rows) - matched,
        "duplicates": duplicates,
    }


def _regional_preview(rows: list[dict[str, Any]], net_amount: Decimal, direct_amount: Decimal) -> list[dict[str, Any]]:
    factor = (net_amount / direct_amount) if direct_amount > 0 else Decimal("1")
    grouped: dict[tuple[str, str, str, str], list[Decimal]] = defaultdict(list)
    for row in rows:
        province = _text(row.get("province"))
        completed = _text(row.get("completedAt"))
        carrier = _text(row.get("carrier"))
        weight_band = _text(row.get("weightBand")) or "未知"
        if not province or not completed:
            continue
        month = completed[:7]
        fee = _d(row.get("fee")) * factor
        grouped[(month, province, carrier, weight_band)].append(fee)

    result: list[dict[str, Any]] = []
    for (month, province, carrier, weight_band), fees in grouped.items():
        avg = sum(fees, Decimal("0")) / Decimal(len(fees))
        med = Decimal(str(statistics.median([float(x) for x in fees])))
        count = len(fees)
        result.append(
            {
                "month": month,
                "province": province,
                "carrier": carrier,
                "weightBand": weight_band,
                "sampleCount": count,
                "avgFee": _money_str(avg),
                "medianFee": _money_str(med),
                "confidence": "high" if count >= 30 else "medium" if count >= 10 else "low",
            }
        )
    result.sort(key=lambda row: (row["month"], row["sampleCount"]), reverse=True)
    return result[:60]


def parse_bill_xlsx(db: Session, content: bytes, filename: str, *, persist: bool = False) -> dict[str, Any]:
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("只支持 .xlsx / .xlsm 物流账单")
    digest = sha256(content).hexdigest()
    for bill in db.query(LogisticsBill).order_by(LogisticsBill.id.desc()).limit(100).all():
        if isinstance(bill.raw, dict) and bill.raw.get("fileHash") == digest:
            if persist:
                raise ValueError(f"该物流账单已导入（账单 #{bill.id}）")
            duplicate_bill_id = bill.id
            break
    else:
        duplicate_bill_id = None

    wb = load_workbook(BytesIO(content), data_only=True, read_only=False)
    missing_sheets = sorted(_REQUIRED_SHEETS - set(wb.sheetnames))
    if missing_sheets:
        raise ValueError("物流账单缺少工作表：" + "、".join(missing_sheets))

    regular_quote = _parse_regular_quote(wb["报价"]) if "报价" in wb.sheetnames else {}
    sf_quote = _parse_sf_quote(wb["顺丰报价"]) if "顺丰报价" in wb.sheetnames else {}
    shipments = _parse_shipments(wb["发货明细"], regular_quote, sf_quote, wb.epoch)
    pickup = _parse_pickup(wb["自提"], wb.epoch) if "自提" in wb.sheetnames else []
    value_added = _parse_value_added(wb["增值"]) if "增值" in wb.sheetnames else []
    all_rows = shipments + pickup

    if not all_rows:
        raise ValueError("未识别到任何物流发货明细")

    shipping_amount = sum((_d(row["fee"]) for row in shipments), Decimal("0"))
    pickup_amount = sum((_d(row["fee"]) for row in pickup), Decimal("0"))
    value_added_amount = sum((_d(row["amount"]) for row in value_added), Decimal("0"))
    direct_amount = shipping_amount + pickup_amount
    computed_gross = direct_amount + value_added_amount

    summary_gross, adjustments, final_amount, summary_warehouse = _summary_amounts(wb["汇总"])
    adjustment_total = sum((_d(row["amount"]) for row in adjustments), Decimal("0"))
    net_amount = final_amount if final_amount is not None else summary_gross + adjustment_total
    if net_amount <= 0:
        net_amount = computed_gross + adjustment_total

    datetimes = [
        datetime.fromisoformat(row["completedAt"])
        for row in all_rows
        if row.get("completedAt")
    ]
    period_start = min(datetimes).date() if datetimes else None
    period_end = max(datetimes).date() if datetimes else None
    if period_start and period_end:
        period_label = f"{period_start:%Y-%m-%d} ~ {period_end:%Y-%m-%d}"
    else:
        period_label = _text(wb["汇总"]["A1"].value) or "物流账单"

    detail_warehouses = Counter(_text(row.get("warehouse")) for row in all_rows if _text(row.get("warehouse")))
    dominant_warehouse = detail_warehouses.most_common(1)[0][0] if detail_warehouses else ""
    carrier_counts = Counter(_text(row.get("carrier")) for row in all_rows if _text(row.get("carrier")))
    dominant_carrier = carrier_counts.most_common(1)[0][0] if carrier_counts else ""
    if len(carrier_counts) == 1:
        carrier_label = dominant_carrier
    elif carrier_counts:
        carrier_label = f"{dominant_carrier} 等 {len(carrier_counts)} 个渠道"
    else:
        carrier_label = "未识别物流渠道"

    match = _match_shipments(db, all_rows)
    warnings: list[str] = []
    if summary_warehouse and dominant_warehouse and summary_warehouse != dominant_warehouse:
        warnings.append(f"汇总仓库“{summary_warehouse}”与明细主要仓库“{dominant_warehouse}”不一致")
    if summary_gross > 0 and abs(summary_gross - computed_gross) > Decimal("0.01"):
        warnings.append(f"汇总金额 {_money_str(summary_gross)} 与明细计算 {_money_str(computed_gross)} 不一致")
    if not regular_quote and shipments:
        warnings.append("未识别到“三通一达”报价表，普通快递费用无法按报价复核")
    if match["unmatched"] > 0:
        warnings.append(f"有 {match['unmatched']} 条运单暂未匹配到本地吉客云发货事实")

    overhead_factor = (net_amount / direct_amount) if direct_amount > 0 else Decimal("1")
    preview = {
        "fileName": filename,
        "fileHash": digest,
        "duplicateBillId": duplicate_bill_id,
        "periodLabel": period_label,
        "periodStart": period_start.isoformat() if period_start else None,
        "periodEnd": period_end.isoformat() if period_end else None,
        "summaryWarehouse": summary_warehouse,
        "detailWarehouse": dominant_warehouse,
        "carrier": carrier_label,
        "carriers": [{"name": name, "count": count} for name, count in carrier_counts.most_common()],
        "shipmentCount": len(shipments),
        "pickupCount": len(pickup),
        "waybillCount": len(all_rows),
        "shippingAmount": _money_str(shipping_amount),
        "pickupAmount": _money_str(pickup_amount),
        "valueAddedAmount": _money_str(value_added_amount),
        "grossAmount": _money_str(summary_gross if summary_gross > 0 else computed_gross),
        "adjustmentAmount": _money_str(adjustment_total),
        "actualAmount": _money_str(net_amount),
        "directChargeAmount": _money_str(direct_amount),
        "overheadFactor": str(overhead_factor.quantize(Decimal("0.0001"))),
        "matchedCount": match["matched"],
        "unmatchedCount": match["unmatched"],
        "duplicateCount": match["duplicates"],
        "abnormalCount": 0,
        "warnings": warnings,
        "regionalModels": _regional_preview(all_rows, net_amount, direct_amount),
        "sampleRows": all_rows[:20],
    }

    if not persist:
        return {"ok": True, "preview": preview}

    start_dt = datetime.combine(period_start, datetime.min.time()) if period_start else None
    end_dt = datetime.combine(period_end, datetime.max.time()) if period_end else None
    unit_price = logistics_service.default_unit_price(db)
    estimated = logistics_service._estimate_for_period(db, start_dt, end_dt, unit_price)
    raw = {
        "source": "xlsx",
        "format": "warehouse_logistics_bill_v1",
        "fileName": filename,
        "fileHash": digest,
        "summary": {k: preview[k] for k in (
            "summaryWarehouse", "detailWarehouse", "shippingAmount", "pickupAmount",
            "valueAddedAmount", "grossAmount", "adjustmentAmount", "actualAmount",
            "directChargeAmount", "overheadFactor",
        )},
        "shipments": shipments,
        "pickup": pickup,
        "valueAdded": value_added,
        "adjustments": adjustments,
    }
    bill = LogisticsBill(
        period_label=period_label,
        period_start=start_dt,
        period_end=end_dt,
        carrier=carrier_label,
        waybill_count=len(all_rows),
        estimated_amount=estimated,
        actual_amount=net_amount,
        status="pending",
        invoice_status="none",
        attachment_name=filename,
        note="Excel 物流账单自动导入",
        matched_count=match["matched"],
        unmatched_count=match["unmatched"],
        duplicate_count=match["duplicates"],
        abnormal_count=0,
        raw=raw,
    )
    db.add(bill)
    db.commit()
    db.refresh(bill)
    return {"ok": True, "bill": logistics_service._bill_dict(bill), "preview": preview}
