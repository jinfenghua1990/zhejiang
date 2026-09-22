"""一次性导入：吉客云「采购入库申请单」导出 XLSX → jackyun_goods_documents。

幂等 upsert（按 (document_type='inbound', goodsdoc_no=申请单号)），可安全重跑：
- 主档：存在则更新，明细行整表重建（与 adapters/jackyun.py::_sync_goods_documents 同模式）
- API 同步链路不会写的维度（金额 / 批次效期 / 申请数量等）由本脚本从 Excel 写入

用法：
    cd backend && ./.venv/bin/python scripts/import_inbound_xlsx.py <path-to.xlsx>
"""
from __future__ import annotations

import sys
from collections import OrderedDict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook  # noqa: E402

from app.db import SessionLocal  # noqa: E402
from app.models.jackyun import JackyunGoodsDocument, JackyunGoodsDocumentItem  # noqa: E402


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _decimal(value: object) -> Decimal | None:
    text = _text(value)
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _datetime(value: object) -> datetime | None:
    text = _text(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


HEADER_TO_FIELD = {
    "货品编号": "goods_no",
    "货品名称": "goods_name",
    "规格": "spec",
    "条码": "sku_barcode",
    "单位": "unit_name",
    "申请数量": "apply_quantity",
    "入库数量": "quantity",
    "剩余数量": "remain_quantity",
    "退回数量": "return_quantity",
    "无税单价": "unit_price_notax",
    "含税单价": "unit_price_tax",
    "无税金额": "amount_notax",
    "含税金额": "amount_tax",
    "批次号": "batch_no",
    "生产批号": "production_lot",
    "生产日期": "production_date",
    "到期日期": "expiry_date",
    "保质期": "shelf_life",
    "保质期单位": "shelf_life_unit",
    "生产厂家": "manufacturer",
    "批准文号": "approval_no",
    "货品入库状态": "goods_status",
}

DOC_HEADERS = (
    "申请时间", "入库仓库", "入库类型", "审核状态", "往来单位", "物流单号",
    "入库金额", "入库费用", "入库原因", "费用分摊", "公司", "申请部门", "申请人",
    "标记", "预计入库日期", "通知状态", "审核人", "备注", "是否揽收", "关联单号",
    "原始单号", "跟踪单号", "创建时间", "修改时间", "外部业务单号", "通知失败原因",
)


def load_groups(path: str) -> "OrderedDict[str, list[dict]]":
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    headers = [_text(c) for c in rows[0]]
    idx = {h: i for i, h in enumerate(headers) if h}
    if "申请单号" not in idx:
        raise SystemExit("表头缺少「申请单号」，不是采购入库申请单导出文件")

    groups: "OrderedDict[str, list[dict]]" = OrderedDict()
    for row in rows[1:]:
        if not any(c not in (None, "") for c in row):
            continue
        doc_no = _text(row[idx["申请单号"]])
        if not doc_no:
            continue
        record: dict[str, object] = {}
        for header, i in idx.items():
            record[header] = row[i] if i < len(row) else None
        groups.setdefault(doc_no, []).append(record)
    return groups


def upsert_document(db, doc_no: str, records: list[dict]) -> str:
    head = records[0]
    row = (
        db.query(JackyunGoodsDocument)
        .filter_by(document_type="inbound", goodsdoc_no=doc_no)
        .first()
    )
    created = row is None
    if row is None:
        row = JackyunGoodsDocument(document_type="inbound", goodsdoc_no=doc_no)
        db.add(row)
        db.flush()

    row.document_at = _datetime(head.get("申请时间")) or _datetime(head.get("创建时间"))
    row.warehouse_name = _text(head.get("入库仓库"))
    row.supplier_name = _text(head.get("往来单位"))
    row.company_name = _text(head.get("公司"))
    row.total_amount = _decimal(head.get("入库金额"))
    row.total_fee = _decimal(head.get("入库费用"))
    # warehouse_code：Excel 无编码，保留原值（不覆盖既有回填）
    row.raw = {h: str(head.get(h) or "") for h in DOC_HEADERS}

    db.query(JackyunGoodsDocumentItem).filter_by(document_id=row.id).delete(synchronize_session=False)
    total = Decimal("0")
    has_quantity = False
    for line_no, rec in enumerate(records, start=1):
        item = JackyunGoodsDocumentItem(document_id=row.id, line_no=line_no)
        for header, field in HEADER_TO_FIELD.items():
            value = rec.get(header)
            if field in ("production_date", "expiry_date"):
                setattr(item, field, _datetime(value))
            elif field in (
                "apply_quantity", "remain_quantity", "return_quantity",
                "unit_price_tax", "unit_price_notax", "amount_tax", "amount_notax",
            ):
                setattr(item, field, _decimal(value))
            else:
                setattr(item, field, _text(value))
        quantity = _decimal(rec.get("入库数量"))
        if quantity is not None:
            total += quantity
            has_quantity = True
        item.raw = {h: str(rec.get(h) or "") for h in list(HEADER_TO_FIELD) + ["申请单号"]}
        db.add(item)
    row.total_quantity = total if has_quantity else None
    return "created" if created else "updated"


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("用法: import_inbound_xlsx.py <path-to.xlsx>")
    path = sys.argv[1]
    groups = load_groups(path)
    stats = {"created": 0, "updated": 0}
    db = SessionLocal()
    try:
        for doc_no, records in groups.items():
            result = upsert_document(db, doc_no, records)
            stats[result] += 1
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()
    print(f"完成：{stats['created']} 新建 / {stats['updated']} 更新，共 {len(groups)} 单，"
          f"{sum(len(v) for v in groups.values())} 行明细")


if __name__ == "__main__":
    main()
